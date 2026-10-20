# -*- coding: utf-8 -*-
"""Профиль оформления документа: снимается с размеченного PDF оригинала (docx_struct.load/layout) и задаёт сборку DOCX
(docx_build.Builder) и её проверку (check). Утверждено Максимом 15.10: «как будто писали по-русски в Word, взяв за
основу оформление оригинала» — шрифты ролей, выключка, интервалы, отступы номеров заголовков и пунктов, подписи,
таблицы, колонтитулы берутся из оригинала, а не зашиваются под один документ.

Правила (номера — как в отчёте МЕТРИКА_struct_docx.md):
  Ф1 поля листа            Ф6 межстрочный                 Ф11 поля ячеек
  Ф2 шрифт ролей           Ф7 отступы списков по уровням   Ф12 линия под шапкой листа
  Ф3 H1 с новой страницы   Ф8 рисунки и подписи            Ф13 абзацы по центру
  Ф4 номер → текст         Ф9 шапка таблицы                Ф14 номер пункта — только автонумерацией
  Ф5 выключка              Ф10 ширина таблиц               Ф15 значок в строке на месте"""
import os
import re
from collections import Counter
from statistics import median

from docx_struct import text_of
import docx_build as B

ROLES = ("body", "list", "td", "th", "caption", "code", "H1", "H2", "H3", "H4", "H5")


def _flat(blocks, ctx="body", depth=0):
    """(роль, блок, глубина списка) в порядке чтения."""
    for b in blocks:
        k = b["k"]
        if k == "h":
            yield f"H{min(b['level'], 5)}", b, depth
        elif k == "p":
            t = text_of(b["runs"]).strip()
            if B._figure(b):
                yield "figure", b, depth
            elif t and B.CAPTION.match(t):
                yield "cap_" + ("fig" if B.CAPTION.match(t).group(2).lower() == "ábra" else "tab"), b, depth
            elif t:
                code = B.is_code(b["runs"]) and ctx in ("body", "list")
                yield ("code" if code else "label" if ctx == "body" and B.is_label(b) else ctx), b, depth
        elif k == "list":
            for it in b["items"]:
                yield from _flat(it, "list", depth + 1)
        elif k == "table":
            yield "table", b, depth
            for r in b["rows"]:
                for c in r:
                    yield from _flat(c["blocks"], "th" if c["th"] else "td", depth)


def _list_firsts(blocks, depth=0, cell=False):
    """(вид списка, первый абзац пункта, глубина, в ячейке ли) по всем спискам, и в ячейках таблиц."""
    for b in blocks:
        if b["k"] == "list":
            for it in b["items"]:
                first = next((x for x in it if x["k"] == "p"), None)
                if first is not None:
                    yield b.get("kind"), first, depth + 1, cell
                yield from _list_firsts([x for x in it if x["k"] != "p"], depth + 1, cell)
        elif b["k"] == "table":
            for r in b["rows"]:
                for c in r:
                    yield from _list_firsts(c["blocks"], depth, True)


def _font(runs):
    c = Counter()
    for r in runs:
        if "t" in r and r["t"].strip():
            c[(r["font"], r["size"], r["b"])] += len(r["t"])
    return c


def _text_w(text, role):
    """Ширина строки шрифтом роли, pt: по файлу шрифта Windows, если он есть, иначе по средней ширине цифры."""
    import fitz
    name = role["font"].lower() + ("bd" if role["b"] else "")
    for f in (f"C:/Windows/Fonts/{name}.ttf", f"C:/Windows/Fonts/{role['font'].lower()}.ttf"):
        if not os.path.exists(f):        # о пропавшем файле PyMuPDF бросает FzErrorSystem, а не FileNotFoundError
            continue
        try:
            return fitz.Font(fontfile=f).text_length(text, role["size"])
        except (RuntimeError, FileNotFoundError, ValueError):   # шрифта нет (сервер Linux)
            continue
    return 0.58 * role["size"] * len(text)


def centered(g, L):
    """Абзац выключен по центру области текста (заголовок предупреждения, формула, подпись) — не во всю ширину."""
    if not g or g.get("x0") is None:
        return False
    left, right = L["left"], L["W"] - L["right"]
    mid = (g["x0"] + g["x1"]) / 2
    # и заметно короче полной строки: строка абзаца с отступом во всю ширину — не по центру (3PR42, OCR, 16.10)
    return abs(mid - (left + right) / 2) < 8 and g["x0"] - left > 24 and g["x1"] - g["x0"] < 0.85 * (right - left)


def extract(path, blocks, L):
    """Профиль оформления по оригиналу: блоки и раскладка из docx_struct. Значения в pt."""
    import pdfplumber
    pdf = pdfplumber.open(path)
    seen = {}

    def page(n):
        """(строки, слова) листа n (с 1): лист разбирается один раз и сразу освобождается — pdfplumber держит
        разобранные объекты каждого прочитанного листа до закрытия файла (у 3PR42 текст нарисован контурами: ~3000
        путей и ~50 МБ на лист, на 146 листах — 6,4 ГБ, 19.10)."""
        if n not in seen:
            pg = pdf.pages[n - 1]
            seen[n] = (pg.extract_text_lines(), pg.extract_words(x_tolerance=1.5))
            pg.close()
        return seen[n]
    items = list(_flat(blocks))
    fonts = {r: Counter() for r in ROLES}
    for role, b, _ in items:
        if role in fonts:
            fonts[role] += _font(b["runs"])
    P = {"page": {k: round(L[k], 1) for k in ("W", "H", "left", "right", "top", "bottom")}}
    P["roles"] = {}
    for r, c in fonts.items():
        if c:
            (f, s, bold), _ = c.most_common(1)[0]
            P["roles"][r] = {"font": f, "size": s, "b": bold}
    # выключка: длинные строки тела кончаются у одного правого края (по ширине) или вразброс (влево); тот же край —
    # правое поле листа (у таблиц рамка выступает за поле на поле ячейки, по ним поле не мерить)
    body = P["roles"].get("body", {"size": 10})
    s0 = body["size"]
    # край меряется по строкам самих абзацев тела и списков (кроме последней строки абзаца — она короче и при
    # выключке по ширине) с любого листа, кроме титула: строк ровно от поля с 3-го листа в коротком документе меньше
    # десятка, а пункт списка с отступом 18 pt не считался ни телом, ни списком — выходило «влево» (синтетика, 16.10)
    x1s = {"body": [], "list": []}
    width = L["W"] - L["left"] - L["right"]
    for role, b, _ in items:
        g = b.get("geom") or {}
        if role not in x1s or not g.get("page") or (g.get("lines") or 1) < 2 \
                or g["page"] == 1 and not L.get("first_is_body"):
            continue
        own = sorted((ln for ln in page(g["page"])[0] if g["top"] - 2 <= float(ln["top"]) <= g["bottom"]
                      and float(ln["x0"]) >= g["x0"] - 2 and float(ln["x1"]) <= g["x1"] + 2), key=lambda ln: float(ln["top"]))
        x1s[role] += [round(float(ln["x1"])) for ln in own[:-1] if float(ln["x1"]) - float(ln["x0"]) > 0.6 * width]
    step = []
    for n in range(3, len(pdf.pages) + 1):
        lines = [ln for ln in page(n)[0] if L["top"] < float(ln["top"]) < L["H"] - L["bottom"]]
        for a, b2 in zip(lines, lines[1:]):
            d = float(b2["top"]) - float(a["top"])
            if 1.1 * s0 < d < 1.6 * s0:
                step.append(round(d, 1))
    P["align"] = {}
    for role, v in x1s.items():
        c = Counter(v)
        edge, n = c.most_common(1)[0] if c else (None, 0)
        near = sum(k for x, k in c.items() if edge is not None and abs(x - edge) <= 1)
        P["align"][role] = "both" if len(v) >= 4 and near > 0.5 * len(v) else "left"
        if role == "body" and P["align"][role] == "both":
            P["page"]["right"] = round(L["W"] - edge, 1)
    P["line"] = Counter(step).most_common(1)[0][0] if step else round(s0 * 1.39, 1)
    # интервалы: зазор между соседними абзацами минус межстрочный зазор
    lead = P["line"] - body["size"]
    gaps = {"body": [], "list": [], "label_before": [], "label_after": [],
            **{f"H{i}_before": [] for i in range(1, 6)}, **{f"H{i}_after": [] for i in range(1, 6)}}
    seq = [(r, b) for r, b, _ in items if r in ("body", "list", "label") or r.startswith("H")]
    for (r1, a), (r2, b2) in zip(seq, seq[1:]):
        ga, gb = a.get("geom") or {}, b2.get("geom") or {}
        if not ga.get("bottom") or not gb.get("top") or ga.get("page") != gb.get("page"):
            continue
        gap = gb["top"] - ga["bottom"] - lead
        if gap < -2 or gap > 40:
            continue
        if r1 == r2 and r1 in ("body", "list"):
            gaps[r1].append(gap)
        if r2.startswith("H") and r1 in ("body", "list"):
            gaps[r2 + "_before"].append(gap)
        if r1.startswith("H") and r2 in ("body", "list"):
            gaps[r1 + "_after"].append(gap)
        if r2 == "label" and r1 in ("body", "list"):
            gaps["label_before"].append(gap)
        if r1 == "label" and r2 in ("body", "list"):
            gaps["label_after"].append(gap)
    P["space"] = {k: round(max(0.0, median(v)), 1) for k, v in gaps.items() if len(v) >= 3}
    # отступ первой строки: у большинства многострочных абзацев тела первая строка правее остальных (методика
    # МТ-ЦАСУТП-14: 36 pt) — абзацным отступом стиля, иначе он пропадал; хватает четырёх таких абзацев (в коротком
    # документе их восемь — при пороге в десять отступ терялся, синтетика 16.10), а двух-трёх — если отступ у всех
    offs = [b["geom"]["first_x0"] - b["geom"]["x0"] for r, b, _ in items if r == "body" and b.get("geom")
            and (b["geom"].get("lines") or 1) >= 2 and b["geom"].get("first_x0") is not None]
    pos = [o for o in offs if o > 5]
    if len(offs) >= 4 and len(pos) > 0.5 * len(offs) or len(offs) >= 2 and len(pos) == len(offs):
        P["first_line"] = round(median(pos), 1)
    # заголовки: с новой страницы, расстояние номер → текст
    P["heads"] = {}
    # «с нового листа» — заголовок, первый на своём листе (а не в 24 pt от верха тела: «2.1.» в 20 pt под «2.» на
    # том же листе считался первым, и весь уровень получал разрыв листа, ÜFK 16.10)
    page_first = {}
    for _, b, _ in items:
        g = b.get("geom") or {}
        if g.get("page") and g.get("top") is not None:
            page_first[g["page"]] = min(page_first.get(g["page"], g["top"]), g["top"])
    for lv in range(1, 6):
        hs = [b for r, b, _ in items if r == f"H{lv}" and b.get("geom") and b["geom"].get("page")]
        if not hs:
            continue
        tabs = []
        for b in hs[:20]:
            g = b["geom"]
            ws = sorted((w for w in page(g["page"])[1] if abs(float(w["top"]) - g["top"]) < 3), key=lambda w: w["x0"])
            if len(ws) >= 2 and re.fullmatch(r"\d+(\.\d+)*\.?", ws[0]["text"]):
                tabs.append(float(ws[1]["x0"]) - float(ws[0]["x0"]))
        first = sum(1 for b in hs if abs(b["geom"]["top"] - page_first.get(b["geom"]["page"], -99)) < 2)
        tab = round(median(tabs), 1) if tabs else None
        # самый длинный номер уровня («13.2» кеглем 14 — 27 pt при табуляции 28,8): текст не ближе 4 pt к номеру,
        # иначе в Word «13.2SYS TEC» читается слитно (в оригинале зазор 1,5 pt)
        role = P["roles"].get(f"H{lv}") or {"font": "Tahoma", "size": 12, "b": False}
        nums = [m[1] for b in hs for m in [re.match(r"\s*(\d+(?:\.\d+)*)", text_of(b["runs"]))] if m]
        if tab and nums:
            tab = max(tab, round(max(_text_w(n, role) for n in nums) + 4, 1))
        P["heads"][str(lv)] = {"page_break": first == len(hs) and len(hs) >= 2, "tab": tab}
    # списки: маркер и текст от поля по уровням — и по виду списка: у ÜFK пункты «a.)» — от поля, а «–» — с отступом,
    # и буквенный список брал отступ тире (17.10)
    lv_x, kind_x = {}, {}
    for kind, b, depth, cell in _list_firsts(blocks):
        if not (b.get("geom") or {}).get("page"):
            continue
        g = b["geom"]
        ws = sorted((w for w in page(g["page"])[1] if abs(float(w["top"]) - g["top"]) < 3
                     and float(w["x0"]) >= g["first_x0"] - 1), key=lambda w: w["x0"])
        # маркер — первое слово строки пункта («•», «1.», «a.», «a.)», «(a)»), текст — следующее
        if len(ws) >= 2 and abs(float(ws[0]["x0"]) - g["first_x0"]) < 2 \
                and re.fullmatch(r"[•▪◦·o\-–\uf0a7\uf0b7]|\(?(?:\d+|[a-zA-Z])\.?\)|\d+\.|[a-zA-Z]\.", ws[0]["text"]):
            x = (float(ws[0]["x0"]) - L["left"], float(ws[1]["x0"]) - L["left"])
            if not cell:        # в ячейке отступ — от края столбца, не от поля (ТЭ357: «•» почти все в таблицах)
                lv_x.setdefault(depth - 1, []).append(x)
                kind_x.setdefault(kind, {}).setdefault(depth - 1, []).append(x)
    mid = lambda v: {"marker": round(median(a for a, _ in v), 1), "text": round(median(t for _, t in v), 1)}
    P["lists"] = {str(lv): mid(v) for lv, v in sorted(lv_x.items()) if len(v) >= 2}
    P["lists_kind"] = {k: {str(lv): mid(v) for lv, v in sorted(d.items()) if len(v) >= 2} for k, d in kind_x.items()}
    # рисунки и подписи
    fig = [b for r, b, _ in items if r == "figure"]
    caps = [b for r, b, _ in items if r.startswith("cap_")]
    order = [r for r, _, _ in items]
    P["figure"] = {"center": sum(centered(b.get("geom"), L) or abs(((b.get("geom") or {}).get("x0", 0) + (b.get("geom") or {}).get("x1", 0)) / 2
                                                                       - (L["left"] + L["W"] - L["right"]) / 2) < 15 for b in fig) > len(fig) / 2}
    P["caption"] = {"center": sum(centered(b.get("geom"), L) for b in caps) > len(caps) / 2 if caps else True,
                    "fig_below": sum(1 for a, b2 in zip(order, order[1:]) if a == "figure" and b2 == "cap_fig") >= order.count("cap_fig") / 2,
                    "tab_above": sum(1 for a, b2 in zip(order, order[1:]) if a == "cap_tab" and b2 == "table") >= order.count("cap_tab") / 2}
    # таблицы: шапка, линии, поля ячеек (по таблицам — в блоке frame; здесь — обычное значение документа)
    frames = [b.get("frame") or {} for r, b, _ in items if r == "table"]
    def med(key):
        v = [f[key] for f in frames if key in f]
        return round(median(v), 1) if v else None
    P["table"] = {"pad_left": med("pad_left"), "pad_top": med("pad_top"), "pad_bottom": med("pad_bottom"),
                  "fill": Counter(f["fill"] for f in frames if f.get("fill")).most_common(1)[0][0]
                  if any(f.get("fill") for f in frames) else None,
                  "line": Counter(f["line"] for f in frames if f.get("line")).most_common(1)[0][0]
                  if any(f.get("line") for f in frames) else "000000"}
    pdf.close()
    return P


# ---------------------------------------------------------------- проверка DOCX против профиля
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
# номер пункта ещё и текстом («1. …», «a) …»); инициалы «Ф. И. О. оператора» — не номер
ITEM_NUM_TEXT = re.compile(r"^\s*(\d+|[A-Za-zА-Яа-я])[.)]\s(?![A-Za-zА-Яа-я][.)]\s)")


def check(docx_path, P, blocks, L):
    """Нарушения оформления DOCX против профиля: [{"rule", "where", "was", "need"}]. Плюс «прямое оформление»
    (метрика 2): прямой кегль и отступы на абзацах вне титула и пунктов списка."""
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(docx_path)
    out = []
    if (P.get("page") or {}).get("right"):
        L = dict(L, right=P["page"]["right"])               # как в сборке (Builder.build)

    def bad(rule, where, was, need):
        out.append({"rule": rule, "where": where, "was": str(was), "need": str(need)})

    st = doc.styles
    sec = doc.sections[-1]
    for k, v in (("left", sec.left_margin), ("right", sec.right_margin)):
        if abs(v.pt - P["page"][k]) > 1.5:
            bad("Ф1", f"поле {k}", round(v.pt, 1), P["page"][k])
    role_style = {"body": "Normal", "caption": "Caption UTS", "code": "Code UTS", "td": "Table Text",
                  **{f"H{i}": f"Heading {i}" for i in range(1, 6)}}

    def eff(s, attr):
        while s is not None:
            v = getattr(s.font, attr) if attr in ("name", "bold") else (s.font.size.pt if s.font.size else None)
            if v is not None:
                return v
            s = s.base_style
        return None
    for role, sname in role_style.items():
        want = P["roles"].get(role)
        if not want:
            continue
        s = st[sname]
        got = (eff(s, "name"), eff(s, "size"), bool(eff(s, "bold")))
        if got[0] != want["font"] or abs((got[1] or 0) - want["size"]) > 0.25 or got[2] != want["b"]:
            bad("Ф2", f"стиль «{sname}»", "%s %s%s" % (got[0], got[1], " ж" if got[2] else ""),
                "%s %s%s" % (want["font"], want["size"], " ж" if want["b"] else ""))
    h = P["heads"].get("1")
    if h and bool(st["Heading 1"].paragraph_format.page_break_before) != h["page_break"]:
        bad("Ф3", "стиль «Heading 1»", st["Heading 1"].paragraph_format.page_break_before, h["page_break"])
    num = doc.part.numbering_part.element
    an = next((a for a in num.findall(f"{W}abstractNum") if a.get(qn("w:abstractNumId")) == "900"), None)
    for lv, hp in P["heads"].items():
        if not hp.get("tab") or an is None:
            continue
        lvl = next((x for x in an.findall(f"{W}lvl") if x.get(qn("w:ilvl")) == str(int(lv) - 1)), None)
        ind = lvl.find(f"{W}pPr/{W}ind") if lvl is not None else None
        got = int(ind.get(qn("w:hanging"))) / 20 if ind is not None else None
        if got is None or abs(got - hp["tab"]) > 1.5:
            bad("Ф4", f"заголовок {lv}: номер → текст", got, hp["tab"])
    jc = {"both": 3, "left": None}
    for role, sname in (("body", "Normal"), ("list", "List UTS")):
        want = P["align"].get(role, "left")
        got = st[sname].paragraph_format.alignment
        got = 3 if got is not None and int(got) == 3 else None
        if got != jc[want]:
            bad("Ф5", f"выравнивание «{sname}»", "по ширине" if got == 3 else "влево", "по ширине" if want == "both" else "влево")
    ls = st["Normal"].paragraph_format.line_spacing
    got = ls.pt if hasattr(ls, "pt") else None
    if got is None or abs(got - P["line"]) > 0.6:
        bad("Ф6", "межстрочный «Normal»", ls, f"не меньше {P['line']} pt")
    # Ф7: отступы пунктов — по оригиналу, одним значением (договор правил, шаг 3): у каждого пункта отступ абзаца тот
    # же, что у его уровня в нумерации, а нумерация — из источника сборки (docx_build.item_ind: место списка на листе,
    # иначе профиль с местом под метку). Прежде проверка видела только нумерацию (а место по листу ставилось только в
    # абзац) и сверяла её с профилем без места под метку — ложные нарушения
    abst = {a.get(qn("w:abstractNumId")): a for a in num.findall(f"{W}abstractNum")}
    of_num = {n.get(qn("w:numId")): n.find(f"{W}abstractNumId").get(qn("w:val")) for n in num.findall(f"{W}num")}

    def num_ind(nid, ilvl):
        a = abst.get(of_num.get(nid))
        lvl = next((x for x in a.findall(f"{W}lvl") if x.get(qn("w:ilvl")) == ilvl), None) if a is not None else None
        ind = lvl.find(f"{W}pPr/{W}ind") if lvl is not None else None
        if ind is None:
            return None
        text = int(ind.get(qn("w:left"))) / 20
        return round(text - int(ind.get(qn("w:hanging")) or 0) / 20, 1), round(text, 1)
    want = set()                                            # отступы пунктов тела по источнику сборки

    def lists(bs, level=None):
        for b in bs:
            if b["k"] == "list":
                lv = b.get("level", 0) if level is None else level
                want.add(tuple(round(v, 1) for v in B.item_ind(b, lv, P, L["left"])))
                for it in b["items"]:
                    lists(it, lv + 1)
    lists(blocks)
    for p in doc.element.body.iter(f"{W}p"):
        np_ = p.find(f"{W}pPr/{W}numPr")
        nid = np_.find(f"{W}numId") if np_ is not None else None
        if nid is None or int(nid.get(qn("w:val"))) < 1000:
            continue
        il = np_.find(f"{W}ilvl")
        got = num_ind(nid.get(qn("w:val")), il.get(qn("w:val")) if il is not None else "0")
        if got is None:
            continue
        t = "".join(e.text or "" for e in p.iter(f"{W}t"))[:40]
        ind = p.find(f"{W}pPr/{W}ind")
        if ind is not None and ind.get(qn("w:left")) is not None:
            left = int(ind.get(qn("w:left"))) / 20
            own = (round(left + (int(ind.get(qn("w:firstLine")) or 0) - int(ind.get(qn("w:hanging")) or 0)) / 20, 1),
                   round(left, 1))
            if abs(own[0] - got[0]) > 0.5 or abs(own[1] - got[1]) > 0.5:
                bad("Ф7", t, f"абзац {own[0]}/{own[1]}", f"как в нумерации {got[0]}/{got[1]}")
        in_cell = any(a.tag == f"{W}tc" for a in p.iterancestors())
        if not in_cell and want and all(abs(got[0] - w[0]) > 2 or abs(got[1] - w[1]) > 2 for w in want):
            bad("Ф7", t, f"{got[0]}/{got[1]}", "по месту списка на листе или по профилю")
    if P["figure"]["center"] and int(st["Figure UTS"].paragraph_format.alignment or 0) != 1:
        bad("Ф8", "стиль рисунка", "не по центру", "по центру")
    if P["caption"]["center"] and int(st["Caption UTS"].paragraph_format.alignment or 0) != 1:
        bad("Ф8", "стиль подписи", "не по центру", "по центру")

    body = doc.element.body
    paras = list(body.iter(f"{W}p"))
    sid = {s.style_id: s.name for s in doc.styles}

    def pst(p):
        ps = p.find(f"{W}pPr/{W}pStyle")
        return sid.get(ps.get(qn("w:val")), "?") if ps is not None else "Normal"

    def vt(p):
        return "".join((e.text or "") if e.tag == f"{W}t" else "\t" for e in p.iter(f"{W}t", f"{W}tab"))
    # подпись рисунка — сразу под рисунком, таблицы — над таблицей
    seq = [e for e in body if e.tag in (f"{W}p", f"{W}tbl")]
    for i, e in enumerate(seq):
        if e.tag != f"{W}p" or pst(e) != "Caption UTS":
            continue
        t = vt(e)
        prev, nxt = (seq[i - 1] if i else None), (seq[i + 1] if i + 1 < len(seq) else None)
        if t.startswith("Рисунок") and P["caption"]["fig_below"] and not (prev is not None and prev.tag == f"{W}p" and pst(prev) == "Figure UTS"):
            bad("Ф8", t[:40], "подпись не под рисунком", "под рисунком")
        if t.startswith("Таблица") and P["caption"]["tab_above"] and not (nxt is not None and nxt.tag == f"{W}tbl"):
            bad("Ф8", t[:40], "подпись не над таблицей", "над таблицей")
    # таблицы: ширина, заливка шапки, шапка не отрывается
    tabs = [b for r, b, _ in _flat(blocks) if r == "table"]
    dts = [t for t in body.iter(f"{W}tbl") if not any(a.tag == f"{W}tbl" for a in t.iterancestors())]
    text_w = L["W"] - L["left"] - L["right"]
    for n, (b, t) in enumerate(zip(tabs, dts), 1):
        xs = (b.get("frame") or {}).get("xs") or []
        grid = sum(int(c.get(qn("w:w"))) for c in t.findall(f"{W}tblGrid/{W}gridCol")) / 20
        tcw = [sum(int(w.get(qn("w:w"))) for w in tr.iter(f"{W}tcW")) / 20 for tr in t.findall(f"{W}tr")]
        width = max(tcw) if tcw and max(tcw) else grid
        if len(xs) >= 2 and abs(width - (xs[-1] - xs[0])) > 0.05 * text_w:
            bad("Ф10", f"таблица {n}", f"{round(width)} pt", f"{round(xs[-1] - xs[0])} pt")
        fill = (b.get("frame") or {}).get("fill")
        first = t.find(f"{W}tr")
        shd = first.find(f".//{W}shd") if first is not None else None
        got = shd.get(qn("w:fill")) if shd is not None else None
        if (fill or None) != (got if got not in (None, "auto") else None) and b.get("borders", True):
            bad("Ф9", f"таблица {n}: заливка шапки", got, fill)
        hp = [p for p in (first.findall(f".//{W}p") if first is not None else []) if vt(p).strip()]
        if len(t.findall(f"{W}tr")) > 1 and hp and all(c["th"] for c in (b["rows"][0] if b["rows"] else [])) \
                and not all(pst(p) in ("Table Head UTS", "Table Head Small") for p in hp):
            bad("Ф9", f"таблица {n}: шапка", "может оторваться от тела", "«не отрывать от следующего»")
    # линия шапки листа — под шапкой (абзац с нижней границей после таблицы шапки)
    for s in doc.sections[1:2]:
        hdr = s.header._element
        kids = [k for k in hdr if k.tag in (f"{W}p", f"{W}tbl")]
        bpos = [i for i, k in enumerate(kids) if k.tag == f"{W}p" and k.find(f".//{W}pBdr/{W}bottom") is not None]
        tpos = [i for i, k in enumerate(kids) if k.tag == f"{W}tbl"]
        if bpos and tpos and bpos[0] < tpos[0]:
            bad("Ф12", "шапка листа", "линия над шапкой", "линия под шапкой")
    # абзацы по центру в оригинале → по центру
    want_c = sum(1 for r, b, d in _flat(B.merge_inline(blocks)) if r in ("body", "code") and d == 0
                 and centered(b.get("geom"), L) and (b["geom"].get("page") or 0) >= 2)
    got_c = sum(1 for p in paras if pst(p) in ("Center UTS", "Code Center UTS"))
    if got_c < want_c:
        bad("Ф13", "абзацы по центру", got_c, want_c)
    # номер пункта текстом, пустые скобки от значка
    for p in paras:
        np_ = p.find(f"{W}pPr/{W}numPr/{W}numId")
        t = vt(p)
        if np_ is not None and int(np_.get(qn("w:val"))) >= 1000 and ITEM_NUM_TEXT.match(t):
            bad("Ф14", t[:50], "номер ещё и текстом", "только автонумерация")
        # значок между скобками — «(\0)»; пустые скобки без значка — значок уехал
        seq = "".join("\0" if e.tag == f"{W}drawing" else (e.text or "") for e in p.iter(f"{W}t", f"{W}drawing"))
        if re.search(r"\(\s*\)", seq):
            bad("Ф15", t[:50], "пустые скобки: значок не на месте", "значок в скобках")
    # прямое оформление (метрика 2)
    title_end = next((i for i, p in enumerate(paras) if p.find(f"{W}pPr/{W}sectPr") is not None), -1)
    for i, p in enumerate(paras):
        if i <= title_end:
            continue
        for r in p.iter(f"{W}r"):
            if r.find(f"{W}rPr/{W}sz") is not None:
                bad("М2", vt(p)[:40], "прямой кегль", "кегль стиля")
                break
        ind = p.find(f"{W}pPr/{W}ind")
        # пункты списка (и в ячейках) несут отступ своего списка прямо — то же значение, что в нумерации (Ф7)
        if ind is not None and pst(p) not in ("List UTS", "List Cell UTS"):
            bad("М2", vt(p)[:40], "прямой отступ", "отступ стиля")
    return out

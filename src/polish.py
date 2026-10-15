# -*- coding: utf-8 -*-
"""Цикл шлифовки перевода по языковым соглашениям (утверждён Максимом 15.10):
проверка → правка кодом всего детерминированного (conventions.fix_code) → правка порталом только формулировок
(перепереводятся лишь задетые сегменты, нарушенное правило — в задании) → повторная проверка. Не больше трёх проходов;
стоп, если нарушений нет или проход ничего не изменил. Каждая правка — строка журнала «было → стало, правило, кто»
(лист «шлифовка» таблицы вычитки).

items: {ключ: {"hu", "ru", "kind": "text" | "shot", "page", "id"}} — ru правится на месте."""
import json
import re

import conventions as C

PROMPT = """Ты — редактор русского перевода эксплуатационной документации АЭС с венгерского языка.
Для каждого сегмента даны оригинал (hu), текущий перевод (ru) и нарушенные правила (rules). Исправь перевод так, чтобы
правила выполнялись. Меняй только то, чего требуют правила, и согласуй с правкой соседние слова (род, падеж);
остальной текст оставь слово в слово. Метки <f1>…</f1> сохраняй на тех же местах; числа, коды, имена файлов, команды —
без изменений.
Ответ — только JSON-массив [{"id": …, "ru": "…"}] без пояснений."""

HINT = {
    "Я2": "ссылку на рисунок пиши «рисунок N» в нужном падеже (не «рис. N»)",
    "Я3": "ссылку на раздел пиши «раздел N» в нужном падеже, в скобках — «(см. раздел N)»; не «глава»",
    "Я8": "предупреждение начинай словом «ВНИМАНИЕ!»",
    "Я11": "переведи оставшиеся венгерские слова; имена-заполнители (PARAMÉTERNÉV) — по-русски, тоже прописными",
}


def _hint(rule, why):
    if rule == "Я5":
        name = re.search(r"«(.+?)»", why)[1]
        return f"«{name}» — название, пиши латиницей «{name}»; не переводи, не расшифровывай и не заменяй сокращением"
    if rule == "Я6":
        return f"{why.split(' — ')[0]} — код или имя файла, перенеси дословно"
    if rule == "Я10":
        return "термин документа: " + why
    return HINT.get(rule.split("/")[0], why)


def _numbers(t):
    return sorted(re.findall(r"\d+(?:[.,]\d+)*", C.TAGS.sub(" ", t)))


def _ok(old, new, hu, rules, kind):
    """Ответ портала принят: метки и числа те же, длина в разумных пределах, нарушений стало меньше и новых нет.
    Длина — против старого перевода; при Я5 — против оригинала, шире: сворачивание раскрытого названия укорачивает
    перевод в разы («Включите ЭВМ системы оперативной информации!» → «Включите ЭВМ UTS!»)."""
    if not new or sorted(re.findall(r"</?f\d+>", old)) != sorted(re.findall(r"</?f\d+>", new)):
        return False
    before = C.check(hu, old, rules, kind)
    if any(v[0] == "Я5" and not v[2] for v in before):
        base, lo, hi = C.TAGS.sub("", hu), 0.4, 2.0
    else:
        base, lo, hi = old, 0.6, 1.6
    if _numbers(old) != _numbers(new) or not lo < len(new) / max(1, len(base)) < hi:
        return False
    after = C.check(hu, new, rules, kind)
    return len([v for v in after if not v[2]]) < len([v for v in before if not v[2]]) and len(after) <= len(before)


def _portal(chunk, chat):
    """chunk: [(ключ, hu, ru, [подсказки])] → {ключ: ru}."""
    ids = {str(i + 1): k for i, (k, _, _, _) in enumerate(chunk)}
    data = [{"id": str(i + 1), "hu": C.TAGS.sub("", hu) if "<f" not in ru else hu, "ru": ru, "rules": hints}
            for i, (_, hu, ru, hints) in enumerate(chunk)]
    r = chat([{"role": "system", "content": PROMPT},
              {"role": "user", "content": json.dumps(data, ensure_ascii=False, indent=0)}],
             max_tokens=8000, think=False, tag="polish")
    txt = (r.get("text") or "").strip()
    m = re.search(r"\[.*\]", txt, re.S)
    try:
        got = json.loads(m[0]) if m else []
    except json.JSONDecodeError:
        got = []
    return {ids[str(g.get("id"))]: g.get("ru") for g in got if isinstance(g, dict) and str(g.get("id")) in ids}


def run(items, rules, chat=None, max_passes=3, batch=6, log=print, progress=None):
    """Шлифовка items на месте. Возвращает (журнал правок, оставшиеся нарушения). progress(доля 0..1) — ход по
    запросам портала (конвейер показывает его в шаге «шлифовка»)."""
    edits, tries = [], {}

    def edit(k, it, rule, who, new, n):
        edits.append({"key": k, "id": it.get("id"), "page": it.get("page"), "rule": rule, "who": who, "pass": n,
                      "hu": C.TAGS.sub("", it["hu"]), "was": C.TAGS.sub("", it["ru"]), "now": C.TAGS.sub("", new)})
        it["ru"] = new

    for n in range(1, max_passes + 1):
        changed = 0
        for k, it in items.items():
            new, done = C.fix_code(it["hu"], it["ru"], rules, it["kind"])
            if done and new != it["ru"]:
                edit(k, it, ", ".join(done), "код", new, n)
                changed += 1
        todo = []
        for k, it in items.items():
            v = [x for x in C.check(it["hu"], it["ru"], rules, it["kind"]) if not x[2]]
            if v and chat and tries.get(k, 0) < 2:
                todo.append((k, it["hu"], it["ru"], sorted({_hint(r, why) for r, why, _ in v})))
        size = batch if n == 1 else 1                      # повтор — по одному: другой запрос, не ответ из кэша
        for i in range(0, len(todo), size):
            chunk = todo[i:i + size]
            got = _portal(chunk, chat)
            for k, hu, ru, _ in chunk:
                tries[k] = tries.get(k, 0) + 1
                new = got.get(k)
                it = items[k]
                if new and new != ru and _ok(ru, new, hu, rules, it["kind"]):
                    rules_hit = sorted({r for r, _, code in C.check(hu, ru, rules, it["kind"]) if not code})
                    edit(k, it, ", ".join(rules_hit), "портал", new, n)
                    changed += 1
            log(f"  шлифовка, проход {n}: портал {min(i + size, len(todo))}/{len(todo)}")
            if progress:
                progress((n - 1 + min(i + size, len(todo)) / len(todo)) / max_passes)
        left = sum(1 for it in items.values() if C.check(it["hu"], it["ru"], rules, it["kind"]))
        log(f"шлифовка, проход {n}: правок {changed}, сегментов с нарушениями {left}")
        if not changed or not left:
            break
    remaining = [{"key": k, "id": it.get("id"), "page": it.get("page"), "rule": r, "why": why,
                  "ru": C.TAGS.sub("", it["ru"])}
                 for k, it in items.items() for r, why, _ in C.check(it["hu"], it["ru"], rules, it["kind"])]
    return edits, remaining


def write_xlsx(out, edits, remaining, base=None, extra_rows=()):
    """Таблица вычитки + листы «шлифовка» и «нарушения» (add_sheets). base — исходная таблица вычитки (копируется),
    иначе новая книга."""
    from openpyxl import Workbook, load_workbook
    wb = load_workbook(base) if base else Workbook()
    if not base:
        wb.remove(wb.active)
    add_sheets(wb, edits, remaining, extra_rows)
    wb.save(out)


def add_sheets(wb, edits, remaining, extra_rows=()):
    """Листы «шлифовка» (каждая правка: было → стало, правило, кто) и «нарушения» (что осталось: языковые
    соглашения и оформление) в книгу wb (assemble.review_xlsx — в общую таблицу вычитки)."""
    from openpyxl.styles import Alignment, Font
    for name in ("шлифовка", "нарушения"):
        if name in wb.sheetnames:
            del wb[name]
    ws = wb.create_sheet("шлифовка")
    ws.append(["№", "стр.", "правило", "кто правил", "проход", "было", "стало", "венгерский (источник)"])
    for e in edits:
        ws.append([e.get("id") if e.get("id") is not None else "", e.get("page") or "", e["rule"], e["who"], e["pass"],
                   e["was"], e["now"], e["hu"]])
    for col, w in zip("ABCDEFGH", (7, 6, 12, 10, 7, 60, 60, 50)):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet("нарушения")
    ws2.append(["стр.", "правило", "где / пояснение", "текст"])
    for r in remaining:
        ws2.append([r.get("page") or "", r["rule"], C.NAMES.get(r["rule"].split("/")[0], "") + ": " + r["why"], r["ru"]])
    for row in extra_rows:
        ws2.append(list(row))
    for col, w in zip("ABCD", (6, 10, 50, 80)):
        ws2.column_dimensions[col].width = w
    for sh in (ws, ws2):
        for c in sh[1]:
            c.font = Font(bold=True)
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"

# -*- coding: utf-8 -*-
"""Возврат правок редактора: колонка «Исправление (редактор)» файла <док>_вычитка.xlsx →
перевод в JSON прогона (правка разносится по всем сегментам с тем же исходным текстом) →
пересборка DOCX/PDF/xlsx → память переводов проекта (венгерский → утверждённый русский).
Память подключается к следующим прогонам (project.json: tm): совпавшая фраза берётся без модели.
Вызывается из pipeline.edits (команда `translate_komplekt.py edits`)."""
import re, json, shutil
from pathlib import Path

from assemble import assemble_all

TAG = re.compile(r"</?f\d+>")


def norm_hu(s):
    return re.sub(r"\s+", " ", TAG.sub("", s)).strip()


def apply(jf, xf, tm_path, out_dir=None):
    """Правки колонки «Исправление (редактор)» из xf → JSON прогона jf, память переводов tm_path, пересборка.
    Возвращает список правок [(венгерский, утверждённый русский)] — для размножения по похожим фразам (rag)."""
    jf, xf = Path(jf), Path(xf)
    from openpyxl import load_workbook
    ws = load_workbook(xf).active
    head = [c.value for c in ws[1]]
    ci, cf = head.index("№"), head.index("Исправление (редактор)")
    fixes = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[cf] and str(row[cf]).strip():
            fixes[str(row[ci])] = str(row[cf]).strip()
    if not fixes:
        return []
    d = json.loads(jf.read_text(encoding="utf-8"))
    bak = jf.with_suffix(".before_review.json")
    if not bak.exists():
        shutil.copy(jf, bak)
    by_src = {}
    for s in d["segs"]:
        by_src.setdefault(norm_hu(s["text"]), []).append(str(s["id"]))
    src_of = {str(s["id"]): norm_hu(s["text"]) for s in d["segs"]}
    tmf = Path(tm_path)
    tm = json.loads(tmf.read_text(encoding="utf-8")) if tmf.exists() else {}
    n_seg = 0
    for sid, ru in fixes.items():
        hu = src_of.get(sid)
        if hu is None:
            continue
        for k in by_src[hu]:                       # одинаковая фраза — одинаковая правка
            r = d["res"][k]
            r.update({"ru": ru, "issues": [], "edited": True})
            for k2 in ("review", "autofix", "rag"):         # пометки о прежнем переводе
                r.pop(k2, None)
            n_seg += 1
        tm[hu] = ru
    jf.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
    tmf.write_text(json.dumps(tm, ensure_ascii=False, indent=0), encoding="utf-8")
    assemble_all(jf, d, out_dir)
    print(f"{jf.stem}: правок {len(fixes)} → сегментов {n_seg}; память переводов {len(tm)} фраз", flush=True)
    return [(src_of[sid], ru) for sid, ru in fixes.items() if sid in src_of]

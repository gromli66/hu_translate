# -*- coding: utf-8 -*-
"""PDF с текстовым слоем (руководство UTS, 15.10): страница со скриншотом не скан, невидимый текст не переводится,
маркер пункта отделён от текста, ячейка таблицы — одним абзацем, DOCX без квадратов вместо маркеров и без
обрезанных строк таблицы, надписи скриншота перерисовываются без порчи соседних пикселей."""
import io
import zipfile

import numpy as np
import fitz

import layout as L
import segments as S
import shots as SH
import docx_src as D

LOREM = ("Az UTS alkalmazás grafikus felhasználói felülete a következő fő részekre bontható a felhasználó szerint, "
         "amelyek a menüsor alatt helyezkednek el és a tesztelés során folyamatosan láthatók maradnak.")


def png(w, h, gray=200):
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, w, h), False)
    pix.clear_with(gray)
    return pix.tobytes("png")


def word(t, box, c=95):
    return {"t": t, "box": list(box), "h": box[3] - box[1], "c": c}


def test_screenshot_words_do_not_make_page_a_scan():
    doc = fitz.open(); page = doc.new_page()
    tw = fitz.TextWriter(page.rect)
    y = 60
    for _ in range(6):
        tw.fill_textbox(fitz.Rect(50, y, 550, y + 40), LOREM, fontsize=9); y += 45
    tw.write_text(page)
    page.insert_image(fitz.Rect(50, 400, 550, 600), stream=png(500, 200))
    words = [word("Eredmények", (60, 410, 110, 420)), word("Adminisztráció", (120, 410, 180, 420)),
             word("Tesztelés", (190, 410, 230, 420)), word("felhasználói", (60, 430, 120, 440))]
    words += [word(w, (50 + i * 10, 62, 58 + i * 10, 70)) for i, w in enumerate(LOREM.split()[:20])]
    share, outside = S.page_missing_share(page, words)
    assert share == 0.0
    assert all(w["box"][1] < 400 for w in outside)          # слова скриншота отброшены


def test_curve_text_still_counts_as_missing():
    doc = fitz.open(); page = doc.new_page()
    page.insert_text((50, 60), "Azonosító: 3PR42 " + " ".join(["Verziószám"] * 8), fontsize=9)
    words = [word("Azonosító", (50, 52, 90, 62))] + [word("blokkvezénylőben", (50 + i * 60, 100, 100 + i * 60, 110))
                                                      for i in range(5)]
    share, _ = S.page_missing_share(page, words)
    assert share > 0.15


def test_hidden_text_under_full_page_image_is_skipped():
    doc = fitz.open(); page = doc.new_page()
    page.insert_text((50, 800), "Létesítmény: Paksi Atomerőmű", fontsize=9)       # под фоном
    page.insert_image(page.rect, stream=png(60, 80, 240))
    page.insert_text((50, 100), "UTS felhasználói kézikönyv", fontsize=14)        # поверх фона
    texts = [f["t"] for f in L.frags_from_text(page)]
    assert texts == ["UTS felhasználói kézikönyv"]


def test_bullet_marker_is_own_fragment_and_text_is_item():
    doc = fitz.open(); page = doc.new_page()
    put = lambda x, y, t: page.insert_htmlbox(fitz.Rect(x, y, 560, y + 14), t, css="* {font-size: 10pt;}")
    put(80, 100, "•"); put(98, 100, "Alfanumerika: a vizsgálandó jel alfanumerikája"); put(98, 113, "folytatás")
    put(80, 140, "•"); put(98, 140, "Művelet: összehasonlító logikai művelet")
    hs, vs = L.rules(page)
    pars = L.group(L.frags_from_text(page, vs), hs, vs)
    texts = [p["text"] for p in pars]
    assert "•" in texts
    assert "Művelet: összehasonlító logikai művelet" in texts       # пункт не приклеился к предыдущему
    assert not any(t.startswith("•") and len(t) > 1 for t in texts)


def test_narrow_cell_lines_stay_one_paragraph():
    doc = fitz.open(); page = doc.new_page()
    for x in (60, 150):
        page.draw_line((x, 80), (x, 200))
    page.insert_text((65, 100), "Teszt", fontsize=10)
    page.insert_text((65, 113), "szüneteltetése", fontsize=10)
    hs, vs = L.rules(page)
    pars = L.group(L.frags_from_text(page, vs), hs, vs)
    assert [p["text"] for p in pars] == ["Teszt szüneteltetése"]


def _docx(body_xml):
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" ' \
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    doc = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W}><w:body>{body_xml}' \
          f'<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="400" w:bottom="500" w:left="800" ' \
          f'w:right="800"/></w:sectPr></w:body></w:document>'
    rels = '<?xml version="1.0" encoding="UTF-8"?><Relationships ' \
           'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'
    types = '<?xml version="1.0" encoding="UTF-8"?><Types ' \
            'xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", doc)
        z.writestr("word/_rels/document.xml.rels", rels)
        z.writestr("[Content_Types].xml", types)
    return buf.getvalue()


def test_docx_normalize_splits_bullets_and_unlocks_rows(tmp_path):
    sym = '<w:r><w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol"/></w:rPr><w:t>•</w:t></w:r>'
    para = ('<w:p><w:pPr><w:tabs><w:tab w:val="left" w:pos="922"/></w:tabs><w:ind w:left="562" w:firstLine="0"/></w:pPr>'
            f'{sym}<w:r><w:tab/></w:r><w:r><w:t>aktuális dátum</w:t></w:r><w:r><w:br/></w:r>'
            f'{sym}<w:r><w:tab/></w:r><w:r><w:t>szimulátor idő</w:t></w:r></w:p>')
    table = '<w:tbl><w:tr><w:trPr><w:trHeight w:hRule="exact" w:val="240"/></w:trPr><w:tc><w:p><w:r><w:t>x</w:t>' \
            '</w:r></w:p></w:tc></w:tr></w:tbl><w:p/>'
    f = tmp_path / "a.docx"
    f.write_bytes(_docx(para + table))
    D.normalize(f)
    ds = S.docx_segments(f)
    assert [d["src"].strip() for d in ds if d["src"].strip() not in ("x",)] == ["aktuális dátum", "szimulátor idő"]
    xml = zipfile.ZipFile(f).read("word/document.xml").decode("utf-8")
    assert 'w:hRule="atLeast"' in xml and 'w:hRule="exact"' not in xml
    assert 'w:hanging="360"' in xml
    D._fix_fonts(f)
    assert 'w:ascii="Symbol"' not in zipfile.ZipFile(f).read("word/document.xml").decode("utf-8")


def test_screenshot_label_repaint_keeps_neighbours():
    img = np.full((40, 200, 3), 120, np.uint8)            # серая шапка
    img[30:32, :] = 255                                   # белая линия под шапкой
    img[12:22, 20:70] = 250                               # «буквы» надписи
    doc = fitz.open(); page = doc.new_page()
    pix = fitz.Pixmap(fitz.csRGB, 200, 40, img.tobytes(), False)
    xref = page.insert_image(fitz.Rect(50, 50, 250, 90), pixmap=pix)
    out = SH.render(doc, xref, [({"box": [20, 12, 69, 21], "h": 9, "t": "Kategória"}, "Категория")])
    a = np.frombuffer(out.samples, np.uint8).reshape(out.height, out.width, 3)
    assert (a[30:32, :] == 255).all()                     # линия цела
    assert (np.abs(a[2:8, :].astype(int) - 120) <= 2).all()   # фон выше надписи не тронут
    assert (a[10:24, 15:120] > 200).any()                 # перевод нарисован светлым, как оригинал


# ---------------------------------------------------------------- находки красной команды 15.10
def _section_docx(tables_xml):
    """DOCX из нескольких «страниц»: у каждой — абзац текста, таблица и абзац-разрыв раздела."""
    sect = '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="400" w:bottom="500" w:left="800" w:right="800"/></w:sectPr>'
    body = ""
    for i, t in enumerate(tables_xml):
        body += f'<w:p><w:r><w:t>Szöveg {i}</w:t></w:r></w:p>{t}'
        if i < len(tables_xml) - 1:
            body += f'<w:p><w:pPr>{sect}</w:pPr></w:p>'
    return _docx(body)


def _tbl(text, h=240, rows=1):
    tr = f'<w:tr><w:trPr><w:trHeight w:hRule="exact" w:val="{h}"/></w:trPr><w:tc><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:tc></w:tr>'
    return f'<w:tbl>{tr * rows}</w:tbl>'


def test_stamp_goes_to_footer_but_repeated_step_table_does_not(tmp_path):
    stamp = [_tbl(f"Létesítmény: Paksi Atomerőmű Verzió, Lapszám: 1.10 {i}/99") for i in range(3)]
    f = tmp_path / "s.docx"; f.write_bytes(_section_docx(stamp)); D.normalize(f)
    z = zipfile.ZipFile(f)
    assert len([n for n in z.namelist() if n.startswith("word/footer")]) == 3
    assert "Létesítmény" not in z.read("word/document.xml").decode("utf-8")
    # таблица шагов TU: шапка одна, тело разное, высокая — остаётся в теле
    steps = [_tbl(f"Sorsz. Feltétel, művelet, ellenőrzés lépés {w}", h=3000, rows=3) for w in ("alfa", "béta", "gamma")]
    g = tmp_path / "t.docx"; g.write_bytes(_section_docx(steps)); D.normalize(g)
    z = zipfile.ZipFile(g)
    assert not [n for n in z.namelist() if n.startswith("word/footer")]
    assert "Sorsz." in z.read("word/document.xml").decode("utf-8")


def test_pages_follow_bookmarks_not_extra_sections(tmp_path):
    from docx import Document
    parts = []
    for i in range(3):
        d = Document(); d.add_paragraph(f"oldal {i} eleje")
        if i == 1:                                  # pdf2docx: лишний разрыв раздела внутри страницы
            d.add_section(); d.add_paragraph(f"oldal {i} vége")
        p = tmp_path / f"p{i}.docx"; d.save(p); parts.append(p)
    from assemble import _merge_docx
    out = tmp_path / "m.docx"; _merge_docx(parts, out)
    pages = {s["src"].strip(): s["page"] for s in D.paged_segments(out)}
    assert pages == {"oldal 0 eleje": 0, "oldal 1 eleje": 1, "oldal 1 vége": 1, "oldal 2 eleje": 2}


def test_watermark_over_whole_page_does_not_hide_text():
    doc = fitz.open(); page = doc.new_page()
    for i in range(10):
        page.insert_text((50, 80 + 20 * i), f"Sor {i}: ellenőrizni kell a nyomást", fontsize=10)
    page.insert_image(page.rect, stream=png(60, 80, 250))          # во весь лист, поверх всего текста
    assert len(L.frags_from_text(page)) == 10


def test_stencil_image_is_not_a_screenshot():
    doc = fitz.open(); page = doc.new_page()
    a = (np.random.rand(200, 400) > 0.5).astype(np.uint8) * 255
    x = page.insert_image(fitz.Rect(100, 450, 400, 600), pixmap=fitz.Pixmap(fitz.csGRAY, 400, 200, a.tobytes(), False))
    doc.xref_set_key(x, "ImageMask", "true")
    assert x not in SH.find_shots(doc)


def test_label_that_does_not_fit_keeps_original():
    img = np.full((30, 60, 3), 230, np.uint8)
    img[10:20, 5:25] = 20                                             # короткая надпись «Ki»
    img[:, 28:30] = 0                                                 # граница кнопки сразу справа
    doc = fitz.open(); page = doc.new_page()
    xref = page.insert_image(fitz.Rect(50, 50, 110, 80), pixmap=fitz.Pixmap(fitz.csRGB, 60, 30, img.tobytes(), False))
    out = SH.render(doc, xref, [({"box": [5, 10, 24, 19], "h": 9, "t": "Ki"}, "Выключено полностью и навсегда")])
    a = np.frombuffer(out.samples, np.uint8).reshape(out.height, out.width, 3)
    assert (a[10:20, 5:25] == 20).all()                               # оригинал не стёрт


def test_docx_toc_block_is_split_and_gets_dot_leader_tabs(tmp_path):
    br = '<w:r><w:br/></w:r>'
    lines = ["Tartalom", "Dokumentumhoz kapcsolódó szerepek ........................ 2",
             "Előzmények ........................................ 13", "A dokumentum célja ............................ 14"]
    para = '<w:p><w:pPr><w:jc w:val="center"/></w:pPr>' + br.join(f'<w:r><w:t xml:space="preserve">{t}</w:t></w:r>' for t in lines) + '</w:p>'
    f = tmp_path / "toc.docx"
    f.write_bytes(_docx(para))
    D.normalize(f)
    assert [d["src"].strip() for d in S.docx_segments(f)] == lines
    assert D._toc_tabs(f) == 3
    xml = zipfile.ZipFile(f).read("word/document.xml").decode("utf-8")
    assert xml.count('w:leader="dot"') == 3 and "....." not in xml

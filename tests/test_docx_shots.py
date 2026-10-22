# -*- coding: utf-8 -*-
"""Надписи на картинках DOCX (22.10): надпись картинки — сегмент после абзаца с картинкой, перевод вписан в ту же
запись word/media (формат и размер в пикселях те же, разметка документа не тронута); язык источника — параметр."""
import io
import json
import re
import zipfile

import fitz
import numpy as np
from PIL import Image

import assemble as A
import llm
import segments as S
import shots as SH


def _picture(lines, size=16, w=420, h=None, under=False):
    """PNG: белый лист, надписи строками сверху вниз; under — подчёркивание последней строки, уходящее вправо и
    касающееся выносных букв (выноска схемы, как в договоре)."""
    h = h or 40 + len(lines) * int(size * 1.3) + 20
    doc = fitz.open(); pg = doc.new_page(width=w, height=h)
    pg.draw_rect(pg.rect, color=None, fill=(1, 1, 1))
    pg.draw_rect(fitz.Rect(w - 60, 10, w - 10, h - 10), color=(0, 0, 0), width=2)    # деталь чертежа
    y = 20
    for t in lines:
        pg.insert_htmlbox(fitz.Rect(20, y, w - 80, y + size * 1.4), t,
                          css=f"* {{font-family: sans-serif; font-size: {size}px; line-height: 1.15; color: black}}")
        y += int(size * 1.3)
    if under:
        pg.draw_line(fitz.Point(20, y - 3), fitz.Point(w - 70, y - 3), color=(0, 0, 0), width=1)
    return pg.get_pixmap(dpi=72, alpha=False).tobytes("png")


def _docx(tmp_path, png, name="src.docx"):
    from docx import Document
    from docx.shared import Cm
    d = Document()
    logo = tmp_path / "logo.png"
    logo.write_bytes(_picture(["LOGO"], w=120, h=60))
    d.sections[0].header.paragraphs[0].add_run().add_picture(str(logo), width=Cm(2))     # колонтитул — не смотрится
    d.add_paragraph("Előtte")
    pic = tmp_path / "pic.png"
    pic.write_bytes(png)
    d.add_picture(str(pic), width=Cm(12))
    d.add_paragraph("Utána")
    out = tmp_path / name
    d.save(out)
    return out


def _echo(lang_of):
    """Портал-заглушка: возвращает надписи как прочитаны, язык — lang_of(текст)."""
    def chat(messages, images=None, max_tokens=0, tag=""):
        labels = re.findall(r"^(\d+): (.*)$", messages[0]["content"], re.M)
        return {"text": json.dumps([{"n": int(n), "t": t, "lang": lang_of(t)} for n, t in labels], ensure_ascii=False),
                "finish": "stop"}
    return chat


def _media(path, name):
    with zipfile.ZipFile(path) as z:
        return z.read(name)


def test_label_translated_in_place_same_size_same_place(tmp_path, monkeypatch):
    src = _docx(tmp_path, _picture(["Szivattyú indítása"], size=18))
    monkeypatch.setattr(llm, "chat", _echo(lambda t: "hu"))
    segs, info = SH.add_docx_shots(S.docx_segments(src), src, tmp_path / "w", "hu")
    shot = [s for s in segs if "shot" in s["loc"]]
    assert [s["text"] for s in shot] == ["Szivattyú indítása"] and info == {"shot_labels": 1}
    # сегмент — между абзацем перед картинкой и после неё; номера — по порядку
    i = segs.index(shot[0])
    assert segs[i - 1]["text"] == "Előtte" and segs[i + 1]["text"] == "Utána"
    assert [s["id"] for s in segs] == list(range(len(segs)))
    name = shot[0]["loc"]["shot"]["media"]

    res = {str(shot[0]["id"]): {"ru": "Запуск насоса", "issues": []}}
    out = tmp_path / "out.docx"
    assert A.build_docx(src, out, segs, res)["shots"] == 1
    assert _media(out, "word/document.xml") == _media(src, "word/document.xml")         # место и размер на листе
    a, b = Image.open(io.BytesIO(_media(src, name))), Image.open(io.BytesIO(_media(out, name)))
    assert (b.format, b.size, b.mode) == (a.format, a.size, a.mode)                   # та же запись, тот же размер
    A_, B_ = np.asarray(a.convert("RGB")).astype(int), np.asarray(b.convert("RGB")).astype(int)
    changed = np.abs(A_ - B_).sum(axis=2) > 30
    assert changed[15:45, 15:340].any() and not changed[:, 345:].any()               # деталь чертежа не тронута
    got = " ".join(lb["t"] for lb in SH.pic_labels(B_.astype(np.uint8), tmp_path / "t.png", "ru"))
    assert "Запуск" in got and "Szivatty" not in got


def test_source_language_is_a_parameter(tmp_path, monkeypatch):
    # русский источник (ru-en): словарь rus+eng, в перевод — только русские надписи, английская остаётся
    img = SH._open_raster(_picture(["Съемный кожух", "Power supply"], size=18))[0]
    labels = SH.pic_labels(img, tmp_path / "t.png", "ru")
    assert any("Съемный кожух" in lb["t"] for lb in labels)
    monkeypatch.setattr(llm, "chat", _echo(lambda t: "ru" if re.search("[а-я]", t) else "en"))
    good = SH.check_labels(b"png", labels, "t", "ru", SH.PICTURE, merge=True)
    assert [lb["t"] for lb in good] == ["Съемный кожух"]


def test_hungarian_prompt_unchanged():
    # ключ кэша портала: уже прочитанные скриншоты PDF не перечитываются
    old = (
        "Это скриншот интерфейса программы. Ниже — надписи с него, прочитанные OCR с венгерским словарём, по одной на "
        "строке в виде «номер: текст».\n{labels}\n\n"
        "Для каждой надписи верни объект JSON {{\"n\": номер, \"t\": \"текст надписи точно как на картинке\", "
        "\"lang\": \"hu\"|\"en\"|\"de\"|\"-\"}}.\n"
        "- t: исправь ошибки OCR по картинке (венгерские буквы á é í ó ö ő ú ü ű только там, где они есть на картинке); "
        "если надписи на картинке нет или это мусор OCR — t пустая строка.\n"
        "- lang: язык надписи на картинке: hu — венгерский (и вместе с кодом или числом: «UTS üzenet», «1 perc»), "
        "en — английский, de — немецкий, \"-\" — без языка: коды и идентификаторы (X0YZ11K500_ZD11, SKRVR_X_TEST), имена "
        "файлов и путей, числа, даты, время, имена людей, сокращения из заглавных латинских букв. OCR читал венгерским "
        "словарём и мог исказить английские слова — решай по картинке.\n"
        "Не добавляй надписей, которых нет в списке. Выведи только массив JSON.")
    assert SH.check_prompt("hu") == old
    assert "русским словарём" in SH.check_prompt("ru", SH.PICTURE) and '"ru"|"en"|"de"|"-"' in SH.check_prompt("ru")


def test_caption_lines_join_menu_items_do_not():
    lab = lambda t, x, y: {"t": t, "box": [x, y, x + 8 * len(t), y + 12], "h": 12, "c": 95}
    labels = [lab("Ось 32\" монитора,", 40, 100), lab("возможно смещение", 40, 114), lab("по направляющей", 40, 128),
              lab("Mentés", 300, 100), lab("Törlés", 300, 114), lab("Kilépés", 300, 128)]
    out = SH._blocks(labels)
    assert [b["t"] for b in out] == ["Ось 32\" монитора, возможно смещение по направляющей", "Mentés", "Törlés", "Kilépés"]
    assert len(out[0]["lines"]) == 3 and out[0]["ts"][1] == "возможно смещение"


def test_caption_rewrapped_underline_kept(tmp_path):
    # подпись в две строки над подчёркиванием-выноской: перевод переносится по ширине подписи, линия цела
    img = SH._open_raster(_picture(["Съемный кожух", "для монтажа монитора."], size=16, under=True))[0]
    labels = SH._blocks(SH.pic_labels(img, tmp_path / "t.png", "ru"))
    block = [lb for lb in labels if lb.get("lines")]
    assert len(block) == 1 and block[0]["t"] == "Съемный кожух для монтажа монитора."
    line_row = int(np.argmax((img[:, 20:300].astype(int).sum(axis=2) < 600).sum(axis=1)))     # строка подчёркивания
    sh = {k: block[0][k] for k in ("box", "h", "t", "lines", "ts")}
    new = SH.render_img(img.copy(), [(sh, "Removable cover for monitor mounting.")], pic=True)
    assert (np.abs(new[line_row, 20:300].astype(int) - img[line_row, 20:300].astype(int)) <= 10).all()
    b = [int(v) for v in block[0]["box"]]
    assert (np.abs(new[b[1]:b[3], b[0]:b[2]].astype(int) - img[b[1]:b[3], b[0]:b[2]].astype(int)).sum(axis=2) > 30).any()
    assert "Removable" in " ".join(lb["t"] for lb in SH.pic_labels(new, tmp_path / "t2.png", "ru"))


def test_vector_pictures_and_header_logo_skipped(tmp_path):
    src = _docx(tmp_path, _picture(["Szivattyú"]))
    pics, skipped = SH.docx_pictures(src)
    assert len(pics) == 1 and not skipped                       # логотип колонтитула не смотрится
    name = pics[0][0]
    emf = tmp_path / "emf.docx"                                 # та же картинка, но записью EMF
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(emf, "w") as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == "word/_rels/document.xml.rels":
                data = data.replace(name.split("/", 1)[1].encode(), name.split("/", 1)[1].replace(".png", ".emf").encode())
            zout.writestr(it.filename.replace(name, name.replace(".png", ".emf")), data)
    pics, skipped = SH.docx_pictures(emf)
    assert not pics and skipped == {name.replace(".png", ".emf"): "векторная (EMF)"}


def test_docx_without_pictures_unchanged(tmp_path):
    from docx import Document
    d = Document(); d.add_paragraph("Szöveg"); d.save(tmp_path / "t.docx")
    segs = S.docx_segments(tmp_path / "t.docx")
    assert SH.add_docx_shots(segs, tmp_path / "t.docx", tmp_path / "w") == (segs, {})
    assert not (tmp_path / "w").exists()


def test_portal_answer_in_its_own_order_realigned():
    # со 2-й надписи ответ сдвинут на одну (договор, 22.10): «2. В конструкцию…» получал «2400-2500» и выпадал
    labels = [{"t": t} for t in ["Вид сбоку", "Пределы", "для монтажа 55\" монитора.", "2. В конструкцию заложено"]]
    got = {1: {"n": 1, "t": "Вид сбоку", "lang": "ru"}, 2: {"n": 2, "t": "для монтажа 55\" монитора.", "lang": "ru"},
           3: {"n": 3, "t": "Пределы", "lang": "ru"}, 4: {"n": 4, "t": "2400-2500", "lang": "-"},
           5: {"n": 5, "t": "2. В конструкцию заложено", "lang": "ru"}}
    out = SH._realign(labels, got)
    assert [out[n]["t"] for n in (1, 2, 3, 4)] == [lb["t"] for lb in labels]


def test_confident_reading_merged_word_by_word():
    assert SH._merge("по горизонту +5мм.", "по горизонту ±5мм.") == "по горизонту ±5мм."      # знак
    assert SH._merge("по направяющей уеза", "по направляющей vesa") == "по направяющей vesa"   # алфавит; опечатка рисунка
    assert SH._merge("Moy", "МФУ") == "МФУ"
    assert SH._merge("ao", "проводов.") == "ao"                                                # чужое слово — нет
    assert SH._merge("FELCSATLAKOZÁS", "LECSATLAKOZÁS") == "FELCSATLAKOZÁS"                    # UTS стр. 26


def test_transparency_and_jpeg_quality_kept():
    # прозрачный PNG: фон остаётся прозрачным, перерисованное — непрозрачно
    rgba = np.zeros((40, 60, 4), np.uint8); rgba[10:20, 10:30] = (0, 0, 0, 255)
    buf = io.BytesIO(); Image.fromarray(rgba, "RGBA").save(buf, "PNG")
    img, im = SH._open_raster(buf.getvalue())
    assert (img[0, 0] == 255).all()                              # прозрачное — на белом листе
    new = img.copy(); new[25:30, 40:50] = 0
    out = Image.open(io.BytesIO(SH._save_raster(new, im, "PNG", (new != img).any(axis=2))))
    al = np.asarray(out)[:, :, 3]
    assert out.mode == "RGBA" and al[0, 0] == 0 and al[27, 45] == 255 and al[15, 15] == 255
    # JPEG: те же таблицы квантования и поворот EXIF
    exif = Image.Exif(); exif[274] = 6
    buf = io.BytesIO(); Image.fromarray(np.full((40, 60, 3), 200, np.uint8)).save(buf, "JPEG", quality=60, exif=exif)
    img, im = SH._open_raster(buf.getvalue())
    out = Image.open(io.BytesIO(SH._save_raster(img, im, "JPEG", np.zeros(img.shape[:2], bool))))
    assert out.quantization == im.quantization and out.getexif().get(274) == 6 and out.size == im.size

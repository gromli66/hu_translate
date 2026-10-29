# Аудит 29.10, причина №7: значения без слов берут перевод из res, идентификаторы кода не переводятся,
# единые термины документа (детектор и подстановка под флагом).
import docx_build as B
import segments as S


def test_value_without_words_takes_res():
    """«1,7 kW» не требует перевода, но в res лежит «1,7 кВт» (fix_units): память отдаёт res, а не исходный текст."""
    assert not S.needs_translation("1,7 kW")
    m = B.Memory([("1,7 kW", "1,7 кВт"), ("6 t/h", "6 т/ч"), ("30÷42oC", "30÷42 °C")])
    assert m.get("1,7 kW") == "1,7 кВт"
    assert m.get("6 t/h") == "6 т/ч"
    assert m.get("30÷42oC") == "30÷42 °C"


def test_value_without_words_not_in_memory_stays():
    """Значения нет в памяти — как раньше: исходный текст, а не None (в missing не попадает)."""
    m = B.Memory([])
    assert m.get("02TD90F601") == "02TD90F601"
    assert m("02TD90F601") == "02TD90F601" and m.missing == []


def test_copied_segment_gets_units():
    """Копируемый сегмент в конвейере: единица по глоссарию при переводе на русский."""
    from translate import fix_units
    assert fix_units("1,7 kW") == "1,7 кВт" and fix_units("36 m3/h") == "36 м³/ч" and fix_units("40÷42 t/h") == "40÷42 т/ч"
    assert fix_units("30TK62P101") == "30TK62P101"


def test_caps_word_from_code_names_is_identifier():
    """RUN рядом с RUN_XSTEP, MODE рядом с get_mode: слово, из которого документ строит имена, — не текст (IS_reko_09)."""
    texts = ["RUN_XSTEP", "REPLAY_ON", "Char400", "RUN", "MODE", "get_mode_val", "END", "KIESETT", "Állapot", "BYTE"]
    cw = S.code_words(texts)
    assert S.is_code_word("RUN", cw) and S.is_code_word("MODE", cw)
    assert not S.is_code_word("END", cw) and not S.is_code_word("KIESETT", cw)    # END нет среди имён — слово как слово
    assert not S.is_code_word("Run", cw)                                          # только прописными
    assert not S.needs_translation("BYTE") and S.needs_translation("Byte")        # тип данных кода


def test_kks_signal_suffix_is_not_a_code_name():
    """TH21S801_NYIT — сигнал с кодом KKS; NYIT в таблице сигналов остаётся венгерским словом (30RVR..VT)."""
    cw = S.code_words(["TH21S801_NYIT", "TJ23S803_NYIT", "NYIT"])
    assert not S.is_code_word("NYIT", cw)
    assert S.code_words(["Zárva", "NYITVA", "A szivattyú"]) == frozenset()


# ------------------------------------------------------------------ единые термины внутри документа (termvar)
import termvar as TV


def _doc(pairs):
    pairs = list(pairs)
    segs = [{"id": i, "text": hu, "loc": {}} for i, (hu, ru) in enumerate(pairs)]
    return {"segs": segs, "res": {str(i): {"ru": ru} for i, (hu, ru) in enumerate(pairs)}}


FILL = [(f"Más mondat {w}", f"Другое предложение {w}") for w in
        "белка волна гроза дерево ежевика желудь заря игла камень лодка мыло нора океан перо река свеча тропа утро фонарь хлеб цапля чайка шарф щека".split()]
MIST = [(f"A vízköddel oltó rendszer {n}. fokozata", f"Система пожаротушения водяным туманом, ступень {n}.") for n in range(1, 7)] + [
    ("A vízköddel oltó rendszer üzembe helyezése", "Ввод системы пожаротушения водяным паром."),
    ("Vízköddel oltó rendszer ellenőrzése", "Проверка системы пожаротушения водяным облаком.")]


def test_term_variants_detected_in_same_slot():
    d = _doc(MIST + FILL)
    g = [g for g in TV.detect(TV.items_of(d["segs"], d["res"])) if g["kind"] == "slot"]
    assert g and any({v["forms"][0] for v in x["variants"]} >= {"туманом", "паром", "облаком"} for x in g)


def test_unify_replaces_minority_with_chosen_form():
    d = _doc(MIST + FILL)
    ed, sk = TV.unify_doc(d)
    assert {e["id"] for e in ed} == {6, 7}
    assert d["res"]["6"]["ru"] == "Ввод системы пожаротушения водяным туманом."
    assert d["res"]["7"]["unified"] == "Проверка системы пожаротушения водяным облаком."


def test_unify_off_by_default_in_pipeline_and_sheet_is_report_only():
    """Лист «Термины» только описывает; res не меняется без вызова unify_doc."""
    from openpyxl import Workbook
    d = _doc(MIST + FILL)
    before = {k: dict(v) for k, v in d["res"].items()}
    wb = Workbook()
    TV.sheet(wb, d)
    assert d["res"] == before and "Термины" in wb.sheetnames
    rows = list(wb["Термины"].iter_rows(values_only=True))
    assert any(r[1] and "паром" in r[1] for r in rows)


def test_same_text_different_translations():
    d = _doc([("Név", "Ф. И. О.")] * 5 + [("Név", "Имя"), ("Név:", "Имя:")])
    g = TV.same_text(TV.items_of(d["segs"], d["res"]))
    assert g and g[0]["variants"][0]["ru"] == "Ф. И. О."
    d2 = _doc([("Név", "Ф. И. О.")] * 5 + [("Név", "Имя")])
    ed, _ = TV.unify(TV.items_of(d2["segs"], d2["res"]), TV.same_text(TV.items_of(d2["segs"], d2["res"])), {})
    assert [e["now"] for e in ed] == ["Ф. И. О."]


def test_unify_keeps_numbers_of_clauses():
    """Сегменты с разными номерами разделов не склеиваются: цифры входят в ключ."""
    d = _doc([("1. Állapot", "1. Состояние"), ("2. Állapot", "2. Статус")])
    assert TV.same_text(TV.items_of(d["segs"], d["res"])) == []

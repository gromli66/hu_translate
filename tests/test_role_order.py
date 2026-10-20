# Договор правил вёрстки, шаги 2–3 (docs/rules_contract/STEP2_3.md): роль строки — одним решателем по старшинству
# источников (факты листа → метка портала → правила вида строки); отступы пунктов — одним источником (место на листе,
# профиль — запас), одно значение в нумерации и в абзаце; проверки Ф7 и М2 — по этому значению.
import fitz
import pytest

import docx_layout as DL

from test_struct_docx import ocr_doc


@pytest.mark.parametrize("facts, role", [
    # факты листа старше метки портала
    ({"table_head": True, "repeat": True}, "body"),                       # шапка таблицы на каждом листе (N2)
    ({"repeat": True, "label": "heading"}, "running"),                    # повтор у края листа
    ({"ocr": True, "label": "header", "stable": True}, "running"),        # скан: «колонтитул» на всех листах
    ({"ocr": True, "label": "header", "text": "1. A turbina üzemzavarai"}, "heading"),   # скан: название раздела в рамке
    ({"title": True, "label": "list", "labeled": True, "text": "1.8. Turbina csapágyak rezgése"}, "heading"),  # Р5
    ({"toc": True, "label": "list", "labeled": True}, "toc"),            # Р1: строка оглавления с меткой «пункт»
    ({"sign": True, "label": "heading", "labeled": True, "text": "Purpose of the work"}, "item"),  # Р6
    ({"toc": True, "title": True, "text": "1.1. ÜV-4 reaktorvédelmi működés 5"}, "toc"),   # оглавление — не название
    ({"pieces": 2, "label": "heading", "labeled": True, "text": "1.1. ÜV-4 reaktorvédelmi működés"}, "body"),
    # метка портала старше правил вида строки
    ({"label": "header", "labeled": True, "text": "File name: 331-165 REV 2.docx"}, "body"),       # Р4: текстовый слой
    ({"label": "caption", "labeled": True, "typo": True}, "body"),       # Р8: подпись — абзац, не заголовок
    ({"label": "table", "labeled": True}, "body"),                       # Р8: таблицы — по линиям, не по метке
    ({"labeled": True, "typo": True, "num_label": True}, "body"),         # портал сказал «текст»
    ({"label": "list", "labeled": True}, "item"),
    ({"label": "heading", "labeled": True, "text": "Bevezetés"}, "heading"),
    ({"label": "heading", "labeled": True, "text": "12"}, "body"),       # заголовок без слов — не заголовок
    # лист без меток — правила вида строки
    ({"typo": True}, "heading"),
    ({"num_label": True}, "item"),                                        # Р9: «a)», «1)» в начале
    ({}, "body"),
])
def test_role_order_table(facts, role):
    assert DL._role(facts) == role


def _text_page(path, lines):
    """Листы текстовым слоем: [(x, y, текст)] или [(лист, x, y, текст)] → PDF и сегменты по строкам (как layout.group
    для текстового слоя)."""
    lines = [ln if len(ln) == 4 else (0,) + tuple(ln) for ln in lines]
    doc = fitz.open()
    for _ in range(max(ln[0] for ln in lines) + 1):
        doc.new_page(width=595, height=842)
    for pno, x, y, t in lines:
        doc[pno].insert_text((x, y), t, fontsize=10)
    doc.save(str(path))
    segs = []
    for pno, pg in enumerate(fitz.open(str(path))):
        for bl in pg.get_text("dict")["blocks"]:
            for ln in bl.get("lines", []):
                t = "".join(sp["text"] for sp in ln["spans"]).strip()
                if t:
                    bb = [round(v, 1) for v in ln["bbox"]]
                    segs.append({"id": len(segs), "page": pno, "text": t, "src": t,
                                 "loc": {"bbox": bb, "frags": [bb], "size": 10.0, "bold": False, "family": "serif",
                                         "align": "left"}})
    return str(path), segs


def _label(segs, labels):
    """Метки портала листа: все сегменты — с границами портала (roles), роль — по куску текста."""
    for s in segs:
        s["loc"]["roles"] = True
        for head, role in labels.items():
            if head in s["text"]:
                s["loc"]["role"] = role
    return segs


TEXT = lambda b: "".join(r.get("t", "") for r in b.get("runs") or [])
BODY = [(85, 100 + 16 * i, f"Ez a bekezdés a dokumentum szövegének {i}. sora, közönséges szöveg.") for i in range(12)]


def test_text_layer_header_label_is_not_heading(tmp_path):
    """Р4: метка портала «колонтитул» на листе текстового слоя (повтор у края листа на трети листов — П) — не заголовок
    раздела: «File name: 331-165 - REV 2 … .docx» внизу титула Eskom шло заголовком H1."""
    src, segs = _text_page(tmp_path / "t.pdf", BODY + [(85, 800, "File name: 331-165 - REV 2. Specification.docx")])
    blocks, _ = DL.load(src, _label(segs, {"File name": "header"}), text=True)
    assert not any(b["k"] == "h" for b in blocks)
    assert any(b["k"] == "p" and TEXT(b).startswith("File name") for b in blocks)


def test_portal_paragraph_after_dash_item_is_not_its_continuation(tmp_path):
    """Р7: абзац, который портал начал новым (метка «текст»), у левого края текста пункта — не продолжение пункта выше:
    «2.1.4. Для курсовых проектов …» НГТУ уходил внутрь пункта «- две последние цифры года …»."""
    src, segs = _text_page(tmp_path / "t.pdf", BODY[:6] + [
        (99, 210, "- две последние цифры года окончания выполнения работы: 01, 02;"),
        (103, 226, "2.1.4. Для курсовых проектов вместо номера по списку указывается номер группы.")] + BODY[6:])
    blocks, _ = DL.load(src, _label(segs, {"- две": "list"}), text=True)
    lst = next(b for b in blocks if b["k"] == "list")
    assert [len(it) for it in lst["items"]] == [1]
    assert any(b["k"] == "p" and TEXT(b).startswith("2.1.4.") for b in blocks)


def test_bullet_sign_beats_portal_heading_label(tmp_path):
    """Р6: знак маркера в начале строки (факт листа) старше метки портала «заголовок»: полужирный пункт «• Purpose …»
    шёл заголовком раздела."""
    src, segs = _text_page(tmp_path / "t.pdf", BODY[:6] + [(99, 210, "• Purpose of the work")] + BODY[6:])
    blocks, _ = DL.load(src, _label(segs, {"Purpose of": "heading"}), text=True)
    assert not any(b["k"] == "h" for b in blocks)
    assert any(b["k"] == "list" and b["kind"] == "Disc" for b in blocks)


def test_scan_dash_beats_portal_heading_label(tmp_path):
    """Р6 на скане: тире текстового слоя перед строкой — пункт, хоть портал и метит строку заголовком."""
    src, segs = ocr_doc(tmp_path)
    for s in segs:
        s["loc"]["roles"] = True
        if s["text"].startswith("Ellenőrizni"):
            s["loc"]["role"] = "heading"
    blocks, _ = DL.load(src, segs)
    assert not any(b["k"] == "h" and "Ellenőrizni" in TEXT(b) for b in blocks)
    assert sum(b["k"] == "list" for b in blocks) >= 1


def test_text_layer_item_text_place_from_its_second_line(tmp_path):
    """Текстовый слой: пункт «•⇥Normal Environmental Conditions …» в две строки — место текста по второй строке пункта
    (висячий отступ оригинала, 22 pt от маркера), а не по ширине «• » (6 pt, Eskom л. 10); пункт в одну строку — по
    ширине знака, как прежде."""
    src, segs = _text_page(tmp_path / "t.pdf", BODY)
    two = {"id": 100, "page": 0, "text": "• Normal Environmental Conditions shall include all operating conditions",
           "src": "", "loc": {"bbox": [72, 300, 520, 324], "frags": [[72, 300, 520, 311], [94, 313, 400, 324]],
                              "size": 10.0, "bold": False, "family": "serif", "align": "left"}}
    blocks, _ = DL.load(src, segs + [two], text=True)
    assert next(b["geo"] for b in blocks if b["k"] == "list") == {"marker": 72, "text": 94}


def test_marker_segment_beside_multiline_item_is_marker(tmp_path):
    """Маркер отдельным сегментом слева от пункта в три строки (Eskom л. 10) — маркер: строка пункта — первая строка
    его сегмента, а не середина всего сегмента; текст — на своём месте (102 pt), а не по ширине знака."""
    src, segs = _text_page(tmp_path / "t.pdf", BODY)
    loc = lambda bb, fr: {"bbox": bb, "frags": fr, "size": 10.0, "bold": False, "family": "serif", "align": "left"}
    segs += [{"id": 100, "page": 0, "text": "•", "src": "", "loc": loc([76, 534, 82, 547], [[76, 534, 82, 547]])},
             {"id": 101, "page": 0, "text": "Normal Environmental Conditions – This section shall include all", "src": "",
              "loc": loc([102, 535, 542, 572], [[102, 535, 542, 547], [102, 547, 541, 560], [102, 560, 542, 572]])}]
    blocks, _ = DL.load(src, segs, text=True)
    lst_ = next(b for b in blocks if b["k"] == "list")
    assert lst_["geo"] == {"marker": 76, "text": 102}
    assert TEXT(lst_["items"][0][0]).startswith("Normal")


def test_unlabeled_letter_items_are_list(tmp_path):
    """Р9: лист без меток портала (портал не ответил) — «a)», «b)» в начале абзаца — пункты списка, как у PDF с
    разметкой (docx_struct.para_lists), а не абзацы."""
    src, segs = _text_page(tmp_path / "t.pdf", BODY[:6] + [(99, 210, "a) első feltétel teljesül"),
                                                          (99, 226, "b) második feltétel teljesül")] + BODY[6:])
    blocks, _ = DL.load(src, segs, text=True)
    lst = [b for b in blocks if b["k"] == "list"]
    assert lst and lst[0]["kind"] == "Label" and sum(len(b["items"]) for b in lst) == 2


def test_toc_number_with_part_at_margin_is_toc_not_item(tmp_path):
    """Р10: номер листа оглавления «56/1» у правого поля — строка оглавления и при метке портала «пункт» (прежде
    у поля узнавалось только число без «/», а «56/1» — только дальше 60 pt от текста)."""
    src, segs = _text_page(tmp_path / "t.pdf", BODY + [(85, 320, "1.3/1. SZBV kazetta leesése"), (480, 320, "56/1")])
    toc = [s for s in segs if s["text"].startswith("1.3/1.")]
    toc[0]["loc"]["bbox"][2] = toc[0]["loc"]["frags"][0][2] = 440        # кусок строки — до 40 pt от номера
    blocks, _ = DL.load(src, _label(segs, {"1.3/1.": "list", "56/1": "list"}), text=True)
    assert not any(b["k"] == "list" for b in blocks)
    assert any(b["k"] == "p" and TEXT(b).startswith("1.3/1.") and TEXT(b).endswith("56/1") for b in blocks)


def test_symbol_markers_keep_their_kind(tmp_path):
    """Маркер-знак Symbol/Wingdings (частная область: «\uf06f» — кружок, «\uf0a7» — квадрат) отдельным сегментом слева —
    вид списка по знаку, как и прежде (при переносе в решатель роли знаки частной области терялись — защитный тест)."""
    src, segs = _text_page(tmp_path / "t.pdf", BODY)
    for k, (ch, y) in enumerate((("\uf06f", 300), ("\uf0a7", 330))):
        segs += [{"id": 100 + 2 * k, "page": 0, "text": ch, "src": ch,
                  "loc": {"bbox": [99, y, 104, y + 10], "frags": [[99, y, 104, y + 10]], "size": 10.0, "bold": False,
                          "family": "serif", "align": "left"}},
                 {"id": 101 + 2 * k, "page": 0, "text": f"pont {k}", "src": f"pont {k}",
                  "loc": {"bbox": [117, y, 200, y + 10], "frags": [[117, y, 200, y + 10]], "size": 10.0, "bold": False,
                          "family": "serif", "align": "left"}}]
    blocks, _ = DL.load(src, segs, text=True)
    assert [b["kind"] for b in blocks if b["k"] == "list"] == ["Circle", "Square"]


def test_bulleted_rows_are_items_not_lineless_grid(tmp_path):
    """Знак маркера — один факт и для роли строки, и для сетки без линеек: строки «• код ⇥ имя ⇥ значение» текстового
    слоя — пункты, а не таблица без рамок (на скане так было по тире слоя, в текстовом слое — нет)."""
    rows = []
    for i in range(3):
        y = 210 + 16 * i
        rows += [(99, y, f"• 30RA0{i}S201"), (220, y, f"NYIT{i}"), (330, y, f"ZÁR{i}")]
    src, segs = _text_page(tmp_path / "t.pdf", BODY[:6] + rows + BODY[6:])
    blocks, _ = DL.load(src, segs, text=True)
    assert not any(b["k"] == "table" for b in blocks)
    assert sum(len(b["items"]) for b in blocks if b["k"] == "list") == 3


def test_same_subheading_in_next_section_is_kept(tmp_path):
    """Подзаголовок с тем же текстом в следующем разделе («1. Ismérvek:» под «1.8. …» и под «1.9. …», 3SZ19) — не
    повтор бегущего названия: заплатка «тот же текст уровня — прочь» выбрасывала его (7 раз), бегущие названия снимают
    правила повтора у края листа (на скане — то же: 3SZ19 разделы с нового листа)."""
    body = lambda pno, y0, w: [(pno, 85, y0 + 16 * i, f"A {w} rezgése megnő, a {i}. jel megjelenik a vezénylőben.")
                               for i in range(4)]
    lines = body(0, 100, "turbina") + [(0, 200, "1.8. Turbina csapágyak rezgése"), (0, 220, "1. Ismérvek:")] +         body(0, 240, "csapágy") + body(1, 100, "gép") + [(1, 300, "1.9. Generátor teherledobás"),
                                                         (1, 320, "1. Ismérvek:")] + body(1, 340, "generátor")
    src, segs = _text_page(tmp_path / "t.pdf", lines)
    blocks, _ = DL.load(src, _label(segs, {"1.8.": "heading", "1.9.": "heading", "Ismérvek": "heading"}), text=True)
    assert [TEXT(b) for b in blocks if b["k"] == "h"].count("1. Ismérvek:") == 2


# ---- шаг 3: один источник отступов пунктов
from docx import Document                                       # noqa: E402
from docx.shared import Pt                                      # noqa: E402

import docx_build as B                                          # noqa: E402
import docx_profile as DP                                       # noqa: E402
import docx_struct as D                                         # noqa: E402
from test_struct_docx import LAYOUT, PROFILE, W, lst, p         # noqa: E402


def _num_ind(doc, par):
    """(маркер, текст) уровня нумерации абзаца-пункта, pt."""
    q = lambda e, a: e.get(W + a)
    numpr = par._p.find(f"{W}pPr/{W}numPr")
    nid, ilvl = q(numpr.find(W + "numId"), "val"), q(numpr.find(W + "ilvl"), "val")
    num = doc.part.numbering_part.element
    aid = next(q(n.find(W + "abstractNumId"), "val") for n in num.findall(W + "num") if q(n, "numId") == nid)
    a = next(a for a in num.findall(W + "abstractNum") if q(a, "abstractNumId") == aid)
    ind = next(lv for lv in a.findall(W + "lvl") if q(lv, "ilvl") == ilvl).find(f"{W}pPr/{W}ind")
    t = int(q(ind, "left")) / 20
    return round(t - int(q(ind, "hanging")) / 20, 1), round(t, 1)


def _par_ind(par):
    pf = par.paragraph_format
    return round(pf.left_indent.pt + (pf.first_line_indent.pt if pf.first_line_indent else 0), 1), \
        round(pf.left_indent.pt, 1)


def test_list_place_on_sheet_always_and_same_in_numbering(tmp_path):
    """Место списка на листе (geo) — всегда, и то же значение в нумерации и в абзаце: прежде место бралось, только если
    расходилось с профилем больше чем на 3 pt, и только в абзац — в нумерации оставался профиль."""
    left = LAYOUT["left"]
    near = dict(lst("Dash", [[p("első", 2, 100, left + 38)]]), geo={"marker": left + 20.0, "text": left + 38.0})
    far = dict(lst("Dash", [[p("második", 2, 200, left + 70)]]), geo={"marker": left + 50.0, "text": left + 70.0})
    out = tmp_path / "l.docx"
    B.Builder(B.Memory([("első", "первый"), ("második", "второй")]), lambda img: None, PROFILE).build(
        [near, far], str(out), LAYOUT)
    doc = Document(str(out))
    for text, want in (("первый", (20.0, 38.0)), ("второй", (50.0, 70.0))):
        par = next(x for x in doc.paragraphs if x.text == text)
        assert _par_ind(par) == want and _num_ind(doc, par) == want


def test_list_text_just_left_of_margin_stays_at_margin(tmp_path):
    """Скан 3PR42: тире на 9,6 pt, текст на 0,8 pt левее поля набора — пункт у поля (место под знак), а не по профилю
    документа (у 3PR42 профиль ставил текст на 155 pt от поля узким столбцом)."""
    left = LAYOUT["left"]
    b = dict(lst("Dash", [[p("A blokk aktuális állapotától", 2, 100, left - 0.8)]]),
             geo={"marker": left - 9.6, "text": left - 0.8})
    P = dict(PROFILE, lists_kind={"Dash": {"0": {"marker": 84.6, "text": 155.4}}})
    out = tmp_path / "m.docx"
    B.Builder(B.Memory([("A blokk aktuális állapotától", "Независимо от состояния")]), lambda img: None, P).build(
        [b], str(out), LAYOUT)
    doc = Document(str(out))
    par = next(x for x in doc.paragraphs if x.text.startswith("Независимо"))
    assert _par_ind(par) == (0.0, 9.0) and _num_ind(doc, par) == (0.0, 9.0)


def test_cell_list_place_from_column_edge(tmp_path):
    """Список в ячейке — по своему месту от края текста столбца на листе (линия столбца и поле ячейки оригинала; без
    линий — левый край текста столбца), прежде — профиль, замеренный от поля листа (ТЭ357, «•» почти все в таблицах).
    Ячейка из одного вложенного списка — не от начала своего текста: тогда маркер всегда вставал к краю (JPDS)."""
    g = lambda x0, x1, top: {"page": 2, "top": top, "bottom": top + 10, "x0": x0, "x1": x1, "first_x0": x0, "size": 10}
    items = dict(lst("Disc", [[p("egy", 2, 120, 315)], [p("kettő", 2, 135, 315)]]), geo={"marker": 300.4, "text": 315.4})
    inner = dict(lst("Dash", [[p("három", 2, 160, 340)]]), geo={"marker": 330.4, "text": 340.4})
    rows = [[{"blocks": [p("Jel", 2, 100, 100)], "th": False, "geom": g(100, 130, 100)},
             {"blocks": [p("Leírás", 2, 100, 300.4), items], "th": False, "geom": g(300.4, 500, 100)}],
            [{"blocks": [p("Jel", 2, 160, 100)], "th": False, "geom": g(100, 130, 160)},
             {"blocks": [inner], "th": False, "geom": g(330.4, 500, 160)}]]
    t = {"k": "table", "rows": rows, "frame": {"xs": [95, 295, 520]}, "geom": dict(g(95, 520, 95), bottom=180)}
    out = tmp_path / "c.docx"
    B.Builder(B.Memory([("egy", "один"), ("kettő", "два"), ("három", "три")]), lambda img: None, PROFILE).build(
        [t], str(out), dict(LAYOUT, first_is_body=True))
    doc = Document(str(out))
    par = next(x for x in doc.tables[0].rows[0].cells[1].paragraphs if x.text == "один")
    assert _par_ind(par) == (0.0, 15.0) and _num_ind(doc, par) == (0.0, 15.0)      # линия 295 + поле 5,4
    par = next(x for x in doc.tables[0].rows[1].cells[1].paragraphs if x.text == "три")
    assert _par_ind(par) == (30.0, 40.0) and _num_ind(doc, par) == (30.0, 40.0)


def test_struct_list_place_inside_table_cells():
    """docx_struct.list_geo — и у списков в ячейках таблиц."""
    it = lambda t, top: {"k": "p", "runs": [{"t": t}], "geom": {"page": 3, "top": top, "x0": 300, "first_x0": 300,
                                                                "rest_x0": 312, "x1": 500}}
    inner = lst("Disc", [[it("• első pont, ami két sorba fut", 100)], [it("• második pont, ami szintén", 130)]])
    t = {"k": "table", "rows": [[{"blocks": [inner], "th": False}]]}
    D.list_geo([t])
    assert inner["geo"] == {"marker": 300, "text": 312}


def _check_profile(**kw):
    return dict(PROFILE, page={"left": LAYOUT["left"], "right": 62.3},
                caption={"center": True, "fig_below": True, "tab_above": True}, **kw)


def test_f7_sees_paragraph_indent_and_place_under_marker(tmp_path):
    """Ф7 сверяет отступ абзаца пункта с нумерацией (прежде видел только нумерацию) и не считает нарушением место под
    знак маркера (тире: текст не ближе 9 pt к маркеру — у профиля 5 pt)."""
    P = _check_profile(lists={"0": {"marker": 18.0, "text": 23.0}})
    blocks = [lst("Dash", [[p("első", 2, 100)], [p("második", 2, 115)]])]
    out = tmp_path / "f.docx"
    B.Builder(B.Memory([("első", "первый"), ("második", "второй")]), lambda img: None, P).build(blocks, str(out), LAYOUT)
    assert not [v for v in DP.check(str(out), P, blocks, LAYOUT) if v["rule"] == "Ф7"]
    doc = Document(str(out))
    par = next(x for x in doc.paragraphs if x.text == "второй")
    par.paragraph_format.left_indent = Pt(50)
    doc.save(str(out))
    assert [v["where"] for v in DP.check(str(out), P, blocks, LAYOUT) if v["rule"] == "Ф7"] == ["второй"]


def test_m2_direct_indent_of_cell_items_is_not_violation(tmp_path):
    """М2 (прямое оформление): отступ пункта списка в ячейке (стиль «List Cell UTS») — значение его списка, не
    нарушение; прямой отступ обычного абзаца — нарушение."""
    g = lambda x0, top: {"page": 2, "top": top, "bottom": top + 10, "x0": x0, "x1": 500, "first_x0": x0, "size": 10}
    items = lst("Disc", [[p("egy", 2, 120, 315)], [p("kettő", 2, 135, 315)]])
    rows = [[{"blocks": [p("Jel", 2, 100, 100)], "th": False, "geom": g(100, 100)},
             {"blocks": [items], "th": False, "geom": g(300, 100)}]] * 2
    blocks = [{"k": "table", "rows": rows, "frame": {"xs": [95, 295, 520]}, "geom": dict(g(95, 95), bottom=160)}]
    out = tmp_path / "m.docx"
    P = _check_profile()
    B.Builder(B.Memory([("egy", "один"), ("kettő", "два")]), lambda img: None, P).build(
        blocks, str(out), dict(LAYOUT, first_is_body=True))
    assert not [v for v in DP.check(str(out), P, blocks, LAYOUT) if v["rule"] == "М2" and "отступ" in v["was"]]


def test_struct_list_of_one_line_items_takes_place_from_runs():
    """Список из однострочных пунктов размеченного PDF — место текста по первому прогону правее метки («•» ⇥ «Текст»):
    прежде место бралось только по второй строке пункта, и такой список шёл по профилю (медиана по документу)."""
    def it(top):
        return {"k": "p", "runs": [{"t": "•", "x": 90.0, "top": top}, {"t": "\t", "x": 94.0, "top": top},
                                   {"t": "Hálózati kapcsolat", "x": 108.0, "top": top}],
                "geom": {"page": 3, "top": top, "x0": 90, "first_x0": 90, "x1": 300, "lines": 1, "rest_x0": None}}
    b = lst("Disc", [[it(100)], [it(115)]])
    D.list_geo([b])
    assert b["geo"] == {"marker": 90, "text": 108.0}
    # метка снята из текста (relist), прогон остался на месте метки — места текста нет, а не верхний индекс в конце
    # строки (ÜFK I «a.) … 3,7 × 10⁶ Bq/dm³»: текст уезжал узким столбцом к правому полю)
    lbl = lambda top: {"k": "p", "runs": [{"t": "A primerköri hűtőközeg összes béta-aktivitása 3,7 × 10", "x": 85.0,
                                           "top": top}, {"t": "⁶", "x": 492.5, "top": top - 0.8},
                                          {"t": " Bq/dm", "x": 495.9, "top": top}],
                       "geom": {"page": 3, "top": top, "x0": 85, "first_x0": 85, "x1": 540, "lines": 1, "rest_x0": None}}
    b = lst("LowerAlpha", [[lbl(100)], [lbl(115)]])
    D.list_geo([b])
    assert "geo" not in b


def test_profile_lists_without_cell_items(tmp_path):
    """Профиль отступов пунктов (lists) — без пунктов в ячейках таблиц, как и профиль по виду списка: в ячейке отступ
    от края столбца, а не от поля листа (ТЭ357: «•» почти все в таблицах — уровень 0 профиля уезжал на 150 pt)."""
    src = fitz.open()
    pg = src.new_page(width=595, height=842)
    for i, (x, y) in enumerate([(90, 100), (90, 120), (300, 300), (300, 320), (300, 340)]):
        pg.insert_text((x, y), "-", fontsize=10)
        pg.insert_text((x + 18, y), f"pont {i} szövege", fontsize=10)
    path = str(tmp_path / "l.pdf")
    src.save(path)
    item = lambda x, y: p("- pont szövege", 1, y - 8, x)
    body = lst("Dash", [[item(90, 100)], [item(90, 120)]])
    cell = lst("Dash", [[item(300, 300)], [item(300, 320)], [item(300, 340)]])
    blocks = [body, {"k": "table", "rows": [[{"blocks": [cell], "th": False}]]}]
    P = DP.extract(path, blocks, dict(LAYOUT, left=62.4))
    assert abs(P["lists"]["0"]["marker"] - (90 - 62.4)) < 1

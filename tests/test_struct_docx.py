# -*- coding: utf-8 -*-
"""Структурный DOCX по профилю оригинала и шлифовка по языковым соглашениям (UTS, 15.10, утверждено Максимом):
списки по месту и номерам из текста (номер не дублируется, нумерация продолжается после рисунка), память без
номеров пунктов, стили и нумерация заголовков по профилю, значок в скобках, соглашения комплекта (ссылки,
латиница, коды, предупреждение, термины документа), цикл шлифовки «код → портал → проверка»; в конвейере —
размеченный PDF узнаётся по покрытию разметкой, абзацы разметки без пары переводятся своими сегментами, шлифовка
правит переводы задачи, сборка даёт DOCX по разметке и таблицу вычитки с листами «шлифовка» и «нарушения»."""
import json
import re

import fitz
import pytest
from docx import Document

import conventions as C
import docx_build as B
import docx_struct as D
import polish as PL

W = B.W


def p(text, page=1, top=100.0, x0=80.0, **kw):
    g = {"page": page, "top": top, "bottom": top + 10, "x0": x0, "x1": 500.0, "first_x0": x0, "size": 10, "lines": 1}
    return dict({"k": "p", "runs": [{"t": text, "font": "Tahoma", "size": 10, "b": False, "i": False, "color": "000000"}],
                 "geom": g}, **kw)


def lst(kind, items):
    return {"k": "list", "kind": kind, "items": items}


RULES = {
    "latin": ["UTS", "Ethernet"], "cyrillic_to_latin": {"УТС": "UTS", "Этернет": "Ethernet"},
    "abbr": {"TU": "ИИ"}, "warning": {"hu": "FIGYELEM", "ru": "ВНИМАНИЕ!"},
    "terms": [{"hu": "[Ii]ndít", "ru": "запуск", "swap": "за",
               "bad": "(?<![А-Яа-яЁё])[Пп](?:уск\\w*|уст(?:ить|ит|ил|ила)\\b|ущен\\w*)"},
              {"hu": "[Tt]esztlépés", "ru": "шаг испытания", "bad": "[Тт]естов\\w*\\s+шаг"}],
}


# ---------------------------------------------------------------- чтение
def test_relist_follows_page_order_and_labels():
    """Дерево Word на стыке страниц: подпункты п. 3 лежат в подсписке п. 2, а п. 3–4 — отдельным списком следом."""
    sub2 = lst("LowerAlpha", [[p("a. első", 44, 600, 116)], [p("b. második", 44, 620, 116)],
                              [p("a. harmadik alatt", 45, 180, 116)]])
    l1 = lst("Decimal", [[p("1. egy", 44, 500)], [p("2. kettő", 44, 580), sub2]])
    l2 = lst("Decimal", [[p("3. három", 45, 150)], [p("4. négy", 45, 260)]])
    out = D.relist([l1, l2])
    assert len(out) == 1 and out[0]["start"] == 1
    items = out[0]["items"]
    assert [D.text_of(it[0]["runs"]) for it in items] == ["egy", "kettő", "három", "négy"]   # номер — не текстом
    assert [D.text_of(x[0]["runs"]) for x in items[1][1]["items"]] == ["első", "második"]
    assert [D.text_of(x[0]["runs"]) for x in items[2][1]["items"]] == ["harmadik alatt"]
    assert items[2][1]["start"] == 1


def test_numbering_continues_after_figure():
    fig = {"k": "p", "runs": [{"img": {"top": 0, "bottom": 100}}], "geom": {"page": 9, "top": 300}}
    out = D.relist([lst("Decimal", [[p("6. hat", 9, 100)]]), fig, lst("Decimal", [[p("7. hét", 9, 500)]])])
    assert out[2]["start"] == 7 and D.text_of(out[2]["items"][0][0]["runs"]) == "hét"


def test_memory_finds_item_without_label():
    mem = B.Memory([("a. BER panel állapota", "a. состояние панели BER"),
                    ("<f1>b. az UTS kiadott </f1><f2>ping x_can</f2>", "<f1>b. команда UTS </f1><f2>ping x_can</f2>")])
    assert mem("BER panel állapota") == "состояние панели BER"
    assert mem("<f1>az UTS kiadott </f1><f2>ping x_can</f2>") == "<f1>команда UTS </f1><f2>ping x_can</f2>"


def test_math_letters_become_plain_text():
    runs = D._runs([{"ch": {"text": c, "fontname": "CambriaMath", "size": 10, "top": 0, "x0": i * 5, "x1": i * 5 + 4,
                            "non_stroking_color": [0]}} for i, c in enumerate("𝑏𝑖𝑛á𝑟𝑖𝑠")])
    assert D.text_of(runs) == "bináris"


# ---------------------------------------------------------------- сборка по профилю
PROFILE = {"page": {"right": 62.3}, "roles": {"body": {"font": "Tahoma", "size": 10.0, "b": False},
                                              "H1": {"font": "Tahoma", "size": 14.0, "b": True}},
           "align": {"body": "both", "list": "left"}, "line": 13.9, "space": {"body": 6.1},
           "heads": {"1": {"page_break": True, "tab": 21.6}, "2": {"page_break": False, "tab": 28.8}},
           "lists": {"0": {"marker": 18.0, "text": 36.0}, "1": {"marker": 54.0, "text": 72.0}},
           "figure": {"center": True}, "caption": {"center": True}, "table": {"pad_left": 5.4}}
LAYOUT = {"W": 595.3, "H": 841.9, "left": 62.4, "right": 59.7, "top": 119, "bottom": 123, "header": 36, "footer": 18,
          "title_left": 45, "head_imgs": [], "head_lines": [], "stamp": [], "stamp_widths": []}


def png():
    return fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 4, 4), False).tobytes("png")


def test_builder_follows_profile(tmp_path):
    icon = {"k": "p", "runs": [{"img": {"x0": 300, "x1": 310, "top": 100, "bottom": 110}}],
            "geom": {"page": 2, "top": 100, "x0": 300, "x1": 310, "first_x0": 300}}
    blocks = [{"k": "h", "level": 1, "runs": [{"t": "1 Bevezetés", "font": "Tahoma", "size": 14, "b": True}],
               "geom": {"page": 2, "top": 90}},
              p("Az indító () gombbal.", 2, 100, 62.4), icon,
              dict(lst("Decimal", [[p("hét", 2, 200)]]), start=7)]
    mem = B.Memory([("Bevezetés", "Введение"), ("Az indító () gombbal.", "Кнопкой запуска ( )."), ("hét", "семь")])
    out = tmp_path / "x.docx"
    B.Builder(mem, lambda img: png(), PROFILE).build(blocks, str(out), LAYOUT)
    doc = Document(str(out))
    st = doc.styles
    assert st["Heading 1"].paragraph_format.page_break_before and int(st["Normal"].paragraph_format.alignment) == 3
    assert abs(st["Normal"].paragraph_format.line_spacing.pt - 13.9) < 0.1
    xml = doc.part.numbering_part.element.xml
    assert 'w:pos="432"' in xml and 'w:hanging="432"' in xml                 # 21,6 pt номер → текст у H1
    assert 'w:startOverride w:val="7"' in xml
    par = next(x for x in doc.paragraphs if "запуска" in x.text)
    seq = "".join("\0" if e.tag == W + "drawing" else (e.text or "") for e in par._p.iter(W + "t", W + "drawing"))
    assert "(\0)" in seq                                                     # значок — в скобках, не в конце


def test_monospace_part_command_verbatim_placeholder_in_brackets():
    assert B.code_part("sync", "синхронизировать") == "sync"
    assert B.code_part("gatewayconfig ipaddr ", "настройка IP-адреса шлюза") == "gatewayconfig ipaddr "
    assert B.code_part(" <beállítani kívánt IP cím>", "IP-адрес для установки") == " <IP-адрес для установки>"
    assert B.keep_code("<f1>Adjuk ki: </f1><f2>sync</f2>", "<f1>Введите: </f1><f2>синхронизация</f2>", [2]) \
        == "<f1>Введите: </f1><f2>sync</f2>"


def test_title_bottom_text_goes_to_title_footer(tmp_path):
    """Юридический текст внизу титула — в подвал титула: иначе русский адрес длиннее и текст уезжал на стр. 2."""
    blocks = [p("UTS felhasználói kézikönyv", 1, 300, 150), p("A dokumentum átadója", 1, 780, 20),
              {"k": "h", "level": 1, "runs": [{"t": "1 Bevezetés", "font": "Tahoma", "size": 14, "b": True}],
               "geom": {"page": 2, "top": 90}}]
    mem = B.Memory([("UTS felhasználói kézikönyv", "Руководство пользователя UTS"),
                    ("A dokumentum átadója", "Передающей стороной"), ("Bevezetés", "Введение")])
    out = tmp_path / "t.docx"
    B.Builder(mem, lambda img: None, PROFILE).build(blocks, str(out), LAYOUT)
    doc = Document(str(out))
    assert [x.text for x in doc.sections[0].footer.paragraphs] == ["Передающей стороной"]
    assert "Передающей стороной" not in [x.text for x in doc.paragraphs]


# ---------------------------------------------------------------- соглашения
def test_refs_quotes_names_codes_warning_terms():
    fix = lambda hu, ru: C.fix_code(hu, ru, RULES)[0]
    assert fix("lásd: 5.5", "(см.: 5.5)") == "(см. раздел 5.5)"
    assert fix("lásd: 7.2 és 7.3", "см. 7.2 и 7.3") == "см. разделы 7.2 и 7.3"
    assert fix("(4. fejezet)", "Запуск UTS (глава 4)") == "Запуск UTS (см. раздел 4)"
    assert fix("a 8. fejezetben", "описано в главе 8") == "описано в разделе 8"
    assert fix("az 5. ábrán", "показано на рис. 5") == "показано на рисунке 5"
    assert fix("Ethernet kábel", "Кабель Этернет") == "Кабель Ethernet"
    assert fix("UTS_A.LIS", "UTS_A.СПИС") == "UTS_A.LIS"
    assert fix("FIGYELEM: újra", "ВНИМАНИЕ: при повторе") == "ВНИМАНИЕ! При повторе"
    assert fix("<f1>FIGYELEM!</f1><f2> újra</f2>", "<f1>Внимание!</f1><f2> при повторе</f2>") \
        == "<f1>ВНИМАНИЕ!</f1><f2> При повторе</f2>"                                # без двух пробелов в DOCX
    assert fix("a TU-k listája", "список TU") == "список ИИ"
    assert fix("az UTS indítása", "Пуск UTS") == "Запуск UTS"
    assert fix("indítás után üres", "после пуска поле пустое") == "после запуска поле пустое"   # «пустое» не трогать
    assert fix("„B” szerver", 'сервер "B"') == "сервер «B»"
    assert fix('Filter=„0x0”', 'Filter="0x0"') == 'Filter="0x0"'                  # значение настройки — как есть
    assert fix("<dátum>", "TxsAttachLog_&lt;дата&gt;") == "TxsAttachLog_<дата>"
    assert C.check("5.sz. melléklet", "Приложение 5", RULES) == []                   # «5.sz» — не код
    v = C.check("Az UTS tesztlépés", "тестовый шаг UTS", RULES)
    assert [(r, code) for r, _, code in v] == [("Я10", False)]
    assert C.check("PARAMÉTERNÉV=PARAMÉTERÉRTÉK", "пара PARAMÉTERNÉV=PARAMÉTERÉRTÉK", RULES)[0][0] == "Я11"
    assert C.check("BLÜGY laptop", "ноутбук НСБ (BLÜGY)", RULES) == []           # первое употребление сокращения


def test_polish_cycle_code_then_portal_and_log(tmp_path):
    items = {"a": {"hu": "Az UTS indítása (lásd: 4.2)", "ru": "Пуск УТС (см.: 4.2)", "kind": "text", "page": 3, "id": 1},
             "b": {"hu": "A tesztlépés vége", "ru": "Конец тестового шага", "kind": "text", "page": 4, "id": 2},
             "c": {"hu": "A tesztlépés eleje", "ru": "Начало тестового шага", "kind": "text", "page": 5, "id": 3}}
    calls = []

    def chat(msgs, **kw):
        data = json.loads(msgs[-1]["content"])
        calls.append(len(data))
        # «b» портал исправил, «c» вернул с лишним числом — правка не принимается
        fixed = {"Конец тестового шага": "Конец шага испытания", "Начало тестового шага": "Начало 2 шага испытания"}
        return {"text": json.dumps([{"id": d["id"], "ru": fixed[d["ru"]]} for d in data], ensure_ascii=False)}
    edits, left = PL.run(items, RULES, chat=chat, log=lambda *a: None)
    assert items["a"]["ru"] == "Запуск UTS (см. раздел 4.2)"
    assert items["b"]["ru"] == "Конец шага испытания" and items["c"]["ru"] == "Начало тестового шага"
    assert {(e["key"], e["who"]) for e in edits} == {("a", "код"), ("b", "портал")}
    assert [x["key"] for x in left] == ["c"] and calls[0] == 2
    out = tmp_path / "v.xlsx"
    PL.write_xlsx(str(out), edits, left)
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb["шлифовка"].max_row == 3 and wb["нарушения"]["B2"].value == "Я10"


def test_polish_accepts_collapsing_expanded_name_rejects_cut_answer():
    """Портал сворачивает раскрытое название в «UTS» — перевод короче в разы, это не обрыв (сервис 15.10: все 12 таких
    правок отбрасывала проверка длины против старого перевода); обрыв ответа по-прежнему не принимается."""
    hu, old = "Kapcsolja be az UTS számítógépet!", "Включите ЭВМ системы оперативной информации!"
    assert PL._ok(old, "Включите ЭВМ UTS!", hu, RULES, "text")
    assert PL._ok("3 Система оперативной информации", "3 Система UTS", "3 Az UTS rendszer", RULES, "text")
    assert not PL._ok(old, "UTS", hu, RULES, "text")


def test_vocab_skips_paths_and_abbreviation_written_once_lowercase():
    """«uts» в путях и именах файлов и один раз строчными в тексте при сотнях «UTS» — не венгерское слово (иначе голое
    «UTS» переводилось и не ловилась потерянная латиница, сервис 15.10); слова через косую и венгерское слово
    прописными, которое встречается и строчными, остаются."""
    import pipeline as P

    class T:
        entries = []
    segs = [{"text": t} for t in ["Az UTS alkalmazás indítása"] * 12 + ["MINDEN JEL"] * 20 + [
        "(/etc/uts/shellenv)", "az uts.ini fájl", "uts_<éééé>.tar.gz", "egyébként pedig az uts naplófájl",
        "a fogyasztás kg/óra", "minden", "minden", "minden"]]
    tr = T()
    P.set_vocab(tr, segs)
    assert "uts" not in tr.vocab and "ini" not in tr.vocab
    assert {"kg", "óra", "minden", "alkalmazás", "naplófájl"} <= tr.vocab


def test_item_number_as_text_but_not_initials():
    import docx_profile as DP
    assert DP.ITEM_NUM_TEXT.match("1. Включите ЭВМ") and DP.ITEM_NUM_TEXT.match("a) пункт")
    assert not DP.ITEM_NUM_TEXT.match("Ф. И. О. оператора: имя пользователя UTS")


def run(t, x, top=140.0, page=31, **kw):
    return dict({"t": t, "font": "Tahoma", "size": 10, "b": False, "i": False, "color": "000000",
                 "x": x, "top": top, "page": page}, **kw)


def test_icon_in_line_goes_to_its_place_not_paragraph_end(tmp_path):
    """Word рисует значок строки после текста: «● : не подключено» становилось «: не подключено●» (UTS, 27 пунктов,
    16.10). Значок встаёт перед текстом своей строки по листу; значок без MCID (оформление) возвращается в абзац своей
    строки; в DOCX значок после маркера пункта — в начале пункта."""
    icon = {"img": {"x0": 98, "x1": 106, "top": 141, "bottom": 149, "page_number": 31}}
    runs = [run("•", 90), run("\t", 94), run(": az UTS nincs felcsatlakozva", 108)]
    got = D.insert_imgs(runs, [icon])
    assert ["img" in r for r in got] == [False, False, True, False]
    far = {"img": {"x0": 400, "x1": 408, "top": 300, "bottom": 308, "page_number": 31}}
    assert "img" in D.insert_imgs(runs, [far])[-1]                       # не на строке абзаца — в конец, как было
    para = {"k": "p", "runs": [run(": nem meghatározható", 108, 300)],
            "geom": {"page": 31, "top": 300, "x0": 108, "x1": 200, "size": 10}}
    D.attach_icons([para], [{"x0": 98, "x1": 106, "top": 301, "bottom": 309, "page_number": 31},
                            {"x0": 62, "x1": 120, "top": 40, "bottom": 80, "page_number": 31}], 842)   # логотип шапки — нет
    assert ["img" in r for r in para["runs"]] == [True, False]
    blocks = [lst("Disc", [[{"k": "p", "runs": got, "geom": {"page": 31, "top": 140, "x0": 90, "x1": 300}}]])]
    mem = B.Memory([(": az UTS nincs felcsatlakozva", ": UTS не подключено")])
    out = tmp_path / "i.docx"
    B.Builder(mem, lambda img: png(), PROFILE).build(blocks, str(out), LAYOUT)
    par = next(x for x in Document(str(out)).paragraphs if "подключено" in x.text)
    seq = "".join("\0" if e.tag == W + "drawing" else (e.text or "") for e in par._p.iter(W + "t", W + "drawing"))
    assert seq.lstrip().startswith("\0") and seq.rstrip().endswith("подключено")


def test_table_continued_with_repeated_header_is_joined():
    """Продолжение таблицы на новом листе, начатое повтором шапки («Повторять как заголовок»), — та же таблица, повтор
    прочь; другая таблица с другой шапкой — своя (синтетика, 16.10)."""
    cell = lambda th, t: {"blocks": [p(t)], "th": th}
    head = lambda: [cell(True, "№"), cell(True, "Параметр")]
    t1 = {"k": "table", "rows": [head(), [cell(False, "1"), cell(False, "a")]], "geom": {"page": 4}}
    t2 = {"k": "table", "rows": [head(), [cell(False, "2"), cell(False, "b")]], "geom": {"page": 5}}
    t3 = {"k": "table", "rows": [[cell(True, "Код"), cell(True, "Имя")], [cell(False, "x"), cell(False, "y")]],
          "geom": {"page": 6}}
    got = D.join_tables([t1, t2, t3])
    assert [len(t["rows"]) for t in got] == [3, 2]


def test_table_cut_by_pages_is_joined():
    """Word при экспорте режет таблицу по листам: продолжение — таблица без шапки сразу за ней (UTS стр. 75–78)."""
    cell = lambda th: {"blocks": [], "th": th}
    t1 = {"k": "table", "rows": [[cell(True), cell(True)], [cell(False), cell(False)]], "geom": {"page": 75}}
    t2 = {"k": "table", "rows": [[cell(False), cell(False)]] * 3, "geom": {"page": 76}}
    t3 = {"k": "table", "rows": [[cell(True), cell(True)], [cell(False), cell(False)]], "geom": {"page": 77}}
    got = D.join_tables([t1, t2, t3])
    assert [len(t["rows"]) for t in got] == [5, 2]


def test_label_tab_kept_and_column_alignment_from_original():
    assert B._keep_tab("Műszaki változás száma\t10776", "Номер технического изменения 10776") == \
        "Номер технического изменения\t10776"
    assert B._keep_tab("Tervező\tSinger Gábor", "Проектант\tСинглер Габор") == "Проектант\tСинглер Габор"
    g = lambda x0, x1: {"page": 2, "x0": x0, "x1": x1, "lines": 1}
    rows = [[{"th": True, "geom": g(160, 202)}, {"th": True, "geom": g(410, 450)}],
            [{"th": False, "geom": g(66, 120)}, {"th": False, "geom": g(500, 555)}]]
    assert B._col_align(rows, [62, 300, 560], 2) == {("head", 0): "center", ("head", 1): "center", ("body", 1): "right"}


def test_latin_name_counted_paths_aside_and_verbatim_style_kept():
    """«…UTS… при запуске СОИ другой среды» проходило проверку: UTS хоть раз есть (или «uts_kmod.conf»)."""
    hu = "Blokki UTS esetén az UTS alkalmazás indítása előtt, lásd /etc/modprobe.d/uts_kmod.conf"
    assert any(v[0] == "Я5" for v in C.check(hu, "В случае СОИ блока перед запуском приложения UTS, см. "
                                                 "/etc/modprobe.d/uts_kmod.conf", RULES))
    assert not any(v[0] == "Я5" for v in C.check(hu, "В случае UTS блока перед запуском приложения UTS, см. "
                                                     "/etc/modprobe.d/uts_kmod.conf", RULES))
    assert not any(v[0] == "Я5" for v in C.check("<UTS környezet>/<év>", "<среда UTS>/<год>", RULES))
    plain, ital = (False, False, False, None), (False, True, False, None)
    parts = ["Ezen kívül a kernel modul ", "/etc/modprobe.d/uts_kmod.conf", " fájlban meg kell adni"]
    groups = [[plain, []], [ital, []], [plain, []]]
    out = B._restyle("Кроме того, в файле /etc/modprobe.d/uts_kmod.conf необходимо указать", groups, parts)
    assert (ital, "/etc/modprobe.d/uts_kmod.conf") in out and "".join(t for _, t in out).startswith("Кроме")


def test_slanted_tahoma_is_italic_and_verbatim_parts_need_no_portal():
    """У Tahoma нет курсива — Word наклоняет знаки матрицей (c = 0,33): это курсив (пути в UTS). Если перевод без меток
    есть и курсивный кусок стоит в нём дословно — перевод с метками у портала не просим."""
    ch = {"fontname": "BCDEEE+Tahoma", "size": 10, "non_stroking_color": (0,), "matrix": (1, 0, 0.33, 1, 200, 600)}
    assert D._fmt(ch)["i"] and not D._fmt(dict(ch, matrix=(1, 0, 0, 1, 200, 600)))["i"]
    plain, ital = (False, False, False, None), (False, True, False, None)
    groups = [[plain, []], [ital, []], [plain, []]]
    parts = ["Ezen kívül a kernel modul ", "/etc/modprobe.d/uts_kmod.conf", " fájlban"]
    assert B._verbatim("Кроме того, в файле /etc/modprobe.d/uts_kmod.conf", groups, parts)
    assert not B._verbatim("Отметьте опцию «Шаг испытания успешен»", groups, ["Jelölje be a ", "„A tesztlépés sikeres”", ""])


def test_toc_like_original_levels_bold_title_and_unnumbered_entry(tmp_path):
    """Оглавление — по оригиналу: уровней столько же (UTS: 3, а собиралось 4), 1-й полужирный, отступы и шаг строк,
    заголовок своим кеглем и цветом, пункт без номера раздела (шапка таблицы) — с уровнем структуры."""
    info = {"title": {"t": "Tartalom", "font": "Tahoma", "size": 14, "b": True, "color": "365f91", "center": False},
            "levels": 3, "bold": [True, False, False], "indent": [0, 21, 40], "pitch": 19.6, "figures": False,
            "extra": [("Dokumentumhoz kapcsolódó szerepek", 1)]}
    th = {"blocks": [p("Dokumentumhoz kapcsolódó szerepek", 2, 100)], "th": True}
    blocks = [{"k": "table", "rows": [[th], [{"blocks": [p("Singler Gábor", 2, 120)], "th": False}]],
               "geom": {"page": 2, "top": 100}}, {"k": "toc", "info": info}]
    mem = B.Memory([("Tartalom", "Содержание"), ("Dokumentumhoz kapcsolódó szerepek", "Роли, связанные с документом"),
                    ("Singler Gábor", "Синглер Габор")])
    out = tmp_path / "t.docx"
    B.Builder(mem, lambda img: png(), PROFILE).build(blocks, str(out), LAYOUT)
    doc = Document(str(out))
    body = doc.element.body.xml
    assert 'TOC \\o "1-3"' in body and '<w:outlineLvl w:val="0"/>' in body
    st = {s.style_id: s for s in doc.styles}
    pf2 = st["TOC2"].paragraph_format                         # начало первой строки уровня — как в оригинале (висячий)
    assert st["TOC1"].font.bold and not st["TOC2"].font.bold and pf2.left_indent.pt + pf2.first_line_indent.pt == 21
    assert st["TOC1"].element.get(W + "customStyle") is None      # встроенный «toc 1», иначе Word берёт свой
    assert str(st["TOCTitleUTS"].font.color.rgb) == "365F91"


# ---------------------------------------------------------------- в конвейере: извлечение → структура → шлифовка → сборка
def tagged_pdf(path, pages, tagged=True, producer="Microsoft® Word 2019"):
    """PDF «как из Word»: pages — [[(тег, текст, x, y, кегль, жирный)]], каждая строка — свой элемент дерева разметки
    со своим MCID. Возвращает рамки строк [(стр., текст, Rect, кегль, жирный)] — из них сегменты страниц."""
    doc = fitz.open()
    elems, lines = [], []
    for pno, items in enumerate(pages):
        page = doc.new_page(width=595, height=842)
        for tag, text, x, y, size, bold in items:
            page.insert_text((x, y), text, fontsize=size, fontname="cour" if bold == "code" else "hebo" if bold else "helv")
            lines.append((pno, text, page.search_for(text)[0], size, bold))
        if not tagged:
            continue
        page.clean_contents()
        xref = page.get_contents()[0]
        s = doc.xref_stream(xref).decode("latin-1")
        out, pos = [], 0
        for i, (b, it) in enumerate(zip(re.findall(r"BT.*?ET", s, re.S), items)):
            k = s.index(b, pos)
            out.append(s[pos:k] + f"/{it[0]} <</MCID {i}>> BDC " + b + " EMC")
            pos = k + len(b)
        doc.update_stream(xref, ("".join(out) + s[pos:]).encode("latin-1"))
        elems += [(it[0], page.xref, i) for i, it in enumerate(items)]
    if tagged:
        root, sect = doc.get_new_xref(), doc.get_new_xref()
        kids = []
        for tag, pxref, mcid in elems:
            x = doc.get_new_xref()
            doc.update_object(x, f"<< /Type /StructElem /S /{tag} /P {sect} 0 R /Pg {pxref} 0 R /K {mcid} >>")
            kids.append(x)
        doc.update_object(sect, f"<< /Type /StructElem /S /Sect /P {root} 0 R /K [{' '.join(f'{k} 0 R' for k in kids)}] >>")
        doc.update_object(root, f"<< /Type /StructTreeRoot /K [{sect} 0 R] >>")
        doc.xref_set_key(doc.pdf_catalog(), "StructTreeRoot", f"{root} 0 R")
    doc.set_metadata({"producer": producer, "creator": producer})
    doc.save(path)
    return lines


def page_segs(lines):
    return [{"id": i, "text": t, "src": t, "page": pno,
             "loc": {"bbox": list(r), "avail": list(r), "frags": [list(r)], "align": "left", "cell": [0, 595],
                     "size": size, "bold": bold, "family": "sans-serif", "ocr": False}}
            for i, (pno, t, r, size, bold) in enumerate(lines)]


PAGES = [[("P", "UTS felhasználói kézikönyv", 150, 300, 18, True)],
         [("H1", "1 Bevezetés", 62, 130, 14, True),
          ("P", "Az UTS indítása után a felület üres.", 62, 160, 10, False),
          ("P", "FIGYELEM!", 62, 190, 10, True), ("P", "A teszt leállítása nem lehetséges.", 115, 190, 10, False),
          ("P", "Adjuk ki az alábbi parancsot:", 62, 220, 10, False), ("P", "sync", 210, 220, 10, "code")]]
# перевод «портала»: с нарушениями соглашений (УТС, «пуск», «ВНИМАНИЕ:»), их правит шлифовка кодом; команда
# моноширинным шрифтом переведена — её возвращает сборка
RU = {"UTS felhasználói kézikönyv": "Руководство пользователя UTS", "1 Bevezetés": "1 Введение",
      "Az UTS indítása után a felület üres.": "После пуска УТС интерфейс пуст.", "FIGYELEM!": "ВНИМАНИЕ!",
      "A teszt leállítása nem lehetséges.": "Остановить испытание нельзя.",
      "<f1>FIGYELEM!</f1><f2> A teszt leállítása nem lehetséges.</f2>":
          "<f1>ВНИМАНИЕ:</f1><f2> остановить испытание нельзя.</f2>",
      "Adjuk ki az alábbi parancsot:": "Введите следующую команду:", "sync": "sync",
      "<f1>Adjuk ki az alábbi parancsot:</f1><f2> sync</f2>": "<f1>Введите следующую команду:</f1><f2> синхронизировать</f2>"}


def test_tagged_pdf_is_recognized_by_coverage(tmp_path):
    lines = tagged_pdf(tmp_path / "t.pdf", PAGES)
    assert B.coverage(str(tmp_path / "t.pdf"), page_segs(lines)) == 1.0
    plain = tagged_pdf(tmp_path / "p.pdf", PAGES, tagged=False)
    assert B.coverage(str(tmp_path / "p.pdf"), page_segs(plain)) == 0.0      # без разметки — прежним путём
    # разметка не из Word (Acrobat ставит её сам, дерево иное: 31RDGT развалился) — прежним путём
    acro = tagged_pdf(tmp_path / "a.pdf", PAGES, producer="Adobe PDF Library 20.1.68")
    assert B.coverage(str(tmp_path / "a.pdf"), page_segs(acro)) == 0.0


def test_pipeline_struct_fill_polish_assemble(tmp_path, monkeypatch):
    import pipeline as P
    import translate as T
    from openpyxl import load_workbook
    pdf = tmp_path / "doc.pdf"
    segs = page_segs(tagged_pdf(pdf, PAGES))
    work = tmp_path / "work"
    (work / "json").mkdir(parents=True)
    res = {str(s["id"]): {"ru": RU[s["text"]], "issues": []} for s in segs}
    (work / "json" / "doc.json").write_text(json.dumps({"file": str(pdf), "segs": segs, "res": res, "info": {},
                                                        "docx": None, "struct": {"coverage": 1.0}}), encoding="utf-8")
    g = tmp_path / "g.md"
    g.write_text("## Т\n\n| Венгерский | Русский | Комментарий |\n|---|---|---|\n| szelep | клапан | |\n", encoding="utf-8")
    rules = tmp_path / "conventions.json"
    rules.write_text(json.dumps({k: RULES[k] for k in ("latin", "cyrillic_to_latin", "warning")}), encoding="utf-8")
    (tmp_path / "doc_terms").mkdir()
    (tmp_path / "doc_terms" / "doc_terms.json").write_text(json.dumps({"terms": RULES["terms"][:1]}), encoding="utf-8")
    prj = {"glossary": str(g), "rules": str(rules), "doc_terms": str(tmp_path / "doc_terms")}
    asked = []

    def call(self, units, *a, **k):
        asked.extend(u["text"] for u in units)
        return {i: RU[u["text"]] for i, u in enumerate(units)}, {}
    monkeypatch.setattr(T.Translator, "_call", call)

    # абзац «FIGYELEM!» жирным + текст на той же строке — в разметке два элемента, в DOCX — один абзац с метками:
    # среди сегментов страниц его нет, он переводится своим сегментом (без меток — не спрашивается)
    P.struct_fill(work, prj)
    d = json.loads((work / "json" / "doc.json").read_text(encoding="utf-8"))
    own = [s for s in d["segs"] if s["loc"].get("struct_only")]
    assert asked == ["<f1>FIGYELEM!</f1><f2> A teszt leállítása nem lehetséges.</f2>",
                     "<f1>Adjuk ki az alábbi parancsot:</f1><f2> sync</f2>"]
    assert [s["page"] for s in own] == [1, 1] and d["struct"]["filled"] == 2
    assert d["res"][str(own[1]["id"])]["ru"] == "<f1>Введите следующую команду:</f1><f2> sync</f2>"   # команда — дословно

    P.polish(work, prj, chat=lambda *a, **k: pytest.fail("всё правится кодом — портал не нужен"))
    d = json.loads((work / "json" / "doc.json").read_text(encoding="utf-8"))
    r = d["res"][str(segs[2]["id"])]
    assert r["ru"] == "После запуска UTS интерфейс пуст." and r["polish"]["who"] == "код"
    assert d["res"][str(own[0]["id"])]["ru"] == "<f1>ВНИМАНИЕ!</f1><f2> Остановить испытание нельзя.</f2>"
    assert d["polish"]["remaining"] == [] and {e["rule"] for e in d["polish"]["edits"]} >= {"Я5, Я10", "Я8"}

    P.assemble(work)
    out = work / "out"
    doc = Document(str(out / "doc_RU.docx"))
    pars = [(p.style.name, p.text) for p in doc.paragraphs if p.text.strip()]
    assert ("Heading 1", "Введение") in pars and ("Normal", "После запуска UTS интерфейс пуст.") in pars
    warn = next(p for p in doc.paragraphs if p.text.startswith("ВНИМАНИЕ!"))
    assert warn.text == "ВНИМАНИЕ! Остановить испытание нельзя." and warn.runs[0].bold
    cmd = next(p for p in doc.paragraphs if p.text.startswith("Введите"))
    assert cmd.text == "Введите следующую команду: sync" and cmd.runs[-1].font.name == "Consolas"
    assert (out / "doc_RU.pdf").exists()
    wb = load_workbook(out / "doc_вычитка.xlsx")
    assert wb.sheetnames == ["вычитка", "шлифовка", "нарушения"]
    main = {row[2]: row for row in wb["вычитка"].iter_rows(min_row=2, values_only=True)}
    assert main["Az UTS indítása után a felület üres."][3] == "После запуска UTS интерфейс пуст."
    assert "Шлифовка (Я5, Я10; код), было: После пуска УТС" in main["Az UTS indítása után a felület üres."][7]
    assert "FIGYELEM! A teszt leállítása nem lehetséges." in main                # абзац разметки — строкой вычитки
    assert wb["шлифовка"].max_row == 1 + len(d["polish"]["edits"])


def test_editor_sheet_found_by_name_when_other_sheet_active(tmp_path):
    """Редактор сохранил таблицу, стоя на листе «шлифовка»: правки всё равно читаются с листа «вычитка»."""
    import apply_review as AR
    from assemble import review_xlsx
    from openpyxl import load_workbook
    src = tmp_path / "d.docx"
    dx = Document()
    dx.add_paragraph("A szelep zárva.")
    dx.save(str(src))
    segs = [{"id": 0, "text": "A szelep zárva.", "src": "A szelep zárva.", "page": None,
             "loc": {"part": "word/document.xml", "p": 0, "groups": [[0]]}}]
    res = {"0": {"ru": "Клапан закрыт.", "issues": []}}
    jf = tmp_path / "d.json"
    jf.write_text(json.dumps({"file": str(src), "segs": segs, "res": res}, ensure_ascii=False), encoding="utf-8")
    xf = tmp_path / "d_вычитка.xlsx"
    review_xlsx(xf, segs, res, "d", polish={"edits": [], "remaining": []})
    wb = load_workbook(xf)
    wb["вычитка"]["G2"] = "Клапан закрыт полностью."
    wb.active = wb.sheetnames.index("шлифовка")
    wb.save(xf)
    assert AR.apply(jf, xf, tmp_path / "tm.json", tmp_path / "out") == [("A szelep zárva.", "Клапан закрыт полностью.")]


def test_runner_weights_cover_new_steps():
    from server.runner import W_REVIEW, W_FAST
    for st in ("извлечение", "перевод", "структура", "шлифовка", "сборка", "термины"):
        assert st in W_REVIEW and st in W_FAST
    assert "проверка" in W_REVIEW


# ---------------------------------------------------------------- без зашивок под UTS (ÜFK, 16.10) и распознанный PDF
def test_spans_autonum_text_labels_and_continuation_note():
    """Шапка над «A. | текст» — по столбцам; автонумерация — только если повторяет номера оригинала; номера и тире
    в тексте пунктов — вид списка; пометка о переносе между частями таблицы — прочь."""
    row = [{"geom": {"x0": 110, "x1": 174}}, {"geom": {"x0": 291, "x1": 398}}, {"geom": {"x0": 471, "x1": 531}}]
    assert B._span_starts(row, [58, 80, 231, 267, 464, 537], 5) == [(0, 1), (2, 3), (4, 4)]
    # вторая строка двухъярусной шапки: «Описание» сверху занимает и её первый столбец (ИН-1172, 17.10)
    sub = [{"geom": {"x0": 407 + 41 * i, "x1": 430 + 41 * i}} for i in range(4)]
    assert B._span_starts(sub, [62, 398, 440, 482, 520, 560], 5) == [(1, 1), (2, 2), (3, 3), (4, 4)]
    h = lambda t, lv: {"k": "h", "level": lv, "runs": [{"t": t}]}
    assert B._autonum_ok([h("1 Bevezetés", 1), h("1.1 Cél", 2), h("2 Rendszer", 1)])
    assert not B._autonum_ok([h("1. Bevezetés", 1), h("DF 1.1. Cél", 2), h("DF 1.2. Más", 2)])
    assert not B._autonum_ok([h("Bevezetés", 1)])                      # номеров нет — автонумерация их не добавит
    item = lambda t: [p(t)]
    alpha = {"k": "list", "kind": "Disc", "items": [item("a. első"), item("b. második"), item("c. harmadik")]}
    dash = {"k": "list", "kind": "Disc", "items": [item("– ha az első"), item("– ha a második")]}
    D.text_labels([alpha, dash])
    assert (alpha["kind"], alpha["suffix"], dash["kind"]) == ("LowerAlpha", ".", "Dash")
    cell = lambda th: {"blocks": [], "th": th}
    t1 = {"k": "table", "rows": [[cell(True)] * 2, [cell(False)] * 2], "geom": {"page": 90}}
    note = {"k": "p", "runs": [{"t": "Folytatás a következő oldalon."}], "geom": {"page": 90}}
    t2 = {"k": "table", "rows": [[cell(False)] * 2], "geom": {"page": 91}}
    got = D.join_tables([t1, note, t2])
    assert len(got) == 1 and len(got[0]["rows"]) == 3


def ocr_doc(tmp_path, pages=8):
    """Распознанный PDF: текст — в сегментах OCR, в текстовом слое — только тире пунктов (как «Печать в PDF» 3SZ19)."""
    doc = fitz.open()
    segs = []
    sid = 0

    def seg(pno, text, x0, y0, x1, align="left", bold=False):
        nonlocal sid
        segs.append({"id": sid, "page": pno, "text": text,
                     "loc": {"bbox": [x0, y0, x1, y0 + 9], "frags": [[x0, y0, x1, y0 + 9]], "size": 8.3 + (sid % 3) / 10,
                             "bold": bold, "align": align, "family": "monospace", "ocr": True}})
        sid += 1
    for pno in range(pages):
        page = doc.new_page(width=595, height=842)
        seg(pno, "Azonosító: 3SZ19", 77, 35, 170)
        seg(pno, f"Oldalszám: {pno + 2}", 460, 48, 530)
        title = ["1. A turbina üzemzavarai"] * 3 + ["2. Szekunderköri berendezések"] * 3 + ["3. Egyéb"] * 2
        seg(pno, title[pno], 250, 86, 395, "center", True)     # бегущее название раздела — на листах раздела
        page.insert_text((100, 140 + 7), "-", fontsize=9)
        seg(pno, f"Ellenőrizni a turbinát {pno}.", 112, 140, 400)
        seg(pno, "1.4.", 100, 200, 120)
        seg(pno, "Tengelyelmozdulás megnövekedése", 160, 200, 330)
        seg(pno, "13", 500, 200, 512)
    doc.save(tmp_path / "ocr.pdf")
    return str(tmp_path / "ocr.pdf"), segs


def test_ocr_pdf_blocks_header_page_number_dash_items_toc_line(tmp_path):
    import docx_layout as DL
    src, segs = ocr_doc(tmp_path)
    blocks, info = DL.load(src, segs)
    assert info["font"] == "Courier New" and info["size"] == 8.5             # один кегль на документ, не оценка OCR
    heads = [b for b in blocks if b["k"] == "h"]
    assert len(heads) == 3                                                   # название раздела — один раз
    lists = [b for b in blocks if b["k"] == "list"]
    assert lists and all(b["kind"] == "Dash" for b in lists)
    toc = next(b for b in blocks if b["k"] == "p" and "Tengely" in "".join(r["t"] for r in b["runs"]))
    assert [r["t"] for r in toc["runs"]] == ["1.4.\t", "Tengelyelmozdulás megnövekedése\t", "13"]
    assert toc["runs"][1].get("rtab") and toc.get("seg")
    L = DL.layout(src, blocks, segs, info)
    assert L["page_tok"] and L["head_grid"]["rows"][0]["cells"][0] == "Azonosító: 3SZ19"
    assert L["first_is_body"] and not any("Azonosító" in "".join(r["t"] for r in b.get("runs", [])) for b in blocks)


# ---------------------------------------------------------------- 16.10, вечер: разделы, линии таблиц, распознанные PDF
def test_front_matter_section_and_page_numbers_restart(tmp_path):
    """Вводная часть ÜFK (оглавление и лист изменений, подвал «Tartalomjegyzék» без номера) — своим разделом;
    нумерация тела — с числа оригинала («1. oldal» на 7-м листе)."""
    blocks = [p("ÜZEMELTETÉSI FELTÉTELEK", 1, 300), p("Módosítás Nyilvántartás", 2, 100), p("Bevezetés", 3, 100)]
    mem = B.Memory([("ÜZEMELTETÉSI FELTÉTELEK", "ПРЕДЕЛЫ"), ("Módosítás Nyilvántartás", "Лист изменений"),
                    ("Bevezetés", "Введение"), ("Tartalomjegyzék", "Содержание"), ("85. oldal", "Страница 85")])
    L = dict(LAYOUT, stamp=[{"cells": ["85. oldal"], "bold": False}], stamp_widths=[400], stamp_aligns=["right"],
             page_tok="85.", num_start={"page": 3, "value": 1},
             front={"to": 2, "stamp": [{"cells": ["Tartalomjegyzék"], "bold": False}], "stamp_widths": [400],
                    "stamp_aligns": ["center"], "stamp_box": True, "stamp_inner": False})
    out = tmp_path / "f.docx"
    B.Builder(mem, lambda img: png(), PROFILE).build(blocks, str(out), L)
    doc = Document(str(out))
    assert len(doc.sections) == 3
    front, body = doc.sections[1], doc.sections[2]
    assert "Содержание" in front.footer._element.xml and "PAGE" not in front.footer._element.xml
    assert "PAGE" in body.footer._element.xml
    assert body._sectPr.find(W + "pgNumType").get(W + "start") == "1"
    assert front._sectPr.find(W + "pgNumType") is None


def test_toc_split_by_pages_is_one_and_headings_outside_it_have_no_level(tmp_path):
    """Оглавление, размеченное по листам, — одно поле (было четыре: 556 пунктов вместо 206); заголовок, которого в
    оглавлении оригинала нет, — без уровня структуры (регистр не важен: «1.1. BEVEZETÉS» / «1.1. Bevezetés»)."""
    toc = lambda keys, lv: {"k": "toc", "info": {"levels": lv, "keys": keys, "extra": [], "bold": [True] * lv,
                                                 "indent": [0] * lv, "figures": False, "page": 2},
                            "geom": {"page": 2}}
    got = D.join_tocs([toc(["1.1.BEVEZETÉS"], 2), toc(["4.1.1.Areaktor"], 3)])
    assert len(got) == 1 and got[0]["info"]["levels"] == 3 and len(got[0]["info"]["keys"]) == 2
    h = lambda t, lv: {"k": "h", "level": lv, "runs": [{"t": t, "font": "Tahoma", "size": 12, "b": True}],
                       "geom": {"page": 3, "top": 100}}
    mem = B.Memory([("1.1. Bevezetés", "1.1. Введение"), ("1.1.1. Megjegyzés", "1.1.1. Примечание"),
                    ("A reaktor", "Реактор")])
    heads = [h("1.1. Bevezetés", 2), h("A reaktor", 3), h("1.1.1. Megjegyzés", 3)]   # «A reaktor» — номер у автонумерации
    out = tmp_path / "t.docx"
    B.Builder(mem, lambda img: png(), PROFILE).build(got + heads, str(out), LAYOUT)
    doc = Document(str(out))
    styles = {x.text.split()[-1]: x.style.name for x in doc.paragraphs if x.text}
    assert styles["Введение"] == "Heading 2" and styles["Реактор"] == "Heading 3"
    assert styles["Примечание"] == "Heading 3 NoTOC"
    assert doc.element.body.xml.count("TOC \\o") == 1
    # пунктов оглавления среди заголовков нашлось меньше 60 % — сверка выключена: поле TOC не пустеет
    B.Builder(mem, lambda img: png(), PROFILE).build(got + [h("1.1.1. Megjegyzés", 3)], str(out), LAYOUT)
    assert not any(x.style.name.endswith("NoTOC") for x in Document(str(out)).paragraphs)
    # а лишние заголовки (абзацы стилем заголовка, ИН-АСУТП-1172) сверку не выключают
    many = heads + [h(f"Abszatz {i} szövege", 2) for i in range(6)]
    B.Builder(mem, lambda img: png(), PROFILE).build(got + many, str(out), LAYOUT)
    styles = [x.style.name for x in Document(str(out)).paragraphs if x.text.startswith("Abszatz")]
    assert styles and all(st.endswith("NoTOC") for st in styles)


def test_cell_edges_follow_original_rules_and_builder_draws_them(tmp_path):
    """Линии вокруг ячеек — как в оригинале: у ÜFK нет внешних вертикалей и линии между «E.1.» и текстом; Word
    рисовал полную сетку."""
    g = lambda x0, x1, top: {"page": 5, "x0": x0, "x1": x1, "top": top, "bottom": top + 10, "first_x0": x0,
                             "size": 10, "lines": 1}
    rows = [[{"blocks": [p("Ellenőrzés", 5, 100, 150)], "th": True, "geom": g(150, 230, 100)},
             {"blocks": [p("Ciklusidő", 5, 100, 420)], "th": True, "geom": g(420, 500, 100)}],
            [{"blocks": [p("E.1.", 5, 120, 80)], "th": False, "geom": g(80, 100, 120)},
             {"blocks": [p("Ellenőrizni kell", 5, 120, 120)], "th": False, "geom": g(120, 300, 120)},
             {"blocks": [p("8 óra", 5, 120, 420)], "th": False, "geom": g(420, 450, 120)}]]
    hs = [(96, 75, 520), (116, 75, 520), (135, 75, 520)]                  # над шапкой, между строками, под таблицей
    vs = [(410, 96, 135)]                                                  # одна внутренняя вертикаль
    D.cell_edges(rows, lambda pn: (hs, vs))
    e = rows[1]
    assert e[0]["edges"] == {"top": True, "bottom": True, "left": False, "right": False}
    assert e[1]["edges"]["right"] and e[2]["edges"]["left"] and not e[2]["edges"]["right"]
    mem = B.Memory([("Ellenőrzés", "Контроль"), ("Ciklusidő", "Периодичность"), ("E.1.", "E.1."),
                    ("Ellenőrizni kell", "Необходимо проверить"), ("8 óra", "8 ч")])
    rows[0][0]["span"], rows[0][1]["span"] = 2, 1
    for c in rows[1]:
        c["span"] = 1
    tbl = {"k": "table", "rows": rows, "geom": {"page": 5, "top": 100}, "borders": True,
           "frame": {"xs": [75, 110, 410, 520]}}
    out = tmp_path / "e.docx"
    B.Builder(mem, lambda img: png(), PROFILE).build([p("x", 5, 80), tbl], str(out), LAYOUT)
    xml = Document(str(out)).tables[0]._tbl.xml
    assert '<w:left w:val="nil"/>' in xml and "w:tcBorders" in xml and '<w:insideV w:val="nil"/>' in xml


def test_ocr_lines_reordered_and_wrapped_toc_number_joined():
    """Строка листа — слева направо, хоть OCR отдал «Turbina…» раньше «1.1.»; номер листа у перенесённого названия —
    на его второй строке — в той же строке оглавления."""
    import docx_layout as DL

    def s(i, t, x0, y0, x1, frags=None):
        frags = frags or [[x0, y0, x1, y0 + 7]]
        return {"id": i, "page": 0, "text": t, "loc": {"bbox": [x0, y0, x1, frags[-1][3]], "frags": frags}}
    segs = [s(11, "Turbina üzemzavari leállítása", 145, 214, 388), s(12, "1.1.", 93, 215, 112),
            s(13, "5", 500, 215, 504), s(17, "1.3.", 93, 238, 112),
            s(18, "Generátor teherledobás üresjáratig", 145, 238, 398, [[145, 238, 398, 245], [145, 250, 230, 257]]),
            s(19, "11", 500, 250, 509)]
    got = [[x["text"] for x in g] for g in DL._lines(segs)]
    assert got == [["1.1.", "Turbina üzemzavari leállítása", "5"], ["1.3.", "Generátor teherledobás üresjáratig", "11"]]


def ocr_page_doc(tmp_path):
    """Лист распознанного PDF: обычный и полужирный набор и курсив (текст тут шрифтом, но сегменты — как от OCR),
    таблица линиями, график — рамкой и кривой с «мусором» OCR."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    segs, y = [], 100
    lines = [[("Check the oil pump and ", "cour"), ("START", "cobo"), (" it then.", "cour")]] + \
        [[(f"Plain words written here as usual {i}.", "cour")] for i in range(12)] + \
        [[("Megjegyzes: ", "cobo"), ("the reactor will be slanted text here", "coit")]]
    for i, parts in enumerate(lines):
        x = 90
        for t, f in parts:
            page.insert_text((x, y + 8), t, fontname=f, fontsize=9)
            x += fitz.get_text_length(t, fontname=f, fontsize=9)
        segs.append({"id": i, "page": 0, "text": "".join(t for t, _ in parts),
                     "loc": {"bbox": [90, y, x, y + 10], "frags": [[90, y, x, y + 10]], "size": 9, "family": "monospace"}})
        y += 14
    for yy in (330, 350, 370):                                             # таблица 2×2
        page.draw_line((90, yy), (400, yy), width=0.7)
    for xx in (90, 250, 400):
        page.draw_line((xx, 330), (xx, 370), width=0.7)
    for i, (t, x0, y0) in enumerate((("Alfa", 100, 336), ("Beta", 260, 336), ("Gamma", 100, 356), ("Delta |", 260, 356))):
        page.insert_text((x0, y0 + 8), t, fontname="cour", fontsize=9)
        segs.append({"id": 100 + i, "page": 0, "text": t,
                     "loc": {"bbox": [x0, y0, x0 + 40, y0 + 10], "frags": [[x0, y0, x0 + 40, y0 + 10]], "size": 9,
                             "family": "monospace"}})
    page.draw_rect(fitz.Rect(150, 420, 450, 700), width=0.8)               # график
    page.draw_polyline([fitz.Point(160 + i * 5, 690 - (i * i) % 250) for i in range(58)], width=1.5)
    segs.append({"id": 200, "page": 0, "text": "oO", "loc": {"bbox": [440, 430, 447, 436], "frags": [[440, 430, 447, 436]],
                                                             "size": 5, "family": "monospace"}})
    doc.save(tmp_path / "o.pdf")
    return str(tmp_path / "o.pdf"), segs


def test_ocr_bold_italic_table_and_figure_from_page(tmp_path):
    """Распознанный лист: полужирное и курсив — по штриху на картинке (текстовый слой у 3SZ19 неполный), таблица — по
    линиям, ячейки — по сетке («|» вертикали прочь), график — картинкой, «мусор» OCR внутри него — не абзац."""
    import docx_layout as DL
    src, segs = ocr_page_doc(tmp_path)
    st = DL._pixel_style(fitz.open(src), segs)
    assert [(t.strip(), b) for t, b, _ in st[0]] == [("Check the oil pump and", False), ("START", True), ("it then.", False)]
    assert st[13][0][1] and not st[13][0][2] and st[13][-1][2] and not st[13][-1][1]
    assert 1 not in st                                                      # обычная строка — без выделения
    blocks, info = DL.load(src, segs)
    tables = [b for b in blocks if b["k"] == "table"]
    assert len(tables) == 1 and [[c["blocks"][0]["runs"][0]["t"] for c in r] for r in tables[0]["rows"]] == \
        [["Alfa", "Beta"], ["Gamma", "Delta"]]
    figs = [b for b in blocks if b.get("figure")]
    assert len(figs) == 1 and figs[0]["runs"][0]["img"]["png"] and figs[0]["geom"]["top"] <= 421
    assert not any("oO" in B.text_of(b.get("runs", [])) for b in blocks)


def test_label_paragraphs_become_list():
    """Абзацы «a.)», «b.)» подряд (дерево Word — простые абзацы) — список с той же меткой; продолжение пункта с
    отступом — в пункт; один «a.)» с отступом под «Korlátozó feltétel:» — тоже список; «f.)» после «a.)–e.)» — в него."""
    bs = [p("Korlátozó feltétel:", 9, 100, 57), p("a.) A fűtőelem tartaléka", 9, 120, 85),
          p("folytatás a pont alatt", 9, 135, 110), p("b.) A kazetták tartaléka", 9, 150, 85), p("Más szöveg", 9, 170, 57),
          p("Korlátozó feltétel:", 10, 100, 57), p("a.) Egyetlen pont", 10, 120, 85)]
    got = D.para_lists(bs)
    assert [b["k"] for b in got] == ["p", "list", "p", "p", "list"]
    assert got[1]["kind"] == "LowerAlpha" and got[1]["suffix"] == ".)" and len(got[1]["items"]) == 2
    assert len(got[1]["items"][0]) == 2                                     # продолжение — в пункте
    lst = {"k": "list", "kind": "LowerAlpha", "suffix": ".)", "items": [[p(f"{c}.) x")] for c in "abcde"]}
    got = D.para_lists([lst, p("f.) Az állapot", 14, 300, 57)])
    assert len(got) == 1 and len(got[0]["items"]) == 6


def test_roman_text_labels_make_roman_list():
    """«• i. …», «• ii. …» — двойной маркер: пункты с римскими метками в тексте — римский список (ÜFK I, EF 2.2.2)."""
    b = lst("Disc", [[p(f"{x}. ellenőrzés")] for x in ("i", "ii", "iii", "iv", "v")])
    D.text_labels([b])
    assert (b["kind"], b["suffix"]) == ("LowerRoman", ".")
    b = lst("Disc", [[p("h. első")], [p("i. második")], [p("j. harmadik")]])
    D.text_labels([b])
    assert b["kind"] == "LowerAlpha"


def test_ocr_line_split_paragraph_is_joined():
    """OCR разрезал абзац по строкам — один абзац: тот же край, вплотную ниже, без точки в конце, со строчной буквы;
    после точки или с заглавной — новый абзац."""
    import docx_layout as DL
    prev = {"k": "p", "runs": [{"t": "célja, hogy a Duna vízhőmérséklete ne", "b": False, "i": False}],
            "geom": {"page": 145, "x0": 127.4, "x1": 547, "top": 274.6, "bottom": 285.6, "lines": 1}, "seg": True}
    seg = lambda t, y: {"id": 1, "page": 144, "text": t, "loc": {"bbox": [127.4, y, 446, y + 11]}}
    assert DL._continues(prev, seg("haladja meg a 29,5 °C-ot", 286.6), 11)
    assert not DL._continues(prev, seg("Haladja meg", 286.6), 11)
    assert not DL._continues(prev, seg("haladja meg", 320), 11)            # не вплотную
    assert DL._continues(prev, seg("2.2 és 2.3 pontok figyelembe vétele mellett:", 286.6), 11)
    assert not DL._continues(prev, seg("2.2. Ha a terhelés", 286.6), 11)    # номер пункта — новый абзац
    DL._join(prev, [{"t": "haladja meg a 29,5 °C-ot", "b": False, "i": False}], {"bottom": 297.6, "x1": 446, "lines": 1})
    assert prev["runs"][0]["t"] == "célja, hogy a Duna vízhőmérséklete ne haladja meg a 29,5 °C-ot"
    assert prev["geom"]["lines"] == 2 and prev["geom"]["bottom"] == 297.6


def test_table_without_cell_padding_in_profile(tmp_path):
    """Профиль без полей ячеек (у документа нет таблиц с текстом в рамке: 3PR42) — таблица собирается, поля обычные."""
    prof = dict(PROFILE, table={"pad_left": None, "pad_top": None, "pad_bottom": None})
    rows = [[{"blocks": [p("Mért érték", 5, 100, 130)], "th": True, "geom": None, "span": 1,
              "edges": {"top": True, "bottom": True, "left": True, "right": True}},
             {"blocks": [p("80 MW", 5, 100, 260)], "th": True, "geom": None, "span": 1,
              "edges": {"top": True, "bottom": True, "left": True, "right": True}}]]
    tbl = {"k": "table", "rows": rows, "geom": {"page": 5, "top": 100}, "borders": True, "frame": {"xs": [120, 250, 400]}}
    out = tmp_path / "t.docx"
    B.Builder(B.Memory([("Mért érték", "Измеренное значение"), ("80 MW", "80 МВт")]), lambda img: png(), prof).build(
        [p("x", 5, 80), tbl], str(out), LAYOUT)
    assert "Измеренное значение" in Document(str(out)).tables[0]._tbl.xml


def test_grid_table_without_full_row_and_courier_body_is_not_code(tmp_path):
    """Сетка, где ни в одной строке нет всех столбцов (разрез в 1-й строке на 200, во 2-й на 250), — столбцов по сетке,
    не падает; основной шрифт Courier (3SZ19) — не «код»: перевод с метками берётся целиком, без Consolas."""
    cell = lambda t, span: {"blocks": [p(t, 5, 100, 100)], "th": False, "geom": None, "span": span,
                            "edges": {"top": True, "bottom": True, "left": True, "right": True}}
    tbl = {"k": "table", "rows": [[cell("Egy", 1), cell("Kettő", 2)], [cell("Három", 2), cell("Négy", 1)]],
           "geom": {"page": 5, "top": 100}, "borders": True, "frame": {"xs": [90, 200, 250, 400]}}
    mono = lambda t, b=False: {"t": t, "font": "Courier New", "size": 9, "b": b, "i": False, "color": "000000",
                               "body_mono": True}
    para = {"k": "p", "runs": [mono("Ha a nyomas magas, "), mono("ZÁRNI", True), mono(" a szelepet.")],
            "geom": {"page": 5, "top": 300, "bottom": 310, "x0": 80, "x1": 400, "first_x0": 80, "size": 9}, "seg": True}
    mem = B.Memory([("Egy", "Один"), ("Kettő", "Два"), ("Három", "Три"), ("Négy", "Четыре"),
                    ("<f1>Ha a nyomas magas, </f1><f2>ZÁRNI</f2><f3> a szelepet.</f3>",
                     "<f1>Если давление высокое, </f1><f2>ЗАКРЫТЬ</f2><f3> клапан.</f3>")])
    out = tmp_path / "g.docx"
    B.Builder(mem, lambda img: png(), PROFILE).build([p("x", 5, 80), tbl, para], str(out), LAYOUT)
    doc = Document(str(out))
    assert len(doc.tables[0].columns) == 3
    par = next(x for x in doc.paragraphs if "ЗАКРЫТЬ" in x.text)
    assert par.text == "Если давление высокое, ЗАКРЫТЬ клапан." and "Consolas" not in par._p.xml
    assert [r.bold for r in par.runs if r.text.strip()] == [None, True, None]


def test_floating_text_box_goes_back_to_its_page():
    """Надпись Word, лежащая в дереве у места привязки (лист 40 между листами 2 и 3, руководство ТЭ357), — к своему
    листу: оглавление снова одно поле."""
    toc = lambda pg: {"k": "toc", "info": {"levels": 2, "keys": [f"k{pg}"], "extra": [], "bold": [True, False],
                                           "indent": [0, 0], "figures": False, "page": pg}, "geom": {"page": pg}}
    bs = [p("Cím", 1, 100), toc(2), p("", 40, 250), p("Текст строки", 40, 254), p("", None), toc(3),
          p("Bevezetés", 4, 100), p("Szöveg", 40, 100), p("Vége", 41, 100)]
    for b in bs:
        if b.get("geom") and b["geom"]["page"] is None:
            b["geom"] = None
    got = D.join_tocs(D.reflow(bs))
    pages = [(b.get("geom") or {}).get("page") for b in got if b["k"] != "toc"]
    assert sum(1 for b in got if b["k"] == "toc") == 1
    assert [x for x in pages if x] == sorted(x for x in pages if x)        # листы по порядку
    texts = [B.text_of(b["runs"]) for b in got if b["k"] == "p"]
    assert texts.index("Текст строки") > texts.index("Bevezetés")


def test_scanned_pdf_keeps_old_path_outline_pdf_goes_by_segments(tmp_path):
    """Распознанный PDF — по сегментам листа, только если это не сканы: у скана формы линии — пиксели, рукопись в
    сегменты не попадает, и форма разваливалась (30RDGT, 30RDIT)."""
    import docx_layout as DL
    scan = fitz.open()
    for _ in range(2):
        pg = scan.new_page(width=595, height=842)
        pg.insert_image(fitz.Rect(40, 60, 555, 800), stream=fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 50, 70), False)
                        .tobytes("png"))
    scan.save(tmp_path / "scan.pdf")
    outline = fitz.open()
    for _ in range(2):
        pg = outline.new_page(width=595, height=842)
        pg.draw_rect(fitz.Rect(100, 100, 300, 120))
    outline.save(tmp_path / "outline.pdf")
    assert not DL.suits(tmp_path / "scan.pdf", [0, 1])
    assert DL.suits(tmp_path / "outline.pdf", [0, 1])
    assert not DL.suits(tmp_path / "outline.pdf", [])                     # не распознан — прежний путь


def _cell(text, page, x0, th=False):
    g = {"page": page, "top": 100.0, "bottom": 110.0, "x0": x0, "x1": x0 + 90, "first_x0": x0, "size": 10, "lines": 1}
    return {"blocks": [p(text, page, x0=x0)] if text else [], "th": th, "geom": g, "attrs": {}}


def test_row_cut_by_page_tail_goes_back_to_its_cells():
    """Строка, разрезанная листом: хвост («необходимо контроль» под 4-й ячейкой шапки) — в ту же ячейку; следующая
    строка «1 | …» нового листа — своя, хотя ячейка выше без точки (синтетика, 16.10)."""
    head = [_cell("Система", 1, 57, True), _cell("Расход", 1, 177, True), _cell("Вода", 1, 298, True),
            _cell("Выполнение уровень", 1, 418, True)]
    tail = [_cell("необходимо контроль", 2, 418)]
    row1 = [_cell("1", 2, 57, True), _cell("Частота датчик.", 2, 177), _cell("Ручной.", 2, 418)]
    got = D._cut_rows([head, tail, row1])
    assert len(got) == 2
    assert B.text_of(got[0][3]["blocks"][0]["runs"]) == "Выполнение уровень необходимо контроль"
    assert B.text_of(got[1][0]["blocks"][0]["runs"]) == "1"


def test_header_without_th_repeated_joins_and_icons_go_to_their_rows():
    """Шапка без TH, повторённая на новом листе, — та же таблица; значки строк прошлого листа, вложенные Word в повтор
    шапки, — в свои ячейки по месту на листе; хвост строки, начатый пунктом списка, — в её список (JPDS табл. 15)."""
    def cell(blocks, page, x0, top):
        return {"blocks": blocks, "th": False, "attrs": {},
                "geom": {"page": page, "top": top, "bottom": top + 12, "x0": x0, "x1": x0 + 60, "size": 10}}
    icon = lambda top: dict(p("", 1, top, x0=12), runs=[{"img": {"x0": 12, "x1": 30}}],
                            geom=dict(p("", 1, top, x0=12)["geom"], x1=30.0))
    hdr = lambda pg: [cell([p("Jel", pg, 50, x0=10)], pg, 10, 50), cell([p("Leírás", pg, 50, x0=100)], pg, 100, 50)]
    item = lambda t, pg, top: lst("bullet", [[p(t, pg, top, x0=100)]])
    a = {"k": "table", "geom": {"page": 1, "top": 50}, "rows": [
        hdr(1), [cell([], 1, 10, 80), cell([p("Ablak címe.", 1, 80, x0=100)], 1, 100, 80)],
        [cell([], 1, 10, 110), cell([p("Rádiógombok:", 1, 110, x0=100), item("– MOD", 1, 124)], 1, 100, 110)]]}
    h2 = hdr(2)
    h2[1]["blocks"] += [icon(80), icon(110)]
    b = {"k": "table", "geom": {"page": 2, "top": 50}, "rows": [
        h2, [cell([item("– ADD", 2, 70)], 2, 100, 70)],
        [cell([], 2, 10, 100), cell([p("Kapcsolók.", 2, 100, x0=100)], 2, 100, 100)]]}
    got = D.adopt(D.join_tables([a, b]))
    assert [x["k"] for x in got] == ["table"]
    rows = got[0]["rows"]
    assert len(rows) == 4                                   # шапка, две строки листа 1, строка листа 2
    assert [sum("img" in q for x in r[0]["blocks"] for q in x["runs"]) for r in rows[1:3]] == [1, 1]
    assert len(rows[2][1]["blocks"][-1]["items"]) == 2      # «– MOD», «– ADD» — один список


def test_caption_in_last_cell_below_text_after_table_leaves_it():
    """Подпись следующего рисунка, вложенная Word в последнюю ячейку таблицы, на том же листе ниже текста после
    таблицы, — прочь из таблицы, на своё место (JPDS л. 69: «Рисунок 60» печаталась в ячейке)."""
    c = {"blocks": [p("Gombok:", 1, 200, x0=100), lst("bullet", [[p("Ábra 60 – Ablak", 1, 735)]])], "th": False,
         "attrs": {}, "geom": {"page": 1, "top": 200, "bottom": 747, "x0": 100, "x1": 500}}
    t = {"k": "table", "geom": {"page": 1, "top": 100}, "rows": [[c]]}
    got = D.evict([t, p("4.3.2.1 Változó", 1, 395), p("Az ablak:", 1, 414)])
    assert [B.text_of(x.get("runs") or []) for x in got[1:]] == ["4.3.2.1 Változó", "Az ablak:", "Ábra 60 – Ablak"]
    assert len(got[0]["rows"][0][0]["blocks"]) == 1


def test_paragraph_cut_by_page_is_one():
    """Абзац, разрезанный листом, Word размечает двумя P — снова один; начатый с прописной после точки — свой."""
    a = p("Пуск регулятор блок уставка время выполняется расход", 3, 700)
    b = p("температура производится частота.", 4, 60)
    c = p("Давление пуск агрегат.", 5, 60)
    got = D.join_pages([a, b, c])
    assert [B.text_of(x["runs"]) for x in got] == [
        "Пуск регулятор блок уставка время выполняется расход температура производится частота.",
        "Давление пуск агрегат."]


def test_numbered_and_bullet_items_in_one_word_list_split():
    """Нумерованный и следом маркированный пункты Word кладёт одним L: вид пункта — по маркеру «•» в его тексте."""
    bullet = p("x", 2, 300)
    bullet["runs"] = [{"t": "•", "font": "Symbol", "size": 10, "b": False, "i": False, "color": "000000"},
                      {"t": " Канал подача", "font": "Arial", "size": 10, "b": False, "i": False, "color": "000000"}]
    got = D.relist([lst("Decimal", [[p("Турбина работа", 2, 260)], [p("Уставка документ", 2, 280)], [bullet]])])
    assert [b["kind"] for b in got] == ["Decimal", "Disc"]
    assert [len(b["items"]) for b in got] == [2, 1]


def test_justified_two_word_line_is_not_tab():
    """Строка из двух слов, растянутая выключкой до правого края ячейки, — не табуляция."""
    def ch(t, x0, top):
        return {"text": t, "x0": x0, "x1": x0 + 5, "top": top, "size": 10, "fontname": "Arial", "ncs": "DeviceGray",
                "non_stroking_color": (0,), "page_number": 1}
    line1 = [ch(c, 177 + 5 * i, 118) for i, c in enumerate("Агрегат")] + [ch(" ", 212, 118)] + \
        [ch(c, 257 + 5 * i, 118) for i, c in enumerate("работа")]
    line2 = [ch(c, 177 + 5 * i, 132) for i, c in enumerate("измерение.")]
    runs = D._runs([{"ch": c} for c in line1 + line2])
    assert "\t" not in B.text_of(runs)


def test_text_from_next_page_nested_in_table_cell_goes_out():
    """Подпись с листа 13, которую Word вложил в список последней ячейки таблицы листа 12, — в поток на своё место."""
    cap = p("Рисунок 1 – Данные", 13, 608)
    cell = {"blocks": [lst("Disc", [[p("Строка таблицы", 12, 190), lst("Disc", [[cap]])]])], "th": False,
            "geom": {"page": 12, "top": 190.0, "bottom": 200.0, "x0": 80, "x1": 300, "first_x0": 80, "size": 10}}
    table = {"k": "table", "rows": [[cell]], "geom": {"page": 12, "top": 180.0, "bottom": 200.0}}
    got = D.evict([table, p("Данные в табличном виде", 13, 76), p("Назначение элементов", 13, 630)])
    texts = [B.text_of(b["runs"]) for b in got if b["k"] == "p"]
    assert texts == ["Данные в табличном виде", "Рисунок 1 – Данные", "Назначение элементов"]
    inner = got[0]["rows"][0][0]["blocks"][0]["items"][0]
    assert [x["k"] for x in inner] == ["p"]                  # пустой вложенный список-обёртка убран


def test_frameless_header_is_paragraphs_with_tabs_and_rule():
    """Шапка без рамки (OAH: номер | лист/всего | версия строкой, название по центру второй строкой, линия снизу) —
    абзацами с частями по местам и нижней границей, а не таблицей в рамке."""
    def chars(text, x0, top, bold=False):
        return [{"text": c, "x0": x0 + 5 * i, "x1": x0 + 5 * i + 5, "top": top, "bottom": top + 9, "doctop": top,
                 "size": 9, "fontname": "OpenSans-Bold" if bold else "OpenSans", "upright": True}
                for i, c in enumerate(text)]
    ch = chars("4.22.sz.", 70, 40) + chars("9/47", 290, 40) + chars("2.Verzió", 486, 40) + \
        chars("Szabványok", 270, 52, bold=True)
    rule = [{"x0": 70.0, "x1": 526.0, "top": 72.0, "bottom": 72.5, "object_type": "rect"}]
    zp = D._zone_paras(ch, rule, 70, 526, page_w=595)
    assert zp is not None
    assert [p[1] for p in zp["rows"][0]["parts"]] == ["left", "center", "right"]
    assert zp["rows"][1]["bold"] and zp["rows"][1]["rule_below"]
    box = rule + [{"x0": 70.0, "x1": 70.5, "top": 36.0, "bottom": 72.0, "object_type": "rect"}]
    assert D._zone_paras(ch, box, 70, 526, page_w=595) is None          # с вертикалью — рамка, таблицей


def test_font_table_gets_font_class(tmp_path):
    """Неустановленный шрифт (Open Sans) — с классом в таблице шрифтов: LibreOffice заменял его шрифтом с засечками."""
    from docx import Document
    d = Document()
    d.add_paragraph("x").runs[0].font.name = "Open Sans"
    f = tmp_path / "f.docx"
    d.save(f)
    B.font_table(f, {"Open Sans": "swiss"})
    import zipfile
    ft = zipfile.ZipFile(f).read("word/fontTable.xml").decode("utf8")
    assert re.search(r'<w:font w:name="Open Sans"><w:charset w:val="00"/><w:family w:val="swiss"/>', ft)


def test_list_item_continued_on_next_page_joins_its_item():
    """Продолжение пункта на следующем листе Word размечает пунктом без метки в новом списке «точками» (OAH N3a.35:
    «d) … csak olyan csőelágazások vagy» | «csövek hegesztésénél…») — это тот же пункт; «e)» дальше — по порядку."""
    d_item = [p("b) Csőcsonkok … csak olyan csőelágazások vagy", 19, 700)]
    cont = [p("csövek hegesztésénél engedhetők meg.", 20, 100)]
    e_item = [p("c) Különböző falvastagságú elemek.", 20, 130)]
    got = D.relist([lst("LowerAlpha", [[p("a) Első.", 19, 600)], d_item]), lst("Disc", [cont, e_item])])
    assert len(got) == 1 and got[0]["kind"] == "LowerAlpha" and len(got[0]["items"]) == 3
    assert B.text_of(got[0]["items"][1][0]["runs"]).endswith("vagy csövek hegesztésénél engedhetők meg.")


def test_symbol_font_dash_marker_makes_dash_list():
    """Маркер шрифта Symbol знаком частной области («\\uf02d» — минус) — тире, а не «•» плюс квадратик."""
    items = [[p("\uf02d a redundáns rendszerek ágainak függetlensége", 39, 600)],
             [p("\uf02d a rendszerelemek függetlensége", 39, 640)]]
    got = D.relist([lst("Disc", items)])
    assert got[0]["kind"] == "Dash"
    assert B.LEAD.match("\uf02d a redundáns")


def test_justified_line_with_real_tab_keeps_tab():
    """Строка «номер ⇥ название…», растянутая по ширине: равные промежутки — растяжка, больший — табуляция."""
    def word(t, x0, top):
        return [{"text": c, "x0": x0 + 5 * i, "x1": x0 + 5 * i + 5, "top": top, "size": 10, "fontname": "Arial",
                 "page_number": 1} for i, c in enumerate(t)]
    line1 = word("MSZ", 130, 100) + word("27003", 165, 100) + word("Nuk", 265, 100) + word("léte", 300, 100) + \
        word("komp", 340, 100)
    line2 = word("elemek", 260, 114)
    runs = D._runs([{"ch": c} for c in line1 + line2])
    t = B.text_of(runs)
    assert t.count("\t") == 1 and t.index("\t") < t.index("Nuk")


def test_size_styles_follow_body_size(tmp_path):
    """Кегль 14 у тела по ГОСТ — «Обычный», а не стиль вводного листа (ИН-ЦАСУТП-1065: 632 абзаца без отступа первой
    строки); у тела 10 pt кегль 14 — по-прежнему вводный."""
    from docx import Document
    def styles(body):
        prof = dict(PROFILE, roles={**PROFILE["roles"], "body": {"font": "Times New Roman", "size": body, "b": False}})
        blocks = [p("Текст тела кеглем четырнадцать пунктов.", 3, 200, x0=62.4), p("Конец.", 3, 230, x0=62.4)]
        for b in blocks:
            b["geom"]["size"] = 14
        out = tmp_path / f"s{body}.docx"
        B.Builder(B.Memory([]), lambda img: None, prof).build(blocks, str(out), dict(LAYOUT, first_is_body=True))
        return [x.style.name for x in Document(out).paragraphs if x.text.startswith("Текст")]
    assert styles(14) == ["Normal"]
    assert styles(10) == ["Front UTS"]


def test_word_2007_pdf_counts_as_word(tmp_path):
    """PDF из Word 2007 («Microsoft® Office Word 2007») — тоже путь по разметке: прежде проверка производителя его не
    узнавала, и руководства OAH шли переводом поверх PDF (живая проверка, 16.10)."""
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Bevezetés a dokumentumhoz")
    doc.set_metadata({"producer": "Microsoft® Office Word 2007", "creator": "Microsoft® Office Word 2007"})
    f = tmp_path / "w2007.pdf"
    doc.save(f)
    seen = []
    orig = D.load
    D.load = lambda path: (seen.append(path), ([], {}))[1]       # до разбора дерева дошло — производитель узнан
    try:
        doc2 = fitz.open(f)
        root = doc2.pdf_catalog()
        doc2.xref_set_key(root, "StructTreeRoot", "<< /Type /StructTreeRoot >>")
        doc2.save(tmp_path / "w2007t.pdf")
        B.coverage(tmp_path / "w2007t.pdf", [])
    finally:
        D.load = orig
    assert seen


def test_word_footnote_goes_to_its_paragraph_as_real_footnote(tmp_path):
    """Сноска Word (Note внизу листа) — настоящей сноской своего абзаца: знак ссылки (мелкая «1») из текста прочь,
    текст сноски — в часть сносок DOCX, а не в абзац (ТЗ ВШЭ, 16.10)."""
    import zipfile
    host = p("Szöveg a lábjegyzettel", 3, 200)
    host["runs"] += [dict(host["runs"][0], t="1", size=6), dict(host["runs"][0], t=" és folytatás.")]
    note = {"k": "note", "runs": [dict(host["runs"][0], t="1 A megjegyzés szövege.", size=8)],
            "geom": dict(host["geom"], top=760)}
    got = D.attach_notes([host, note])
    assert len(got) == 1 and B.text_of(got[0]["runs"]) == "Szöveg a lábjegyzettel és folytatás."
    assert B.text_of(got[0]["notes"][0]["runs"]).strip() == "A megjegyzés szövege."
    out = tmp_path / "fn.docx"
    B.Builder(B.Memory([]), lambda img: None, PROFILE).build(got, str(out), dict(LAYOUT, first_is_body=True))
    with zipfile.ZipFile(out) as z:
        assert "A megjegyzés szövege." in z.read("word/footnotes.xml").decode("utf-8")
        assert "footnoteReference" in z.read("word/document.xml").decode("utf-8")


def test_code_rows_are_cells_and_title_after_number_starts_new_paragraph():
    """Перечень кодов без линеек (3PR42 л. 112): строки столбца — каждая сама по себе, код и значение — разные куски;
    оглавление с номером отдельным куском (3SZ19): название пункта не приклеивается к названию пункта выше, а перенос
    названия без номера слева — приклеивается."""
    import layout as LY
    assert LY.code_line("03/A TJ22D001 KM2/1") and LY.code_line("30BW07")
    assert not LY.code_line("FOT szintcsökkenés") and not LY.code_line("RENDELLENES ÁLLAPOTOK")
    f = lambda t, x0, y0, x1: {"t": t, "bbox": [x0, y0, x1, y0 + 8], "size": 10, "bold": False}
    pars = LY.group([f("04.TH10D001", 100, 158, 168), f("02.TK42D001", 100, 171, 168)], [], [])
    assert len(pars) == 2
    toc = [f("1.3.", 93, 238, 112), f("Generátor teherledobás névleges teljesítményről", 145, 238, 398),
           f("üresjáratig", 145, 250, 200), f("1.5.", 93, 273, 112), f("Gravitációs olajtartályban a szint", 145, 273, 372),
           f("1.6.", 93, 285, 112), f("FOT szintcsökkenés", 145, 284, 242)]
    texts = [x["text"] for x in LY.group(toc, [], [])]
    assert "Generátor teherledobás névleges teljesítményről üresjáratig" in texts
    assert "FOT szintcsökkenés" in texts and "Gravitációs olajtartályban a szint" in texts


def test_grid_without_rules_becomes_borderless_table_toc_line_does_not():
    """Строки из нескольких коротких кусков в общих столбцах (от трёх подряд) — таблица без рамок с кеглем, при котором
    куски влезают в столбцы; строки оглавления (в конце номер листа) — нет."""
    import docx_layout as DL
    def seg(i, x0, y0, t, page=0):
        return {"id": i, "page": page, "text": t, "loc": {"bbox": [x0, y0, x0 + 6 * len(t), y0 + 9]}}
    rows = [[seg(10 * r + k, x, 100 + 14 * r, t) for k, (x, t) in enumerate(
        [(100, "04.TH10D001"), (190, "KM1/1"), (261, "02.TK42D001"), (353, "KM1/1")])] for r in range(3)]
    items = [(g[0]["id"], 0, g) for g in rows]
    got = DL._grids(items, {}, "Times New Roman", 12.0, False)
    assert len(got) == 1 and got[0][1] == 1 and got[0][2]["k"] == "table" and len(got[0][2]["rows"]) == 3
    assert len(got[0][2]["rows"][0]) == 4 and got[0][2]["font_size"] < 12
    toc = [[seg(10 * r, 93, 200 + 12 * r, f"1.{r}."), seg(10 * r + 1, 145, 200 + 12 * r, "Turbina üzemzavar"),
            seg(10 * r + 2, 500, 200 + 12 * r, str(5 + r))] for r in range(4)]
    assert all(it[1] == 0 for it in DL._grids([(g[0]["id"], 0, g) for g in toc], {}, "Times New Roman", 12.0, False))


def test_heading_drawn_as_images_takes_text_of_its_toc_entry():
    """Заголовок картинками (номер и текст — растры; JPDS: 71 заголовок) — текстом своего пункта оглавления по порядку;
    уровень должен совпасть."""
    toc = {"k": "toc", "info": {"figures": False, "entries": [("1 Bevezetés", 1), ("1.1 A cél", 2), ("1.2 Hatály", 2)]}}
    img = {"img": {"x0": 1, "x1": 2, "top": 1, "bottom": 2}}
    h = lambda lv, runs: {"k": "h", "level": lv, "runs": runs, "geom": {"page": 5, "top": 100}}
    blocks = [toc, h(1, [{"t": "1 Bevezetés"}]), h(2, [img, {"t": " "}, img]), h(2, [img])]
    D.toc_headings(blocks)
    assert B.text_of(blocks[2]["runs"]) == "1.1 A cél" and B.text_of(blocks[3]["runs"]) == "1.2 Hatály"


def test_list_at_margin_keeps_its_own_indent(tmp_path):
    """Буквенный список у поля, когда тот же вид в документе обычно с отступом (ÜFK: «a.)» и у поля, и на 30 pt) —
    отступы абзацу по своему месту: метка у поля, текст на висячем отступе."""
    from docx import Document
    def item(t, top):
        x = p(t, 4, top, x0=62.4)
        x["geom"].update(rest_x0=90.4, lines=2)
        return [x]
    b = lst("LowerAlpha", [item("Ha egy korlátozó feltétel nem teljesül, akkor be kell lépni.", 200),
                           item("Az előírt művelet bármikor befejezhető, függetlenül attól.", 240)])
    b["suffix"] = ".)"
    D.list_geo([b])
    assert b["geo"] == {"marker": 62.4, "text": 90.4}
    out = tmp_path / "l.docx"
    B.Builder(B.Memory([]), lambda img: None, PROFILE).build([b], str(out), dict(LAYOUT, first_is_body=True))
    par = next(x for x in Document(out).paragraphs if x.text.startswith("Ha egy"))
    assert abs(par.paragraph_format.left_indent.pt - 28) < 0.5
    assert abs(par.paragraph_format.first_line_indent.pt + 28) < 0.5


def test_ocr_sub_sup_tags_become_indices(tmp_path):
    """«N<sub>R</sub>» распознавания портала — нижний индекс, а не теги в тексте (3PR42)."""
    from docx import Document
    out = tmp_path / "s.docx"
    B.Builder(B.Memory([]), lambda img: None, PROFILE).build([p("Az N<sub>R</sub> = 10 % szinten, m<sup>3</sup>.", 3)],
                                                             str(out), dict(LAYOUT, first_is_body=True))
    par = next(x for x in Document(out).paragraphs if x.text.startswith("Az N"))
    assert par.text == "Az NR = 10 % szinten, m3."
    assert [r.text for r in par.runs if r.font.subscript] == ["R"] and [r.text for r in par.runs if r.font.superscript] == ["3"]


def test_notes_of_both_halves_survive_page_join():
    """Абзац, разрезанный листом, — сноски обеих половин остаются, место — долей общего текста (ревью 17.10)."""
    a = dict(p("Első fele", 3, 700), notes=[{"runs": [{"t": "A"}], "pos": 1.0}])
    b = dict(p("második fele.", 4, 100), notes=[{"runs": [{"t": "B"}], "pos": 0.5}])
    got = D.join_notes(a, b)["notes"]
    assert [B.text_of(n["runs"]) for n in got] == ["A", "B"]
    assert got[0]["pos"] < got[1]["pos"] < 1.0
    assert D.join_notes(p("x"), p("y")) == {}


def test_cell_picture_ignores_page_background_takes_small_icon(tmp_path):
    """Картинка ячейки — только своя и мелкая: фон во весь лист (водяной знак) текст ячеек не заменяет, значок без
    разметки в левой ячейке — картинкой ячейки (ревью 17.10)."""
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8), False)
    pix.clear_with(120)
    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(fitz.Rect(40, 80, 560, 760), stream=pix.tobytes("png"))     # фон
    page.insert_image(fitz.Rect(70, 290, 100, 310), stream=pix.tobytes("png"))    # значок строки 2
    f = tmp_path / "bg.pdf"
    doc.save(f)
    def cell(t, x0, x1, top):
        g = {"page": 1, "top": top, "bottom": top + 12, "x0": x0, "x1": x1, "first_x0": x0, "size": 10}
        return {"blocks": [dict(p(t, 1, top, x0=x0), geom=g)], "th": False, "geom": g}
    rows = [[cell("Kapcsoló", 110, 160, 200), cell("A kapcsoló hosszú leírása, amely a cellát kitölti.", 300, 540, 200)],
            [cell("Jelző", 110, 150, 295), cell("A jelző hosszú leírása, amely a cellát kitölti egészen.", 300, 540, 295)],
            [cell("Gomb", 110, 150, 400), cell("A gomb hosszú leírása, amely a cellát kitölti egészen.", 300, 540, 400)]]
    t = {"k": "table", "rows": rows, "geom": {"page": 1, "top": 200, "bottom": 412, "x0": 110, "x1": 540}}
    with fitz.open(f) as d2:
        D._cell_pictures(t, d2, [t])
    clip = lambda c: [r["img"]["clip"] for x in c["blocks"] for r in x["runs"] if "img" in r]
    assert not clip(rows[0][0]) and not clip(rows[2][0]) and not clip(rows[0][1])
    assert clip(rows[1][0]) and B.text_of(rows[2][1]["blocks"][0]["runs"]).startswith("A gomb")


def test_list_tail_below_next_paragraph_goes_after_it():
    """Рисунок и подпись в последнем пункте, а на листе — ниже абзаца за списком: за этот абзац (JPDS л. 42); пункт,
    чей рисунок выше абзаца, — на месте."""
    pic = lambda top: dict(p("", 2, top), runs=[{"img": {"x0": 70, "x1": 570}}], geom=dict(p("", 2, top)["geom"], bottom=top + 250))
    bs = [lst("bullet", [[p("– Fájlok.", 2, 67)], ]), p("Az adatok megfelelése:", 2, 82), p("Folytatás", 2, 400)]
    bs[0]["items"][0] += [pic(103), lst("num", [[p("Ábra 37 – Megfelelés", 2, 369)]])]
    got = D.list_tails(bs)
    assert [x["k"] for x in got] == ["list", "p", "p", "list", "p"]
    assert B.text_of(got[1]["runs"]) == "Az adatok megfelelése:" and len(got[0]["items"][0]) == 1
    keep = [lst("bullet", [[p("– Fájlok.", 2, 67), pic(80)]]), p("Utána", 2, 340)]
    assert len(D.list_tails(keep)[0]["items"][0]) == 2


def test_figures_take_their_place_among_text():
    """Рисунки с подписями пачкой перед абзацами, что на листе выше них, — по месту; порядок текста не меняется;
    значок рядом по высоте с абзацем — как в разметке (JPDS л. 109)."""
    pic = lambda top, h=50: dict(p("", 2, top), runs=[{"img": {"x0": 200, "x1": 400}}],
                                 geom=dict(p("", 2, top)["geom"], bottom=top + h))
    bs = [p("Hozzáadás", 2, 82), pic(118), p("105. ábra – Sor", 2, 184), pic(241), p("106. ábra – Kijelölt", 2, 305),
          p("Kijelölés", 2, 206), pic(203, 12), p("Másolás", 2, 327)]
    got = [B.text_of(x["runs"]) or "IMG%d" % x["geom"]["top"] for x in D.figures_in_place(bs)]
    assert got == ["Hozzáadás", "IMG118", "105. ábra – Sor", "Kijelölés", "IMG203", "IMG241", "106. ábra – Kijelölt",
                   "Másolás"]


def test_paragraph_inside_table_row_goes_into_its_cell():
    """Надпись ячейки, размеченная Word отдельным абзацем до таблицы (образцы «Текст строки» ТЭ357), — в ячейку по
    строке и столбцу; абзац под таблицей — на месте."""
    def cell(t, top, x0):
        return {"blocks": [p(t, 40, top, x0=x0)] if t else [], "th": False, "attrs": {}}
    t = {"k": "table", "geom": {"page": 40, "top": 204},
         "rows": [[cell("Megjelenés", 204, 70), cell("Leírás", 204, 300)],
                  [cell("", 0, 70), cell("háttér színe", 230, 300)], [cell("", 0, 70), cell("dőlt betű", 300, 300)]]}
    loose = dict(p("Sor szövege", 40, 305, x0=100), geom=dict(p("", 40, 305, x0=100)["geom"], x1=180.0))
    got = D.adopt([loose, t, p("Utána", 40, 400)])
    assert [x["k"] for x in got] == ["table", "p"]
    assert B.text_of(got[0]["rows"][2][0]["blocks"][0]["runs"]) == "Sor szövege"


def test_line_hyphen_joins_word_keeps_suspended_and_compound():
    """Перенос по слогам — слово целиком («акку-» + «ратно», NSU); дефис перед союзом («fűtő- és») и в составном
    («loss-of-» + «coolant») — на месте."""
    import layout as LY
    import docx_layout as DL
    assert LY.join_lines(["написан акку-", "ратно от руки"]) == "написан аккуратно от руки"
    assert LY.join_lines(["a fűtő-", "és hűtőrendszer"]) == "a fűtő- és hűtőrendszer"
    assert LY.join_lines(["loss-of-", "coolant"]) == "loss-of-coolant"
    par = lambda t: {"k": "p", "runs": [{"t": t}], "geom": {"x1": 300, "bottom": 100}}
    for a, b, want in [("написан акку-", "ратно", "написан аккуратно"), ("a fűtő-", "és hűtő", "a fűtő- és hűtő"),
                       ("ÜFK-", "nak", "ÜFK-nak"), ("szöveg", "folytatás", "szöveg folytatás")]:
        p0 = par(a)
        DL._join(p0, [{"t": b}], {"x1": 300, "bottom": 112})
        assert p0["runs"][0]["t"] == want


def test_hanging_indent_gap_is_tab():
    """Промежуток в первой строке до висячего отступа — табуляция («MSZ 27003-0:2013 ⇥ Nukleáris…», OAH), хотя он
    уже порога «1,8 кегля»; обычный пробел — пробел."""
    def ch(t, x0, top):
        return {"text": t, "x0": x0, "x1": x0 + 6, "top": top, "bottom": top + 12, "size": 12.0, "fontname": "Arial"}
    first = [ch(c, 120 + 6 * i, 100) for i, c in enumerate("MSZ 1:2013")] + \
            [ch(c, 241 + 6 * i, 100) for i, c in enumerate("Nukl")]
    rest = [ch(c, 241 + 6 * i, 114) for i, c in enumerate("szab")]
    assert D._hanging_tab(first + rest) == {id(first[10])}
    flat = [ch(c, 120 + 6 * i, 114) for i, c in enumerate("szab")]
    assert not D._hanging_tab(first + flat)


def test_body_text_tagged_heading_is_paragraph_and_levels_by_number():
    """Абзац тела со стилем заголовка (длинный, не полужирный) — абзац; короткий заголовок — заголовок; уровни тегов,
    расходящиеся с номерами («1.1.» — H1, «6.» — H3, ИН-1172), — по глубине номера."""
    h = lambda lv, t, b=True, lines=1: {"k": "h", "level": lv, "runs": [{"t": t, "b": b}],
                                        "geom": {"page": 3, "top": 100, "bottom": 112, "lines": lines}}
    bs = [h(3, "6.\tAlgoritmusok."), h(1, "1.1. A dokumentum célja."), h(2, "Ez az utasítás a turbina szabályozójának "
          "működését írja le részletesen, minden üzemállapotban.", b=False, lines=2),
          h(1, "1.2. Az utasítás ismerete kötelező a következő személyeknek:", b=False), h(2, "Rövid cím", b=False)]
    got = D.real_headings(bs)
    assert [b["k"] for b in got] == ["h", "h", "p", "p", "h"]
    assert got[0]["level"] == 1 and got[1]["level"] == 2


def test_icon_after_list_marker_tab_stays_at_item_start():
    """Значок на табуляции сразу после маркера пункта («•⇥[значок] : …», UTS) — в начале текста, не «на табуляции»
    (маркер с табуляцией сборка снимает — значок уезжал в конец пункта); на табуляции внутри текста — at_tab."""
    item = dict(p("", 42, 219), runs=[{"t": "•"}, {"t": "\t", "tabs": [98.0]}, {"t": ": a tesztlépés"}])
    icon = dict(p("", 42, 219), runs=[{"img": {"x0": 98.4, "x1": 106, "top": 219, "bottom": 227}}])
    got = D.inline_pictures([item, icon])
    assert len(got) == 1 and "img" in got[0]["runs"][2] and not got[0]["runs"][2].get("at_tab")
    par = dict(p("", 82, 300), runs=[{"t": "Нажать\t", "tabs": [160.0]}, {"t": "для выхода"}])
    esc = dict(p("", 82, 300), runs=[{"img": {"x0": 140, "x1": 160, "top": 300, "bottom": 310}}])
    got = D.inline_pictures([par, esc])
    assert got[0]["runs"][1].get("at_tab")


def test_portal_roles_split_merge_and_mark_paragraphs():
    """Гибрид (roles.apply): абзац правил делится на строке, которую портал начал новым абзацем; два абзаца правил
    склеиваются, если портал продолжает строкой подряд; роль — первой строки; колонтитул и таблица не делятся и не
    склеиваются (роль колонтитула — для распознанных листов)."""
    import roles as R
    from layout import join_lines
    ln = lambda i, y, t: {"id": i, "x0": 70.0, "x1": 500.0, "top": y, "bottom": y + 12, "size": 10, "bold": False, "text": t}
    fr = lambda y, t: {"t": t, "bbox": [70.0, y, 500.0, y + 12], "size": 10, "bold": False}
    lines = [ln(0, 100, "4.2 Kezelés"), ln(1, 114, "A kezelő a panelen"), ln(2, 128, "dolgozik."),
             ln(3, 142, "a) első pont"), ln(4, 156, "folytatás"), ln(5, 800, "12. oldal")]
    labels = {0: {"new": True, "role": "heading"}, 1: {"new": True, "role": "body"}, 2: {"new": False, "role": "body"},
              3: {"new": True, "role": "list"}, 4: {"new": False, "role": "list"}, 5: {"new": True, "role": "header"}}
    pars = [{"frags": [fr(100, "4.2 Kezelés"), fr(114, "A kezelő a panelen")]},          # правила склеили заголовок
            {"frags": [fr(128, "dolgozik.")]},                                         # и разорвали абзац
            {"frags": [fr(142, "a) első pont"), fr(156, "folytatás")]}, {"frags": [fr(800, "12. oldal")]}]
    got = R.apply(pars, lines, labels, join_lines)
    assert [p["text"] for p in got] == ["4.2 Kezelés", "A kezelő a panelen dolgozik.", "a) első pont folytatás",
                                        "12. oldal"]
    assert [p.get("role") for p in got] == ["heading", None, "list", "header"]      # колонтитул — для сканов
    assert R.apply(pars, lines, {}, join_lines) is pars                     # нет меток — как у правил


def test_portal_roles_running_header_needs_same_height():
    """Колонтитул по повтору — только строка на одной высоте у края на многих листах; подзаголовок «Korlátozó feltétel:»
    вверху разных листов на разной высоте — не колонтитул."""
    import roles as R
    doc = fitz.open()
    for k in range(6):
        pg = doc.new_page()
        pg.insert_text((70, 40), "Üzemeltetési Feltételek és Korlátok", fontsize=9)
        pg.insert_text((70, 70 + 9 * k), "Korlátozó feltétel:", fontsize=10)
        pg.insert_text((70, 400), f"Szöveg {k}", fontsize=10)
    rep = R.repeats(doc)
    lines = R.page_lines(doc[2])
    H = doc[2].rect.height
    flags = {l["text"]: R.is_rep(l, rep, H) for l in lines}
    assert flags["Üzemeltetési Feltételek és Korlátok"] and not flags["Korlátozó feltétel:"]


def test_code_paragraph_keeps_keys_translates_comments(tmp_path):
    """Абзац-листинг моноширинным: строка без венгерских букв («LogFile=…») — как в оригинале, комментарий с венгерскими
    — переводом; так же, если весь абзац — одна моноширинная группа (UTS л. 94, 19.10)."""
    assert B.code_part("LogFile=/tmp/canframes.log", "Файл журнала=/tmp/canframes.log") == "LogFile=/tmp/canframes.log"
    src = "[Connections]\n;Új alapértékek\nLocalPort=12000"
    ru = "[Подключения]\n;Новые значения по умолчанию\nЛокальный порт=12000"
    assert B.code_part(src, ru) == "[Connections]\n;Новые значения по умолчанию\nLocalPort=12000"
    run = {"t": "LogFile=/tmp/canframes.log", "font": "Courier New", "size": 9, "b": False, "i": False, "color": "000000"}
    out = tmp_path / "c.docx"
    B.Builder(B.Memory([("LogFile=/tmp/canframes.log", "Файл журнала=/tmp/canframes.log")]), lambda img: None,
              PROFILE).build([dict(p("x", 3), runs=[run])], str(out), dict(LAYOUT, first_is_body=True))
    assert any(par.text == "LogFile=/tmp/canframes.log" for par in Document(out).paragraphs)


def test_vector_tree_with_labels_is_figure_text_paragraph_is_not():
    """Схема из линий с подписями (дерево Eskom) на текстовом слое — рисунок областью листа; абзац текста, задетый
    скоплением, — не рисунок."""
    import docx_layout as DL
    doc = fitz.open()
    page = doc.new_page()
    page.draw_line((150, 100), (150, 300))
    for y in (100, 150, 200, 250, 300):
        page.draw_line((150, y), (250, y))
    page.draw_line((130, 200), (150, 200))
    seg = lambda i, r, t: {"id": i, "page": 0, "text": t, "loc": {"bbox": list(r)}}
    segs = [seg(0, (60, 194, 120, 206), "1. Scope"), seg(1, (175, 94, 240, 106), "General"),
            seg(2, (175, 244, 260, 256), "Normative"), seg(3, (60, 400, 500, 412), "Body text paragraph.")]
    got = DL._vector_figures(page, 0, page.get_cdrawings(), segs, segs, [], [])
    assert len(got) == 1 and got[0][1]["page_number"] == 1 and got[0][0].y1 < 400
    segs.append(seg(4, (100, 140, 400, 160), "A long body line crossing the drawing"))
    assert DL._vector_figures(page, 0, page.get_cdrawings(), segs, segs, [], []) == []


def test_cell_fill_legend_stays_in_its_row(tmp_path):
    """Легенда цветов: заливки соседних строк касаются — картинка ячейки только своей строки, не весь столбец (JPDS:
    9 копий столбца, +9 листов, 17.10)."""
    doc = fitz.open()
    page = doc.new_page()
    for k, col in enumerate([(1, 0, 0), (0, 1, 0), (0, 0, 1)]):
        page.draw_rect(fitz.Rect(72, 200 + 31 * k, 318, 230 + 31 * k), fill=col, color=None)
        page.draw_rect(fitz.Rect(72, 230 + 31 * k, 318, 231 + 31 * k), fill=(0, 0, 0), color=None)   # граница строк
    f = tmp_path / "legend.pdf"
    doc.save(f)
    def cell(t, top):
        g = {"page": 1, "top": top, "bottom": top + 12, "x0": 330, "x1": 400, "first_x0": 330, "size": 10}
        return {"blocks": [dict(p(t, 1, top, x0=330), geom=g)], "th": False, "geom": g}
    rows = [[{"blocks": [], "th": False}, cell(t, 209 + 31 * k)] for k, t in enumerate(["Nyitott", "Kész", "Lezárt"])]
    t = {"k": "table", "rows": rows, "geom": {"page": 1, "top": 200, "bottom": 293, "x0": 72, "x1": 400}}
    with fitz.open(f) as d2:
        D._cell_pictures(t, d2, [t])
    clips = [[r["img"]["clip"] for x in r0[0]["blocks"] for r in x["runs"] if "img" in r] for r0 in rows]
    assert all(len(c) == 1 for c in clips)
    assert all(c[0][3] - c[0][1] < 40 for c in clips)                       # строка, а не столбец в 93 pt


def test_heading_number_dot_follows_original(tmp_path):
    """Точка после номера заголовка — по уровню, как в оригинале: «1. Introduction», а «2.1 Scope» — без (Eskom)."""
    from docx import Document
    h = lambda lv, t, top: {"k": "h", "level": lv, "runs": [{"t": t, "font": "Arial", "size": 12, "b": True}],
                            "geom": {"page": 2, "top": top, "bottom": top + 12, "x0": 62.4, "x1": 300, "first_x0": 62.4,
                                     "size": 12}}
    blocks = [h(1, "1. Introduction", 100), h(2, "1.1 Scope", 140), h(1, "2. Supporting Clauses", 180),
              h(2, "2.1 Purpose", 220)]
    out = tmp_path / "h.docx"
    B.Builder(B.Memory([]), lambda img: None, PROFILE).build(blocks, str(out), dict(LAYOUT, first_is_body=True))
    num = Document(out).part.numbering_part.element
    lv = {l.get(W + "ilvl"): l.find(W + "lvlText").get(W + "val")
          for a in num.findall(W + "abstractNum") if a.get(W + "abstractNumId") == "900" for l in a.findall(W + "lvl")}
    assert lv["0"] == "%1." and lv["1"] == "%1.%2"
    assert [p.text for p in Document(out).paragraphs if p.text][:2] == ["Introduction", "Scope"]


def test_vector_chart_with_axis_labels_becomes_one_picture(tmp_path):
    """График Word из векторов (области, рамки) с подписями осей абзацами — одной картинкой области листа; подписи из
    потока прочь; рамка вокруг обычного абзаца — не рисунок (ÜFK I, л. 48)."""
    doc = fitz.open()
    page = doc.new_page()
    page.draw_rect(fitz.Rect(100, 300, 400, 500), color=(0.5, 0.5, 0.5), width=1)
    page.draw_rect(fitz.Rect(100, 300, 400, 420), color=None, fill=(0.6, 0.8, 0.3))
    page.draw_rect(fitz.Rect(100, 420, 400, 500), color=None, fill=(0.9, 0.95, 0.85))
    page.draw_rect(fitz.Rect(90, 600, 500, 700), color=(0.5, 0.5, 0.5), width=1)        # рамка вокруг абзаца
    f = tmp_path / "chart.pdf"
    doc.save(f)
    blocks = [p("Az alábbi ábra szerinti tartományon belül kell lennie.", 1, 270, x0=90),
              p("250", 1, 296, x0=80), p("0", 1, 495, x0=88), p("1,E-04 10 20 30 40 50 60 70 80 90 100", 1, 503, x0=95),
              p("Hosszú szöveg a keretben, amely nem ábra, hanem közönséges bekezdés a dokumentumban.", 1, 640, x0=95)]
    for b in blocks:
        b["geom"]["x1"] = b["geom"]["x0"] + 6 * len(B.text_of(b["runs"])) if len(B.text_of(b["runs"])) < 40 else 480
        b["geom"]["bottom"] = b["geom"]["top"] + 9
    got = D.compose_figures(blocks, str(f))
    figs = [b for b in got if b.get("vector")]
    assert len(figs) == 1 and figs[0]["geom"]["top"] <= 300 and figs[0]["geom"]["bottom"] >= 505
    texts = [B.text_of(b["runs"]) for b in got if not b.get("vector")]
    assert "250" not in texts and texts[0].startswith("Az alábbi") and texts[-1].startswith("Hosszú")


def test_table_column_not_narrower_than_its_picture(tmp_path):
    """Значок в ячейке шире столбца по левым краям текста — столбец расширяется (JPDS: значки окна срезались)."""
    from docx import Document
    img = {"img": {"x0": 130, "x1": 154, "top": 540, "bottom": 564, "page_number": 3}}
    def cell(runs, x0, x1):
        g = {"page": 3, "top": 540, "bottom": 564, "x0": x0, "x1": x1, "first_x0": x0, "size": 10}
        return {"blocks": [{"k": "p", "runs": runs, "geom": g}], "th": False, "geom": g}
    rows = [[cell([img], 130, 154), cell([{"t": "«развернуть» – окно на весь экран", "font": "Arial", "size": 10}],
                                         148, 500)] for _ in range(2)]
    t = {"k": "table", "rows": rows, "borders": False, "frame": {"xs": []},
         "geom": {"page": 3, "top": 540, "bottom": 600, "x0": 130, "x1": 500, "size": 10}}
    out = tmp_path / "t.docx"
    B.Builder(B.Memory([]), lambda im: png(), PROFILE).build([t], str(out), dict(LAYOUT, first_is_body=True))
    grid = Document(out).tables[0]._tbl.tblGrid.findall(W + "gridCol")
    assert int(grid[0].get(W + "w")) / 20 >= 24 + 12 - 0.5


def test_raised_and_lowered_small_chars_become_unicode_indices():
    """Индекс — мелкий знак выше или ниже строки: «545 ºс» → «545 °с» (ИН-1065), «dm3» → «dm³»; обычные знаки — как
    есть; соседние элементы разметки без промежутка («Core i» + «5», ТЭ357) — без пробела."""
    def ch(t, x0, top, size=12.0, w=6.0):
        return {"text": t, "x0": x0, "x1": x0 + w, "top": top, "bottom": top + size, "size": size,
                "fontname": "Times New Roman", "page_number": 1}
    items = [{"ch": ch("5", 100, 100)}, {"ch": ch("4", 106, 100)}, {"ch": ch("5", 112, 100)},
             {"ch": ch(" ", 118, 100, w=3)}, {"ch": ch("о", 121, 96, size=8, w=4)}, {"ch": ch("с", 125, 100)},
             {"ch": ch(" ", 131, 100, w=3)}, {"ch": ch("d", 134, 100)}, {"ch": ch("m", 140, 100)},
             {"ch": ch("3", 146, 97, size=7, w=4)}]
    assert B.text_of(D._runs(items)) == "545 °с dm³"
    prev = dict(p("процессор Core", 8, 600), runs=[{"t": "процессор Core", "x": 200.0, "x1": 256.0, "top": 616.7,
                                                    "page": 8, "size": 12.0}])
    nxt = dict(p("i5 c тактовой частотой", 8, 600), inline=True,
               runs=[{"t": "i5 c тактовой частотой", "x": 259.0, "x1": 380.0, "top": 616.7, "page": 8, "size": 12.0}])
    glued = dict(nxt, runs=[dict(nxt["runs"][0], x=256.5)])
    assert B.text_of(B.merge_inline([prev, nxt])[0]["runs"]) == "процессор Core i5 c тактовой частотой"
    assert B.text_of(B.merge_inline([dict(prev, runs=[dict(prev["runs"][0])]), glued])[0]["runs"]) == \
        "процессор Corei5 c тактовой частотой"


def test_icon_after_paragraph_goes_to_its_tab_in_text(tmp_path):
    """Значок, размеченный отдельным рисунком после абзаца, но стоящий в его строке на месте табуляции (клавиша ESC,
    ИН-1065), — в абзац на место табуляции: текст целиком, значок между словами."""
    from docx import Document
    para = p("при помощь клавиши\tна встроенной клавиатуре, с помощью кнопок", 82, 462)
    para["geom"]["bottom"] = 563
    para["runs"][0]["tabs"] = [217.61]
    icon = {"k": "p", "runs": [{"img": {"x0": 168, "x1": 217, "top": 474, "bottom": 505, "page_number": 82}}],
            "geom": {"page": 82, "top": 474, "bottom": 505, "x0": 168, "x1": 217, "first_x0": 168, "size": 10}}
    got = D.inline_pictures([para, icon])
    assert len(got) == 1 and any(r.get("at_tab") for r in got[0]["runs"])
    out = tmp_path / "esc.docx"
    B.Builder(B.Memory([]), lambda im: png(), PROFILE).build(got, str(out), dict(LAYOUT, first_is_body=True))
    par = next(x for x in Document(out).paragraphs if "клавиши" in x.text)
    assert par.text.startswith("при помощь клавиши") and "встроенной клавиатуре" in par.text and "\t" not in par.text
    body = par._p.xml
    assert body.index("клавиши") < body.index("<w:drawing") < body.index("встроенной")


def test_heading_from_memory_keeps_its_color(tmp_path):
    """Заголовок, текст которого взят из памяти целиком (русский документ, перевод), — цвет общий для его текста в
    оригинале (синие «Приложение 1 …» ИН-1172 выходили чёрными)."""
    from docx import Document
    h = {"k": "h", "level": 1, "runs": [{"t": "Приложение 1 ПОРЯДОК ПОДГОТОВКИ", "font": "Calibri", "size": 13.0,
                                         "b": False, "i": False, "color": "2e74b5"}],
         "geom": {"page": 3, "top": 100, "bottom": 112, "x0": 62.4, "x1": 400, "first_x0": 62.4, "size": 13}}
    out = tmp_path / "hc.docx"
    B.Builder(B.Memory([]), lambda im: None, PROFILE).build([h], str(out), dict(LAYOUT, first_is_body=True))
    run = next(x for x in Document(out).paragraphs if x.text).runs[0]
    assert str(run.font.color.rgb) == "2E74B5"


def test_open_sided_table_edges_from_rules():
    """Таблица без боковых рамок (ÜFK, скан): внутренние вертикали есть, крайних нет — края по горизонталям, что
    одинаково выходят за вертикали с обеих сторон; без таких горизонталей (две линии и черта между ними) — не таблица."""
    import docx_layout as DL
    hs = [(221, 57, 542), (252, 57, 542), (473, 57, 542), (694, 56, 542)]
    vs = [(228, 220, 695), (461, 220, 695)]
    t = DL._tables(hs, vs)
    assert len(t) == 1 and [round(x) for x in t[0][0]] == [56, 228, 461, 542]
    assert [round(y) for y in t[0][1]] == [220, 252, 473, 694]
    assert DL._tables([(221, 57, 542), (473, 200, 300)], vs) == []


def test_blank_continuation_joins_with_row_lines():
    """Продолжение таблицы без текста (пустые графы листа ознакомления ИН-1065) — к своей таблице, линии строк обеих
    частей — одной рамкой: по ним высота пустых строк."""
    def cell(t=""):
        return {"blocks": [{"k": "p", "runs": [{"t": t}]}] if t else [], "th": False}
    a = {"k": "table", "rows": [[cell("№"), cell("Фамилия")], [cell(), cell()]], "geom": {"page": 84, "top": 149},
         "frame": {"xs": [70, 106, 248], "ys": [140, 175, 200]}}
    b = {"k": "table", "rows": [[cell(), cell()], [cell(), cell()]], "geom": {"page": 85, "top": 60},
         "frame": {"xs": [70, 106, 248], "ys": [60, 85, 110]}}
    got = D.join_tables([a, b])
    assert len(got) == 1 and len(got[0]["rows"]) == 4
    assert got[0]["frame"]["ys"] == [140, 175, 200, 225, 250]


def test_padded_labels_become_sub_rows():
    """ÜFK «Beavatkozások»: метки «A.1.», «A.2.1.» выровнены пустыми абзацами против своих действий — подстроки по
    меткам; ячейка условия — одна на все подстроки (vcont), линий между подстроками нет."""
    def p(t, top):
        return {"k": "p", "runs": [{"t": t}], "geom": {"top": top, "page": 70}}
    blank = {"k": "p", "runs": []}
    e = {"top": True, "bottom": True, "left": True, "right": True}
    row = [{"blocks": [p("A.", 256)], "th": False, "edges": e},
           {"blocks": [p("A PVS egyetlen", 256), p("VAGY", 295)], "th": False, "edges": e},
           {"blocks": [p("A.1.", 256), blank, blank, p("A.2.1.", 295)], "th": False, "edges": e},
           {"blocks": [p("Helyre kell", 256), p("ÉS", 282), p("A reaktor", 295)], "th": False, "edges": e},
           {"blocks": [p("A.1. 72 óra", 256), blank, p("A.2.1. 2 óra", 295)], "th": False, "edges": e}]
    t = {"k": "table", "rows": [row]}
    D.sub_rows([t])
    txt = lambda c: [D.text_of(x["runs"]) for x in c["blocks"]]
    assert len(t["rows"]) == 2
    assert [txt(c) for c in t["rows"][0]] == [["A."], ["A PVS egyetlen", "VAGY"], ["A.1."], ["Helyre kell", "ÉS"],
                                              ["A.1. 72 óra"]]
    assert [c.get("vcont", False) for c in t["rows"][1]] == [True, True, False, False, False]
    assert [txt(c) for c in t["rows"][1][2:]] == [["A.2.1."], ["A reaktor"], ["A.2.1. 2 óra"]]
    assert not t["rows"][0][3]["edges"]["bottom"] and not t["rows"][1][3]["edges"]["top"]
    assert t["rows"][1][0]["edges"]["bottom"] and not t["rows"][0][0]["edges"]["bottom"]


def test_page_split_paragraph_joins_on_full_last_line():
    """Абзац, разрезанный листом, когда продолжение начато прописной (имя собственное: «…az Országos | Atomenergia
    Hivatal…», OAH 1.1): последняя строка на листе во всю ширину — склейка; короткая последняя строка — нет."""
    def p(t, page, x1, last_x1, lines=3):
        return {"k": "p", "runs": [{"t": t}], "geom": {"page": page, "top": 700 if page == 6 else 85, "x0": 99,
                                                        "first_x0": 99, "x1": x1, "last_x1": last_x1, "lines": lines}}
    got = D.join_pages([p("feladatköréről, az Országos", 6, 524, 524), p("Atomenergia Hivatal hatósági", 7, 524, 498)])
    assert len(got) == 1 and D.text_of(got[0]["runs"]) == "feladatköréről, az Országos Atomenergia Hivatal hatósági"
    got = D.join_pages([p("A telephely engedély", 6, 524, 300), p("Az engedély hatálya", 7, 524, 498)])
    assert len(got) == 2


def test_underline_from_thin_rule_under_glyphs():
    """Подчёркивание — тонкая линия ниже базовой линии внутри рамки знака (OAH 4.22: 297,8–298,4 при знаках 287,7–
    299,7); граница ячейки по низу рамки знака (ÜFK) и зачёркивание посередине — не подчёркивание."""
    ch = lambda x, top=287.7: {"text": "a", "x0": x, "x1": x + 6, "top": top, "bottom": top + 12, "size": 12}
    a, b, c = ch(121), ch(400), ch(121, top=208.3)
    D._underline([a, b, c], [{"top": 297.8, "bottom": 298.4, "height": 0.6, "width": 234.8, "x0": 120.6, "x1": 355.4},
                             {"top": 220.0, "bottom": 220.5, "height": 0.5, "width": 300, "x0": 57, "x1": 357}])
    assert a.get("u") and not b.get("u") and not c.get("u")


def test_flush_paragraph_without_red_line(tmp_path):
    """Документ с красной строкой: абзац, первая строка которого у поля, — стилем без отступа первой строки
    («Пояснительная записка должна содержать…» НГТУ); абзац с красной строкой — обычным стилем тела."""
    blocks = [{"k": "h", "level": 1, "runs": [{"t": "1 Bevezetés", "font": "Tahoma", "size": 14, "b": True}],
               "geom": {"page": 2, "top": 90}},
              p("A dokumentum tartalmazza:", 2, 120, x0=62.4), p("Az UTS alkalmazás leírása.", 2, 140, x0=80.0)]
    mem = B.Memory([("Bevezetés", "Введение"), ("A dokumentum tartalmazza:", "Документ содержит:"),
                    ("Az UTS alkalmazás leírása.", "Описание приложения UTS.")])
    out = tmp_path / "f.docx"
    B.Builder(mem, lambda img: None, dict(PROFILE, first_line=17.0)).build(blocks, str(out), LAYOUT)
    st = {x.text: x.style.name for x in Document(str(out)).paragraphs}
    assert st["Документ содержит:"] == "Body Flush UTS" and st["Описание приложения UTS."] == "Normal"


def test_hanging_tab_and_stretch_after_it():
    """«MSZ EN 571-1:2001 ⇥ Roncsolásmentes ⟷ vizsgálatok.» (OAH 4.22): табуляция — к висячему отступу (там начаты
    прочие строки), широкий промежуток после неё — растяжка по ширине, не вторая табуляция."""
    def word(t, x0, top, size=11.0):
        out, x = [], x0
        for ch in t:
            w = 3.2 if ch == " " else 6.8
            out.append({"text": ch, "x0": x, "x1": x + w, "top": top, "size": size}); x += w
        return out
    line1 = word("MSZ EN 571-1:20", 99.3, 228) + word("Roncsolás ", 241.1, 228) + word("vizsgálatok.", 460.3, 228)
    line2 = word("Folyadékbehatolásos vizsgálat.", 241.1, 244)
    line3 = word("alapelvek", 241.1, 260)
    chars = line1 + line2 + line3
    for i, c in enumerate(chars):
        c["fontname"] = "Open Sans"; c["bottom"] = c["top"] + c["size"]
    runs = D._runs([{"ch": c} for c in chars])
    t = D.text_of(runs)
    assert t.startswith("MSZ EN 571-1:20\tRoncsolás vizsgálatok.") and t.count("\t") == 1


def test_running_section_title_not_in_header_grid(tmp_path):
    """Название раздела в рамке (3SZ19) портал метит колонтитулом, но оно меняется от раздела к разделу — не в сетку
    шапки: в шапке всех листов стояло название другого раздела, шапка ломалась на три столбца."""
    import docx_layout as DL
    src, segs = ocr_doc(tmp_path)
    for s in segs:
        if s["text"][:2] in ("1.", "2.", "3.") and s["loc"]["bbox"][1] == 86:
            s["loc"]["role"] = "header"
    blocks, info = DL.load(src, segs)
    L = DL.layout(src, blocks, segs, info)
    cells = [c for r in L["head_grid"]["rows"] for c in r["cells"]]
    assert "Azonosító: 3SZ19" in cells and not any("üzemzavarai" in c or "Szekunderköri" in c for c in cells)


def test_running_section_title_once_with_portal_labels(tmp_path):
    """Название раздела в рамке с меткой портала на всех листах и с разночтениями OCR — в тексте один раз, у начала
    раздела (3SZ19: заголовок второго раздела пропадал в начале и оставался посреди)."""
    import docx_layout as DL
    src, segs = ocr_doc(tmp_path)
    for s in segs:
        if s["loc"]["bbox"][1] == 86:
            s["loc"]["role"] = "header"
            if s["page"] in (2, 4):
                s["text"] = s["text"].replace("üzemzavarai", "üzemzavarai.").replace("Szekunderköri", "Szekunderkori")
    blocks, info = DL.load(src, segs)
    text = lambda b: "".join(r.get("t", "") for r in b.get("runs") or [])
    titles = [text(b) for b in blocks if b["k"] in ("h", "p") and ("üzemzavar" in text(b) or "Szekunderk" in text(b)
                                                                    or text(b).startswith("3. Egyéb"))]
    assert len(titles) == 3, titles


def test_scan_dash_list_geometry_from_layer(tmp_path):
    """Скан: тире пункта — чёрточка текстового слоя, текст — начало куска OCR: место маркера и текста по оригиналу
    (по профилю текст стоял на 60 pt правее тире, 3PR42 +23 листа)."""
    import docx_layout as DL
    src, segs = ocr_doc(tmp_path)
    blocks, info = DL.load(src, segs)
    geo = next(b["geo"] for b in blocks if b["k"] == "list" and b.get("geo"))
    assert 95 <= geo["marker"] <= 101 and abs(geo["text"] - 112) < 1


def test_list_marker_hanging_into_margin_keeps_original_place(tmp_path):
    """Тире пункта левее поля набора (скан: поле — по началу текста, тире на 4,75 pt левее) — пункт у поля с местом
    текста по оригиналу, а не с отступами профиля (текст уезжал на 64 pt)."""
    item = p("Ha az RTSZ aut. üzemmódban volt", 1, 300, x0=81.5)
    blocks = [{"k": "list", "kind": "Dash", "items": [[item]], "geo": {"marker": 57.7, "text": 81.5}}]
    out = tmp_path / "l.docx"
    B.Builder(B.Memory([]), lambda img: None, PROFILE).build(blocks, str(out), LAYOUT)
    par = next(x for x in Document(str(out)).paragraphs if x.text.startswith("Ha az"))
    assert abs(par.paragraph_format.left_indent.pt - (81.5 - 62.4)) < 0.6
    assert abs(par.paragraph_format.first_line_indent.pt + (81.5 - 62.4)) < 0.6


def test_page_split_does_not_join_next_list_label():
    """«ca) … maradnak,» внизу листа и «cb) …» вверху следующего — разные пункты (OAH N3a), хоть «cb)» и со строчной;
    продолжение фразы со строчной — склеивается."""
    def p(t, page):
        return {"k": "p", "runs": [{"t": t}], "geom": {"page": page, "top": 700 if page == 44 else 85, "x0": 99,
                                                        "first_x0": 99, "x1": 524, "last_x1": 300, "lines": 2}}
    got = D.join_pages([p("ca) a felaktiválódásra hajlamosak, maradnak,", 44), p("cb) a sugárzás hatására,", 45),
                        p("d) a neutronsugárzásnak kitett", 45)])
    assert len(got) == 3
    got = D.join_pages([p("a felaktiválódásra hajlamosak", 44), p("szerkezetük pedig olyan", 45)])
    assert len(got) == 1


def test_scan_toc_lines_not_list_and_running_title_heading(tmp_path):
    """Скан: строки оглавления («1.3/1. SZBV kazetta leesése ⇥ 8») по метке портала «пункт» — не список; бегущее
    название раздела в рамке без метки портала (повтор на следующем листе) — заголовок на первом листе, один раз."""
    import docx_layout as DL
    src, segs = ocr_doc(tmp_path)
    toc = [s for s in segs if s["text"].startswith(("1.4.", "Tengelyelmozdulás", "13"))]
    for s in toc:
        s["loc"]["role"] = "list"
    for s in segs:
        if s["loc"]["bbox"][1] == 86:
            s["loc"]["role"] = None
            s["text"] = s["text"].replace("  ", " ")
    blocks, info = DL.load(src, segs)
    text = lambda b: "".join(r.get("t", "") for r in b.get("runs") or [])
    assert not any(b["k"] == "list" and any("Tengely" in text(x) for it in b["items"] for x in it) for b in blocks)
    heads = [text(b) for b in blocks if b["k"] == "h"]
    assert sum("üzemzavarai" in h for h in heads) == 1 and sum("Szekunderköri" in h for h in heads) == 1


def test_line_under_header_frame_is_body_not_running_title(tmp_path):
    """Скан с рамкой шапки (3SZ19: линии 35/68/82/110, название подраздела в рамке): строка под рамкой, что стоит на
    том же месте и на прошлом листе («1. Ismérvek:» — раздел с нового листа), — тело, а не бегущее название: снималась
    на 39 листах из 68; и строка под рамкой, повторённая ниже на том же листе («Turbinát kiütni.» в двух ветках), —
    тоже; заголовок под рамкой с тем же текстом, что прошлый заголовок того же уровня (метка портала), — не повтор.
    Бегущее название в рамке на листе продолжения раздела по-прежнему снимается."""
    import docx_layout as DL
    doc = fitz.open()
    segs = []

    def seg(pno, text, x0, y0, x1, align="left", role=None):
        segs.append({"id": len(segs), "page": pno, "text": text,
                     "loc": {"bbox": [x0, y0, x1, y0 + 7], "frags": [[x0, y0, x1, y0 + 7]], "size": 9.7, "bold": False,
                             "align": align, "family": "monospace", "ocr": True, "roles": True, "role": role}})
    subs = ["1.1. Turbina leállítása vákuumrontással"] * 2 + ["1.2. Generátor teherledobás", "1.3. Vízütés a turbinánál"] \
        + ["1.4. Csapágy hőmérséklet emelkedés"] * 2
    for pno, sub in enumerate(subs):
        page = doc.new_page(width=595, height=842)
        sh = page.new_shape()
        for y in (35, 68, 82, 110):
            sh.draw_line((70, y), (559, y))
        sh.finish(color=(0, 0, 0), width=0.6)
        sh.commit()
        seg(pno, "Azonosító: 3SZ19", 77, 38, 172, role="header")
        seg(pno, f"Oldalszám: {pno + 2}", 461, 57, 538, role="header")
        seg(pno, "1. A turbina üzemzavarai", 241, 88, 390, "center", "header")
        seg(pno, sub, 160, 98, 470, "center", "heading")        # название подраздела в рамке
        if pno == 0 or sub != subs[pno - 1]:                    # лист раздела: под рамкой — «1. Ismérvek:»
            seg(pno, "1. Ismérvek:", 89, 122, 171, role="heading")
            seg(pno, "Turbinát kiütni.", 157, 145, 241)
            seg(pno, f"Védelmi lefutást ellenőrizni, elmaradt működéseket pótolni {pno}.", 157, 200, 473)
            seg(pno, "Turbinát kiütni.", 157, 250, 241)
        else:
            seg(pno, f"Tengelyelmozdulás növekedése SA10V125, lap {pno}.", 157, 140, 397)
    doc.save(tmp_path / "frame.pdf")
    blocks, _ = DL.load(str(tmp_path / "frame.pdf"), segs)
    texts = [(b["geom"]["page"], B.text_of(b["runs"]).strip()) for b in B._paras(blocks)]
    assert sorted(pg for pg, t in texts if t == "1. Ismérvek:") == [1, 3, 4, 5]
    assert sum(t == "Turbinát kiütni." for _, t in texts) == 8
    assert sum("Turbina leállítása" in t for _, t in texts) == 1 and sum("Csapágy hőmérséklet" in t for _, t in texts) == 1


def test_scan_scheme_of_arrows_and_short_labels_is_figure(tmp_path):
    """Распознанный лист: схема из стрелок с короткими подписями (3PR42 л. 113, кто где при расхолаживании) — рисунком
    с переведённого PDF, подписи — не абзацами (рассыпалась на «Turbi Reaktor operátor», «na operátor»); стрелки
    разнесены дальше 20 pt и соединяются через подписи; подпись, слитая OCR из трёх («Turbina főgépész Primerköri
    főgépész Villamos művezető»), схеме не мешает. Список условий с тире и метками в рамочках «[3/2]» (3SZ19 л. 4, без
    вертикалей) — текстом."""
    import docx_layout as DL
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    sh = page.new_shape()

    def arrow(x, y0, y1):
        sh.draw_line((x, y0), (x, y1))
        sh.finish(color=(0, 0, 0), width=0.7)
        sh.draw_polyline([(x - 3, y1 - 1), (x + 3, y1 - 1), (x, y1 + 5)])
        sh.finish(color=(0, 0, 0), fill=(0, 0, 0), closePath=True)
    for a, b in (((125, 318), (282, 318)), ((228, 332), (285, 332))):
        sh.draw_line(a, b)
        sh.finish(color=(0, 0, 0), width=0.7)
    for x, y0, y1 in ((125, 318, 384), (228, 332, 383), (330, 354, 385), (459, 354, 385), (229, 426, 457),
                      (330, 426, 457), (459, 426, 457)):
        arrow(x, y0, y1)
    for y in (600, 630, 660):                                       # условия: тире и метка «3/2» в рамочке
        sh.draw_line((92, y + 5), (110, y + 5))
        sh.finish(color=(0, 0, 0), width=0.7)
        sh.draw_rect(fitz.Rect(110, y, 135, y + 10))
        sh.finish(color=(0, 0, 0), width=0.7)
    sh.commit()
    segs = []

    def seg(text, x0, y0, x1):
        segs.append({"id": len(segs), "page": 0, "text": text,
                     "loc": {"bbox": [x0, y0, x1, y0 + 10], "frags": [[x0, y0, x1, y0 + 10]], "size": 9.7, "bold": False,
                             "align": "left", "family": "serif", "ocr": True}})
    seg("Lehűtés során az operatív személyzet tartózkodási helye:", 100, 160, 372)
    for y, t in ((302, "Blokkügyeletes"), (314, "IÜO művezető"), (329, "Turbina operátor"), (341, "Reaktor operátor")):
        seg(t, 291, y, 369)
    seg("Vezető elektrikus", 414, 342, 495)
    for x, t in ((92, "TARTÓZKODÓ"), (218, "KÜV"), (294, "TARTÓZKODÓ"), (414, "TARTÓZKODÓ")):
        seg(t, x, 394, x + 80)
    seg("Műszerészek", 100, 409, 185)
    seg("Turbina főgépész Primerköri főgépész Villamos művezető", 195, 409, 510)
    for x, t in ((195, "TARTÓZKODÓ"), (306, "HELYSZÍNEN"), (414, "HELYSZÍNEN")):
        seg(t, x, 461, x + 80)
    for x, t in ((209, "Gépészek"), (306, "Gépészek"), (426, "Elektrikusok")):
        seg(t, x, 476, x + 60)
    for y, t in ((600, "Generátor tengelyzár hőmérséklet-növekedése t > 110 °C."), (630, "Generátor 8. csapágy"),
                 (660, "Tengelyzár olajrendszer nyomása csökken, Δp < 0,2 bar.")):
        seg("3/2", 113, y, 132)
        seg(t, 140, y, 140 + 5 * len(t))
    doc.save(tmp_path / "scheme.pdf")
    blocks, _ = DL.load(str(tmp_path / "scheme.pdf"), segs)
    texts = [B.text_of(b["runs"]) for b in B._paras(blocks)]
    figs = [b for b in blocks if b.get("figure")]
    assert len(figs) == 1 and figs[0]["runs"][0]["img"]["clip"][1] < 330 < 470 < figs[0]["runs"][0]["img"]["clip"][3]
    assert not any(t in x for x in texts for t in ("Reaktor operátor", "Primerköri főgépész", "Gépészek", "KÜV"))
    assert any("Lehűtés során" in x for x in texts) and any("Generátor 8. csapágy" in x for x in texts)
    assert sum("3/2" in x for x in texts) == 3


def test_label_column_splits_table_row_into_sub_rows():
    """Таблица-бланк без внутренних горизонталей (31RDGT): номер, текст и отметка — каждый против своего; строка сетки
    режется по верху меток первого столбца, линий между подстроками нет."""
    import docx_layout as DL
    gh = [(90, 60, 560), (300, 60, 560)]
    gv = [(60, 90, 300), (110, 90, 300), (430, 90, 300), (560, 90, 300)]
    t = (DL._cluster([v[0] for v in gv]), [90, 300], gh, gv)

    def seg(text, x0, y0, x1):
        return {"id": 0, "page": 0, "text": text, "loc": {"bbox": [x0, y0, x1, y0 + 12], "frags": [[x0, y0, x1, y0 + 12]],
                                                          "size": 10, "bold": False, "align": "left"}}
    segs = [seg("1.1.", 70, 100, 90), seg("Az alábbi armatúrák", 120, 100, 400), seg("31RA31S401", 200, 114, 260),
            seg("O", 490, 114, 498), seg("1.2.", 70, 160, 90), seg("A 31RD31P702 kör", 120, 160, 400),
            seg("O", 490, 160, 498)]
    b = DL._table_block(t, segs, 0, None, "Times New Roman", 10, False)
    text = lambda c: [ "".join(r.get("t", "") for r in p["runs"]) for p in c["blocks"]]
    assert len(b["rows"]) == 2
    assert [text(c) for c in b["rows"][0]] == [["1.1."], ["Az alábbi armatúrák", "31RA31S401"], ["O"]]
    assert [text(c) for c in b["rows"][1]] == [["1.2."], ["A 31RD31P702 kör"], ["O"]]
    assert not b["rows"][0][1]["edges"]["bottom"] and not b["rows"][1][1]["edges"]["top"]
    assert len(b["frame"]["ys"]) == 3


def test_case_suffix_and_dotted_abbreviation_are_not_codes():
    """«4.1.9.-ben» — номер пункта с падежным суффиксом, «Oszt.vez.» — сокращение (не имя файла): шлифовка не должна
    возвращать их в верный перевод; настоящее имя файла в конце предложения — код по-прежнему (30RR2126GT, 28.10)."""
    assert C._codes("Ha a 4.1.10. és 4.1.9.-ben rögzített nyomásértékek") == []
    assert C._codes("Oszt.vez. Oszt.vez.") == [] and C._codes("Oszt.vez aláírása") == []
    assert "config.ini" in C._codes("lásd a config.ini.")
    ru = "ЕСЛИ разность значений в 4.1.10. и 4.1.9., <5 бар"
    assert C.fix_code("Ha a 4.1.10. és 4.1.9.-ben rögzített", ru, RULES)[0] == ru
    assert not [v for v in C.check("Oszt.vez.", "Нач. отдела", RULES) if v[0] == "Я6"]


def test_hungarian_leftover_ignores_document_codes_and_first_use_abbreviation():
    """«TLÜ302» (обозначение документа) и «НСБ (BLÜGY)» (первое употребление сокращения) — не остаток венгерского в
    готовом DOCX; настоящее венгерское слово по-прежнему замечается (30RR2126GT, 28.10)."""
    ok = lambda t: not [v for v in C.check_text(t, RULES) if v[0] == "Я11"]
    assert ok("в соответствии с TLÜ302 провел совещание") and ok("Начало теста разрешается НСБ (BLÜGY).")
    assert not ok("ЗАКРЫТЬ szelepeket и ellenőrizni")


def test_short_capital_abbreviations_are_not_hungarian_leftovers():
    """ÜBO, ÍV, ÜIO — сокращения подразделений (латиницей, как есть по глоссарию): не остаток венгерского, шлифовка не
    выдумывает «УБО»/«ИВ» (30BL150FÜT: 30 строк); слово из пяти и более прописных букв — по-прежнему замечается."""
    assert C._hungarian("Подразделение: ÜBO, ÍV, ÜIO") == [] and C._hungarian("ЗАКРЫТО ZÁRVA") == ["ZÁRVA"]
    assert C._hungarian("с идентификатором 30BLHÜT, TLÜ302") == []                           # обозначения с цифрами

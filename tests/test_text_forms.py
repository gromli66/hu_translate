# Бланки текстовым слоем (формы RDGT, 20.10): рамки четырёхугольниками, шапка таблицы на каждом листе, таблица через
# лист, отметки кривыми, подложки полей штампа, нумерация листов при таблице в начале раздела.
import fitz
from docx import Document

import docx_build as B
import docx_layout as DL
import docx_struct as D

from test_struct_docx import LAYOUT, PROFILE, W, p, png


def test_quad_frame_is_four_rules():
    """Рамка путём из четырёх отрезков («qu») — линии, как прямоугольник «re»: без неё таблица бланка не находилась."""
    q = ((42, 81), (574, 81), (42, 769), (574, 769))
    hs, vs = DL._rules([{"rect": (42, 81, 574, 769), "type": "s", "color": (0, 0, 0), "items": [("qu", q)]}])
    assert sorted(round(h[0]) for h in hs) == [81, 769] and sorted(round(v[0]) for v in vs) == [42, 574]


def test_vector_marks_of_form():
    """Кружок из дуг и квадратик обводкой (6–20 pt) — отметки бланка; залитый кружок и крупная рамка — нет."""
    circle = {"rect": (486, 199, 502, 215), "type": "s", "color": (0, 0, 0), "items": [("c",)] * 5}
    square = {"rect": (486, 300, 498, 312), "type": "s", "color": (0, 0, 0), "items": [("re", (486, 300, 498, 312))]}
    dot = {"rect": (486, 400, 494, 408), "type": "f", "fill": (0, 0, 0), "items": [("c",)] * 4}
    box = {"rect": (90, 300, 200, 314), "type": "s", "color": (0, 0, 0), "items": [("re", (90, 300, 200, 314))]}
    assert [ch for _, ch in DL._form_marks([circle, square, dot, box])] == ["○", "□"]


def _form_pdf(path, rows=None):
    """Два листа бланка: штамп в рамке-«qu» с белыми подложками под словами, таблица (рамка-«qu», две вертикали,
    линия под шапкой), шапка таблицы на обоих листах, строки с кружками отметок, подвал."""
    doc = fitz.open()
    rows = rows or [[("1.1.", "Az elso feltétel teljesül."), ("1.2.", "A második armatúra zárva.")],
                    [("1.3.", "A harmadik kapcsoló élesítve."), ("1.4.", "A negyedik próba lezárva.")]]
    for pno in range(2):
        pg = doc.new_page(width=595, height=842)
        sh = pg.new_shape()
        for r in ((54, 42, 102, 55), (137, 42, 180, 55), (405, 42, 459, 55), (54, 55, 112, 68), (501, 55, 552, 68)):
            sh.draw_rect(fitz.Rect(r))
        sh.finish(color=None, fill=(1, 1, 1))
        sh.draw_quad(fitz.Rect(42, 21, 574, 69).quad)
        sh.draw_quad(fitz.Rect(42, 81, 574, 769).quad)
        sh.draw_line((90, 81), (90, 769))
        sh.draw_line((469, 81), (469, 769))
        sh.draw_line((42, 117), (574, 117))
        sh.finish(color=(0, 0, 0), width=0.8)
        for k in range(len(rows[pno])):
            sh.draw_circle((494, 152 + 50 * k), 8)
            sh.finish(color=(0, 0, 0), width=0.8)
        sh.commit()
        put = lambda x, y, t: pg.insert_text((x, y), t, fontsize=9, fontname="tiro")
        put(250, 37, "TESZTELÉSI UTASÍTÁS")
        for x, y, t in ((54, 52, "Azonosító:"), (137, 52, "30RDGT"), (405, 52, "Verziószám:"), (547, 52, "3"),
                        (54, 65, "Megnevezés:"), (137, 65, "3-as blokk"), (405, 65, "Dátum:"), (501, 65, "2017-04-20")):
            put(x, y, t)
        put(48, 96, "Sorsz.")
        put(214, 96, "Feltétel, muvelet, ellenorzés")
        put(478, 96, "Teljesítés igazolás")
        for k, (num, text) in enumerate(rows[pno]):
            put(48, 155 + 50 * k, num)
            put(94, 155 + 50 * k, text)
        put(46, 812, "Nyomtatva: 2017-04-28 11:48")
    doc.save(str(path))
    segs = []
    for pno, pg in enumerate(fitz.open(str(path))):
        for bl in pg.get_text("dict")["blocks"]:
            for ln in bl.get("lines", []):
                t = "".join(s["text"] for s in ln["spans"]).strip()
                if t:
                    bb = [round(v, 1) for v in ln["bbox"]]
                    segs.append({"id": len(segs), "page": pno, "text": t, "src": t,
                                 "loc": {"bbox": bb, "frags": [bb], "size": 9.0, "bold": False, "family": "Times",
                                         "align": "left"}})
    return segs


def test_form_table_header_is_not_running_and_table_joins_across_pages(tmp_path):
    """Шапка таблицы на каждом листе — строка таблицы, не колонтитул (30RDGT: уходила в шапку DOCX, 15 листов → 68);
    таблица через лист — одна, шапка повторяется (w:tblHeader), хвост строки до рамки листа не тянется; кружки
    кривыми — знаком «○»; штамп в колонтитуле — не по краям белых подложек полей."""
    src = tmp_path / "form.pdf"
    segs = _form_pdf(src)
    blocks, info = DL.load(str(src), segs, text=True)
    tabs = [b for b in blocks if b["k"] == "table"]
    assert len(tabs) == 1
    t = tabs[0]
    cell = lambda c: " ".join("".join(r.get("t", "") for r in x["runs"]) for x in c["blocks"]).strip()
    texts = [[cell(c) for c in r] for r in t["rows"]]
    assert texts[0][:2] == ["Sorsz.", "Feltétel, muvelet, ellenorzés"]
    assert sum("Feltétel" in c for r in texts for c in r) == 1                  # повтор шапки — прочь
    assert t.get("repeat_rows") == 1 and t.get("head_repeat")
    body = " ".join(c for r in texts for c in r)
    assert all(n in body for n in ("1.1.", "1.2.", "1.3.", "1.4.")) and body.count("○") == 4
    ys = t["frame"]["ys"]
    assert len(ys) == len(t["rows"]) + 1
    last1 = max(i for i, r in enumerate(t["rows"]) if any((c.get("geom") or {}).get("page") == 1 for c in r))
    assert ys[last1 + 1] - ys[last1] < 80                                       # не до рамки листа (769)
    L = D.layout(str(src), blocks, content=info.get("content_boxes"))
    head = L.get("head_grid") or {}
    assert len(head.get("widths") or []) <= 5
    assert "Feltétel" not in str(head) and "Azonosító" in str(head)
    out = tmp_path / "form.docx"
    mem = B.Memory([(s["text"], s["text"]) for s in segs])
    B.Builder(mem, lambda img: png(), PROFILE).build(blocks, str(out), dict(LAYOUT, **{k: v for k, v in L.items()
                                                                                      if k in ("head_grid",)}))
    tbl = Document(str(out)).tables[0]
    assert tbl.rows[0]._tr.find(".//" + W + "tblHeader") is not None


def test_page_numbers_restart_before_leading_table(tmp_path):
    """Раздел с нумерацией с числа оригинала, начатый таблицей: LibreOffice такое начало теряет (31RDGT: «1» вместо
    «2»), перед таблицей — пустой абзац в 1 pt."""
    g = lambda top: {"page": 1, "top": top, "bottom": top + 10, "x0": 80, "x1": 300, "first_x0": 80, "size": 10,
                     "lines": 1}
    cellp = lambda t, top: {"blocks": [{"k": "p", "runs": [{"t": t, "font": "Tahoma", "size": 10, "b": False,
                                                            "i": False, "color": "000000"}], "geom": g(top)}],
                            "th": False, "geom": g(top)}
    tab = {"k": "table", "rows": [[cellp("1.1.", 100), cellp("Első", 100)], [cellp("1.2.", 120), cellp("Második", 120)]],
           "borders": True, "geom": dict(g(100), x1=500)}
    out = tmp_path / "n.docx"
    B.Builder(B.Memory([("Első", "Первый"), ("Második", "Второй")]), lambda img: png(), PROFILE).build(
        [tab], str(out), dict(LAYOUT, page_tok="2", num_start={"page": 1, "value": 2}, first_is_body=True))
    body = Document(str(out)).element.body
    first = next(el for el in body if el.tag in (W + "p", W + "tbl"))
    assert first.tag == W + "p"
    assert Document(str(out)).sections[0]._sectPr.find(W + "pgNumType").get(W + "start") == "2"


def test_scan_page_inside_text_layer_goes_as_translated_snapshot(tmp_path):
    """Скан посреди текстового слоя (лист регистрации 32RDGT) — снимком листа с переведённого PDF (img["clip"]),
    а не россыпью сегментов; его полосы-картинки — не логотип колонтитула."""
    import numpy as np
    src = tmp_path / "mixed.pdf"
    doc = fitz.open()
    for pno in range(3):
        pg = doc.new_page(width=595, height=842)
        if pno == 1:                                    # скан: две полосы-картинки на весь лист
            a = np.full((421, 595, 3), 255, np.uint8)
            a[100:102, 50:545] = 0
            for k, y in enumerate((0, 421)):
                pix = fitz.Pixmap(fitz.csRGB, 595, 421, a.tobytes(), False)
                pg.insert_image(fitz.Rect(0, y, 595, y + 421), pixmap=pix)
        else:
            pg.insert_text((72, 100 + 20 * pno), f"Szöveg a {pno + 1}. lapon, rendes bekezdés.", fontsize=10)
    doc.save(str(src))
    segs = []
    for pno, pg in enumerate(fitz.open(str(src))):
        lines = [(ln["bbox"], "".join(s["text"] for s in ln["spans"])) for bl in pg.get_text("dict")["blocks"]
                 for ln in bl.get("lines", [])]
        if pno == 1:                                    # распознанные строки скана
            lines = [((60, 80, 300, 92), "NYILVÁNTARTÓ LAP"), ((60, 120, 400, 132), "1./ A tesztelési utasítás")]
        for bb, t in lines:
            bb = [round(v, 1) for v in bb]
            segs.append({"id": len(segs), "page": pno, "text": t, "src": t,
                         "loc": {"bbox": bb, "frags": [bb], "size": 10.0, "bold": False, "family": "Times",
                                 "align": "left", **({"ocr": True} if pno == 1 else {})}})
    blocks, info = DL.load(str(src), segs, text=True)
    snaps = [b for b in blocks if b["k"] == "p" and b["runs"] and "img" in b["runs"][0]]
    assert len(snaps) == 1 and snaps[0]["runs"][0]["img"]["page_number"] == 2 and snaps[0]["runs"][0]["img"]["clip"]
    text = " ".join("".join(r.get("t", "") for r in b.get("runs") or []) for b in blocks)
    assert "NYILVÁNTARTÓ" not in text and "1. lapon" in text
    L = D.layout(str(src), blocks, content=info.get("content_boxes"))
    assert not any((z.get("head_grid") or {}).get("paras") and any("img" in r for r in z["head_grid"]["paras"]["rows"])
                   for z in (L, L.get("front") or {}, L.get("first") or {}))


def test_first_page_with_table_running_on_is_not_cover(tmp_path):
    """Первый лист бланка с одной строкой таблицы, которая идёт дальше, — не титул (30RDIT: таблица уходила в раздел
    титула, у пяти листов не было шапки)."""
    src = tmp_path / "f1.pdf"
    many = [(f"2.{k}.", f"A {k}. feltétel teljesül, a kapcsoló élesítve van.") for k in range(1, 11)]
    segs = _form_pdf(src, rows=[[("1.1.", "Egy sor.")], many])
    blocks, info = DL.load(str(src), segs, text=True)
    assert [b["k"] for b in blocks] == ["table"] and blocks[0].get("split")
    L = D.layout(str(src), blocks, content=info.get("content_boxes"))
    assert L["first_is_body"]


def test_form_column_alignment_votes_on_all_pages():
    """Столбцы таблицы бланка на всех листах те же: голосуют все листы — один заголовок по центру на первом листе
    не ставит по центру весь столбец (30RDIT)."""
    xs = [40, 90, 420, 530]
    cell = lambda page, x0, x1, th=False: {"blocks": [], "th": th, "geom": {"page": page, "x0": x0, "x1": x1,
                                                                             "lines": 1}}
    head = [cell(1, 50, 80, True), cell(1, 200, 310, True), cell(1, 440, 510, True)]
    rows = [head, [cell(1, 50, 60), cell(1, 200, 310), cell(1, 470, 480)]]
    rows += [[cell(2, 50, 60), cell(2, 94, 300 + 10 * k), cell(2, 470, 480)] for k in range(4)]
    assert B._col_align(rows, xs, 1).get(("body", 1)) == "center"                # голос только первого листа
    assert B._col_align(rows, xs, 1, any_page=True).get(("body", 1)) is None       # влево — по умолчанию


def test_private_use_mark_by_its_look():
    """Значок частной области (Wingdings, вшитый в шрифт текста) — по виду на листе: пустой квадрат, залитый кружок."""
    doc = fitz.open()
    pg = doc.new_page(width=200, height=100)
    sh = pg.new_shape()
    sh.draw_rect(fitz.Rect(20, 20, 32, 32))
    sh.finish(color=(0, 0, 0), width=1)
    sh.draw_circle((80, 26), 6)
    sh.finish(color=(0, 0, 0), fill=(0, 0, 0))
    sh.commit()
    assert DL._pua_glyph(pg, (16, 16, 36, 36)) == "□"
    assert DL._pua_glyph(pg, (70, 16, 90, 36)) == "●"
    assert DL._pua_glyph(pg, (120, 16, 140, 36)) is None                         # пусто


def test_narrow_columns_without_room_shrink_by_style(tmp_path):
    """Много узких столбцов, расширить нечем: слово перевода не влезает — кегль ячейки меньше производным стилем
    (не прямым кеглем), а не слово по буквам (таблица параметров 30RDIT)."""
    xs = [40 + 24 * i for i in range(9)]
    g = lambda x0: {"page": 1, "top": 100, "bottom": 110, "x0": x0, "x1": x0 + 20, "first_x0": x0, "size": 10,
                    "lines": 1}
    cell = lambda t, x0: {"blocks": [{"k": "p", "runs": [{"t": t, "font": "Tahoma", "size": 10, "b": False, "i": False,
                                                          "color": "000000"}], "geom": g(x0)}], "th": False,
                          "geom": g(x0), "span": 1, "edges": {"top": True, "bottom": True, "left": True, "right": True}}
    rows = [[cell("ÉRTÉK", x) for x in xs[:-1]] for _ in range(4)]
    tab = {"k": "table", "rows": rows, "borders": True, "frame": {"xs": xs, "ys": [100, 112, 124, 136, 148],
                                                                   "row_lines": True},
           "geom": {"page": 1, "top": 100, "bottom": 148, "x0": xs[0], "x1": xs[-1], "first_x0": xs[0], "size": 10}}
    out = tmp_path / "s.docx"
    B.Builder(B.Memory([("ÉRTÉK", "ЗНАЧЕНИЕ")]), lambda img: png(), PROFILE).build(
        [tab], str(out), dict(LAYOUT, first_is_body=True))
    names = {p.style.name for c in Document(str(out)).tables[0].rows[1].cells for p in c.paragraphs}
    assert any(n.endswith("pt") for n in names)
    assert not Document(str(out)).tables[0]._tbl.xpath(".//w:r/w:rPr/w:sz")


def test_short_values_against_labels_are_not_one_paragraph():
    """Столбик коротких значений против подписей слева («BSZG állásjelzések … rendben» ×4, 31RDGT) — по строке;
    длинные строки второй колонки сплошного текста — как прежде, абзацем."""
    import layout as LY
    fr = lambda x0, y0, x1, t: {"bbox": (x0, y0, x1, y0 + 12), "t": t, "size": 10, "bold": False}
    labels = ["BSZG állásjelzések", "BSZG online eseménylista", "BV állásjelzések", "Helyszín"]
    frags = [fr(156, 288 + 13 * k, 280, t) for k, t in enumerate(labels)] + \
        [fr(375, 288 + 13 * k, 414, "rendben") for k in range(4)]
    pars = LY.group(frags, [], [])
    assert sum(p["text"] == "rendben" for p in pars) == 4
    prose = [fr(60, 100 + 13 * k, 290, "bal hasáb szövege, hosszú sor sok szóval itt") for k in range(3)] + \
        [fr(310, 100 + 13 * k, 540, "jobb hasáb folyó szövege, hosszú sor sok szóval") for k in range(3)]
    assert len(LY.group(prose, [], [])) == 2


def test_title_scan_without_text_is_the_title_not_background(tmp_path):
    """Титул — один скан с подписями (ИН-1065): это сам титул, картинкой; фон обложки под текстом (UTS) — как прежде,
    не картинкой в потоке."""
    big = {"k": "p", "runs": [{"img": {"x0": 57, "x1": 567, "top": 92, "bottom": 812}}],
           "geom": {"page": 1, "top": 92, "bottom": 812, "x0": 57, "x1": 567, "first_x0": 57}}
    body = p("Bevezetés", 2, 100)
    mem = B.Memory([("Bevezetés", "Введение"), ("Címlap", "Титул")])
    count = lambda f: len(Document(str(f)).element.body.xpath(".//w:drawing"))
    out1 = tmp_path / "scan.docx"
    B.Builder(mem, lambda img: png(), PROFILE).build([big, body], str(out1), dict(LAYOUT))
    out2 = tmp_path / "bg.docx"
    B.Builder(mem, lambda img: png(), PROFILE).build([big, p("Címlap", 1, 300), body], str(out2), dict(LAYOUT))
    assert count(out1) == 1 and count(out2) == 0
    # логотип шапки — векторный (xref нет): снимок скана титула (тоже без xref) — не логотип (ред-тим 20.10: лист 1 пустой)
    out3 = tmp_path / "vec.docx"
    logo = {"xref": None, "x0": 60, "x1": 120, "top": 30, "bottom": 60}
    B.Builder(mem, lambda img: png(), PROFILE).build([big, body], str(out3), dict(LAYOUT, head_imgs=[logo]))
    assert count(out3) == 1


def test_soft_hyphens_by_syllables():
    """Места переноса русского слова (правила П. Христова) и мягкие переносы прогона — элементом w:softHyphen."""
    split = lambda w: "-".join(w[a:b] for a, b in zip([0] + B._hyphen_points(w), B._hyphen_points(w) + [len(w)]))
    assert split("ЖУРНАЛ") == "ЖУР-НАЛ" and split("ЗНАЧЕНИЕ") == "ЗНА-ЧЕ-НИЕ" and split("СОБЫТИЙ") == "СО-БЫ-ТИЙ"
    assert split("установки") == "ус-та-нов-ки" and split("мА") == "мА"
    run = Document().add_paragraph().add_run("ЖУРНАЛ СОБЫТИЙ")
    B._soft_hyphens(run)
    assert len(run._r.findall(W + "softHyphen")) == 3 and run.text.replace("-", "") == "ЖУРНАЛ СОБЫТИЙ"


def test_form_fields_become_bordered_runs(tmp_path):
    """Поля ввода бланка (обведённые прямоугольники в ячейке): код в поле — в рамке (w:bdr), пустое поле — неразрывные
    пробелы в рамке и невидимый соединитель после (рамку пробелов в конце строки LibreOffice не рисует), 30RDGT."""
    q = lambda x0, y0, x1, y1: {"rect": (x0, y0, x1, y1), "type": "s", "color": (0, 0, 0),
                                "items": [("qu", ((x0, y0), (x1, y0), (x0, y1), (x1, y1)))]}
    assert DL._form_fields([q(92, 266, 206, 282), q(209, 266, 298, 282), q(42, 81, 574, 769)]) == \
        [(92, 266, 206, 282), (209, 266, 298, 282)]
    gh = [(81, 42, 574), (117, 42, 574), (769, 42, 574)]
    gv = [(42, 81, 769), (90, 81, 769), (469, 81, 769), (574, 81, 769)]
    t = ([42, 90, 469, 574], [81, 117, 769], gh, gv)
    seg = {"id": 7, "page": 0, "text": "31RD21S201", "loc": {"bbox": [95, 267, 155, 280], "frags": [[95, 267, 155, 280]],
                                                            "size": 9, "bold": False}}
    b = DL._table_block(t, [seg], 0, None, "Times New Roman", 9, False, fields=[(92, 266, 206, 282), (209, 266, 298, 282)])
    runs = [r for row in b["rows"] for c in row for p_ in c["blocks"] for r in p_["runs"]]
    assert [bool(r.get("box")) for r in runs if r.get("t", "").strip(" \t")] == [True, True]
    out = tmp_path / "f.docx"
    B.Builder(B.Memory([("31RD21S201", "31RD21S201")]), lambda img: png(), PROFILE).build(
        [b], str(out), dict(LAYOUT, first_is_body=True))
    body = Document(str(out)).element.body
    assert len(body.xpath(".//w:r[w:rPr/w:bdr]")) == 2 and "⁠" in "".join(body.itertext())


def test_soft_hyphens_keep_tabs_and_breaks():
    """Мягкие переносы не трогают табуляцию и разрыв строки прогона («ВНИМАНИЕ!↵На» 3PR42 слипалось)."""
    run = Document().add_paragraph().add_run("ВНИМАНИЕ!\nНаблюдение\tЖУРНАЛ")
    B._soft_hyphens(run)
    assert len(run._r.findall(W + "br")) == 1 and len(run._r.findall(W + "tab")) == 1
    assert run.text.replace("-", "") == "ВНИМАНИЕ!\nНаблюдение\tЖУРНАЛ"


# ---- красная команда 20.10 (docs/rules_contract/UNIVERSALITY.md): ложные срабатывания правил N3, N4, N8, N16, N17

def _segs(path):
    segs = []
    for pno, pg in enumerate(fitz.open(str(path))):
        for bl in pg.get_text("dict")["blocks"]:
            for ln in bl.get("lines", []):
                t = "".join(s["text"] for s in ln["spans"]).strip()
                if t:
                    bb = [round(v, 1) for v in ln["bbox"]]
                    segs.append({"id": len(segs), "page": pno, "text": t, "src": t,
                                 "loc": {"bbox": bb, "frags": [bb], "size": 10.0, "bold": False, "family": "Times",
                                         "align": "left"}})
    return segs


def _grid(sh, xs, ys):
    for x in xs:
        sh.draw_line((x, ys[0]), (x, ys[-1]))
    for y in ys:
        sh.draw_line((xs[0], y), (xs[-1], y))
    sh.finish(color=(0, 0, 0), width=0.6)


def test_cells_drawn_one_by_one_are_not_form_marks_and_ticked_box_is_marked(tmp_path):
    """N4: таблица из клеток 16×16 pt, каждая — своим прямоугольником (график по дням): пустые клетки — не «□»;
    квадратик с галочкой внутри — «☒», а не пустой «□»."""
    src = tmp_path / "g.pdf"
    doc = fitz.open()
    pg = doc.new_page(width=595, height=842)
    sh = pg.new_shape()
    for r in range(4):
        cells = [fitz.Rect(60, 100 + 16 * r, 180, 116 + 16 * r)]
        cells += [fitz.Rect(180 + 16 * m, 100 + 16 * r, 196 + 16 * m, 116 + 16 * r) for m in range(10)]
        for c in cells:
            sh.draw_rect(c)
            sh.finish(color=(0, 0, 0), width=0.5)
    sh.commit()
    for r in range(4):
        pg.insert_text((63, 110 + 16 * r), f"Szivattyú {r + 1}", fontsize=8, fontname="tiro")
    doc.save(str(src))
    blocks, _ = DL.load(str(src), _segs(src), text=True)
    text = " ".join("".join(q.get("t", "") for q in x["runs"]) for b in blocks if b["k"] == "table"
                    for r in b["rows"] for c in r for x in c["blocks"])
    assert "Szivattyú 4" in text and "□" not in text
    box = {"rect": (410, 138, 422, 150), "type": "s", "color": (0, 0, 0), "items": [("re", (410, 138, 422, 150))]}
    tick = {"rect": (412, 139, 421, 148), "type": "s", "color": (0, 0, 0), "items": [("l", (412, 144), (415, 148))]}
    assert [ch for _, ch in DL._form_marks([box, tick])] == ["☒"]
    assert [ch for _, ch in DL._form_marks([box])] == ["□"]


def test_inset_cell_frame_is_not_form_field():
    """N17: внутренняя рамка ячейки «двойной» таблицы (отступ 1,5 pt со всех сторон) — не поле ввода; поле в большой
    ячейке — поле; клетка, совпавшая с сеткой, — не отметка."""
    xs, ys = [57, 250, 400], [100, 120, 160]
    assert not DL._in_cell((58.5, 101.5, 248.5, 118.5), xs, ys, whole=True)
    assert DL._in_cell((260, 125, 360, 141), xs, ys, whole=True)
    assert not DL._in_cell((250, 100, 400, 120), xs, ys)


def _two_tables(path, xs2):
    """Лист 1: текст и внизу таблица A; лист 2 начат таблицей с тем же числом столбцов (столбцы xs2)."""
    doc = fitz.open()
    pg = doc.new_page(width=595, height=842)
    for k in range(25):
        pg.insert_text((57, 80 + 14 * k), f"Bevezeto szöveg {k + 1}. sora.", fontsize=10, fontname="tiro")
    sh = pg.new_shape()
    ys1 = [520 + 25 * k for k in range(11)]
    _grid(sh, [57, 200, 400, 545], ys1)
    sh.commit()
    for r in range(10):
        for c, x in enumerate((57, 200, 400)):
            pg.insert_text((x + 4, ys1[r] + 16), f"A{r}{c}", fontsize=10, fontname="tiro")
    pg2 = doc.new_page(width=595, height=842)
    sh = pg2.new_shape()
    ys2 = [60 + 25 * k for k in range(9)]
    _grid(sh, xs2, ys2)
    sh.commit()
    for r in range(8):
        for c, x in enumerate(xs2[:3]):
            pg2.insert_text((x + 4, ys2[r] + 16), f"B{r}{c}", fontsize=10, fontname="tiro")
    for k in range(20):
        pg2.insert_text((57, 300 + 14 * k), f"Záró szöveg {k + 1}. sora.", fontsize=10, fontname="tiro")
    doc.save(str(path))


def test_next_page_starting_with_other_table_is_not_continuation(tmp_path):
    """N3: лист 2, начатый другой таблицей (те же три столбца, другие ширины), — не продолжение таблицы листа 1;
    продолжение с теми же столбцами — одна таблица."""
    a, b = tmp_path / "a.pdf", tmp_path / "b.pdf"
    _two_tables(a, [57, 260, 420, 545])
    _two_tables(b, [57, 200, 400, 545])
    tabs = lambda f: [x for x in DL.load(str(f), _segs(f), text=True)[0] if x["k"] == "table"]
    assert len(tabs(a)) == 2
    assert len(tabs(b)) == 1 and tabs(b)[0].get("split")


def test_text_page_with_big_figure_is_not_scan_snapshot(tmp_path):
    """N8: лист текстового слоя с крупным рисунком (картинка > ½ листа) и своим текстом — текст абзацами, а не снимок
    листа: снимок — только у листа, распознанного OCR."""
    import numpy as np
    src = tmp_path / "fig.pdf"
    doc = fitz.open()
    for pno in range(3):
        pg = doc.new_page(width=595, height=842)
        if pno == 1:
            a = np.full((520, 515, 3), 230, np.uint8)
            a[50:470, 50:465] = 40
            pg.insert_image(fitz.Rect(40, 50, 555, 570), pixmap=fitz.Pixmap(fitz.csRGB, 515, 520, a.tobytes(), False))
            for k in range(8):
                pg.insert_text((57, 625 + 16 * k), f"Magyarázat a rajzhoz, {k + 1}. sor.", fontsize=10, fontname="tiro")
        else:
            for k in range(30):
                pg.insert_text((57, 80 + 16 * k), f"{pno + 1}. lap, {k + 1}. sor: szöveg.", fontsize=10, fontname="tiro")
    doc.save(str(src))
    assert DL.is_scan(fitz.open(str(src))[1])
    blocks, _ = DL.load(str(src), _segs(src), text=True)
    snaps = [x for x in blocks if x["k"] == "p" and x.get("runs") and x["runs"][0].get("img", {}).get("clip")]
    text = " ".join("".join(q.get("t", "") for q in x.get("runs") or []) for x in blocks)
    assert not snaps and "Magyarázat a rajzhoz, 8. sor." in text


def test_title_box_in_head_zone_is_not_table_head_but_short_continuation_is(tmp_path):
    """N2: штамп листа — таблица в зоне шапки, во второй строке текст меняется по листам, — не «шапка таблицы тела»
    (она уходила в тело один раз); продолжение таблицы в одну строку вверху листа, чья шапка — шапка таблицы с телом
    на прошлом листе, — шапка таблицы."""
    src = tmp_path / "box.pdf"
    doc = fitz.open()
    for pno in range(3):
        pg = doc.new_page(width=595, height=842)
        sh = pg.new_shape()
        _grid(sh, [98, 166, 564], [21, 66, 81])
        put = lambda x, y, t: pg.insert_text((x, y), t, fontsize=9, fontname="tiro")
        put(170, 45, "KKS NUMBERING PHILOSOPHY")
        put(170, 77, "x" * (20 + 3 * pno))                 # меняется по листам
        if pno == 0:
            _grid(sh, [57, 250, 545], [600 + 25 * k for k in range(8)])
            put(60, 616, "Megnevezés")
            put(253, 616, "Művelet")
            for k in range(1, 7):
                put(60, 616 + 25 * k, f"Sor {k}")
                put(253, 616 + 25 * k, "ZÁRVA")
        if pno == 1:
            _grid(sh, [57, 250, 545], [90, 115, 135])
            put(60, 106, "Megnevezés")
            put(253, 106, "Művelet")
            put(60, 128, "Sor 7")
            put(253, 128, "NYITVA")
        if pno == 2:                                       # продолжение без шапки, две строки вверху листа
            _grid(sh, [57, 250, 545], [86, 100, 114])
            put(60, 96, "Sor 8")
            put(253, 96, "NYITVA")
            put(60, 110, "Sor 9")
            put(253, 110, "ZÁRVA")
        for k in range(10):
            put(57, 200 + 18 * k, f"{pno + 1}. lap, {k + 1}. sor: folyó szöveg.")
        sh.commit()
    doc.save(str(src))
    segs = _segs(src)
    d = fitz.open(str(src))
    H = d[0].rect.height
    keys = DL._running(segs, H, len(d))
    heads = DL._head_rows(d, DL._page_rules(d, True), segs, keys, H)
    got = {(s["page"], s["text"]) for s in segs if s["id"] in heads}
    assert not any("KKS" in t for _, t in got)
    assert (1, "Megnevezés") in got and (0, "Megnevezés") in got and (2, "Sor 8") in got


def test_soft_hyphens_skip_abbreviations():
    """N16: аббревиатура заглавными посреди строчного текста — без мягких переносов; слово заглавными в тексте
    заглавными (шапка) — с ними."""
    run = Document().add_paragraph().add_run("Шкаф ЦАСУТП блока")
    B._soft_hyphens(run)
    assert not run._r.findall(W + "softHyphen")
    run = Document().add_paragraph().add_run("ЗНАЧЕНИЕ ПАРАМЕТРА")
    B._soft_hyphens(run)
    assert run._r.findall(W + "softHyphen")
    assert B._pieces("ЦАСУТП", "Шкаф ЦАСУТП") == ["ЦАСУТП"] and len(B._pieces("ЗНАЧЕНИЕ", "ЗНАЧЕНИЕ")) == 3
    assert B._pieces("контуре(⁶)") == ["кон-", "ту-", "ре(⁶)"]                   # слово со скобкой — по буквенной части


def _narrow_table(texts, xs):
    """Таблица по сетке линий (docx_layout): строки texts, столбцы по xs."""
    g = lambda x0: {"page": 1, "top": 100, "bottom": 110, "x0": x0, "x1": x0 + 10, "first_x0": x0, "size": 10,
                    "lines": 1}
    cell = lambda t, x0: {"blocks": [{"k": "p", "runs": [{"t": t, "font": "Tahoma", "size": 10, "b": False, "i": False,
                                                          "color": "000000"}], "geom": g(x0)}], "th": False,
                          "geom": g(x0), "span": 1, "edges": {"top": True, "bottom": True, "left": True, "right": True}}
    rows = [[cell(t, x) for t, x in zip(r, xs)] for r in texts]
    ys = [100 + 12 * k for k in range(len(rows) + 1)]
    return {"k": "table", "rows": rows, "borders": True, "frame": {"xs": xs, "ys": ys, "row_lines": True},
            "geom": {"page": 1, "top": ys[0], "bottom": ys[-1], "x0": xs[0], "x1": xs[-1], "first_x0": xs[0], "size": 10}}


def test_narrow_number_columns_keep_size_and_only_overflowing_paragraph_shrinks(tmp_path):
    """N14: узкие столбцы чисел (по 24 pt при полях ячейки 2,4 pt, ÜFK I л. 32) — кегль тот же: числа влезают; в столбце,
    где шапка не влезает и слогом (код без мест переноса), мельче только шапка, а не «–» под ней (л. 9, 20.10)."""
    xs = [40, 90] + [90 + 24 * i for i in range(1, 9)] + [342]
    nums = [str(100 + 25 * i) for i in range(8)]
    tab = _narrow_table([["Ág TH (°C)"] + nums + ["31RD31S201"], ["pI (bar)"] + nums + ["–"]], xs)
    out = tmp_path / "n.docx"
    mem = B.Memory([(t, t) for t in nums + ["31RD31S201", "–"]] + [("Ág TH (°C)", "Ветвь ТН (°C)"),
                                                                  ("pI (bar)", "pI (бар)")])
    B.Builder(mem, lambda img: png(), dict(PROFILE, table={"pad_left": 2.4})).build(
        [tab], str(out), dict(LAYOUT, first_is_body=True))
    t = Document(str(out)).tables[0]
    sized = [[p.style.name.endswith("pt") for p in c.paragraphs] for r in t.rows for c in r.cells]
    assert not any(any(s) for s in sized[1:9]) and not any(any(s) for s in sized[11:19])   # числа — своим кеглем
    assert any(sized[9]) and not any(sized[19])                                            # код шапки — да, «–» — нет


def test_running_counts_pages_not_pieces(tmp_path):
    """Колонтитул — повтор на листах, а не кусков: строка таблицы из трёх кусков у верха одностраничного документа —
    содержимое (30RDGT modj: шапка таблицы уходила в шапку DOCX); та же строка на двух листах — колонтитул."""
    def pdf(path, pages):
        doc = fitz.open()
        for pno in range(pages):
            pg = doc.new_page(width=595, height=842)
            sh = pg.new_shape()
            _grid(sh, [42, 98, 466, 570], [64, 88, 112, 300])
            sh.commit()
            for x, t in ((48, "Verzió"), (201, "A módosítás megnevezése"), (471, "Jóváhagyás dátuma")):
                pg.insert_text((x, 80), t, fontsize=9, fontname="tiro")
            for x, t in ((48, "0"), (102, "Ismételt új kiadás."), (473, "2010-03-29")):
                pg.insert_text((x, 104), t, fontsize=9, fontname="tiro")
            for k in range(8):
                pg.insert_text((102, 130 + 18 * k), f"{pno + 1}. lap, {k + 1}. sor.", fontsize=9, fontname="tiro")
        doc.save(str(path))
    one, two = tmp_path / "one.pdf", tmp_path / "two.pdf"
    pdf(one, 1)
    pdf(two, 2)
    assert not DL._running(_segs(one), 842, 1)
    keys = DL._running(_segs(two), 842, 2)
    assert ("head", "Verzió") in keys
    blocks, _ = DL.load(str(one), _segs(one), text=True)
    text = " ".join("".join(q.get("t", "") for q in x["runs"]) for b in blocks if b["k"] == "table"
                    for r in b["rows"] for c in r for x in c["blocks"])
    assert "Verzió" in text and "Jóváhagyás dátuma" in text


def test_cell_rules_short_of_row_lines_still_make_table(tmp_path):
    """Таблица без боковых рамок: вертикали ячеек не доходят до линеек строк на 3–4 pt, шапка залита серым (ÜFK I
    л. 83 текстовым слоем, «Beavatkozások»): сетка распадалась на куски по одной линейке, и таблица вместе с подписью
    над ней уходила рисунком. Должна быть таблица 3 столбца с шапкой, подпись — абзацем."""
    src = tmp_path / "t.pdf"
    doc = fitz.open()
    pg = doc.new_page(width=595, height=842)
    sh = pg.new_shape()
    sh.draw_rect(fitz.Rect(57, 255, 542, 286))
    sh.finish(color=None, fill=(0.85, 0.85, 0.85))
    ys = [255, 286, 368, 489]
    for y in ys:
        sh.draw_line((57, y), (542, y))
    for x in (228, 461):
        for a, b in zip(ys, ys[1:]):
            sh.draw_line((x, a + 3.5), (x, b - 0.5))
    sh.finish(color=(0, 0, 0), width=0.5)
    sh.commit()
    for x, y, t in ((57, 248, "Beavatkozások:"), (100, 275, "Eltérés"), (290, 275, "Szükséges beavatkozás"),
                    (466, 275, "Végrehajtási határ"), (62, 300, "A. A korlátozó feltétel"),
                    (62, 312, "NEM TELJESÜL."), (232, 300, "A.1. Helyre kell állítani."), (466, 300, "A.1. 360 óra"),
                    (62, 380, "B. Az A.1. időtartam LETELT."), (232, 380, "B.1. A blokkot le kell állítani."),
                    (466, 380, "B.1. 288 óra")):
        pg.insert_text((x, y), t, fontsize=9, fontname="tiro")
    doc.save(str(src))
    blocks, _ = DL.load(str(src), _segs(src), text=True)
    assert not any(b.get("figure") for b in blocks)
    tab = [b for b in blocks if b["k"] == "table"]
    assert len(tab) == 1 and max(len(r) for r in tab[0]["rows"]) == 3 and len(tab[0]["rows"]) == 3
    cell = lambda c: " ".join("".join(q.get("t", "") for q in x.get("runs") or []) for x in c["blocks"])
    assert [cell(r[-1]) for r in tab[0]["rows"][1:]] == ["A.1. 360 óra", "B.1. 288 óra"]
    assert [cell(c) for c in tab[0]["rows"][0]] == ["Eltérés", "Szükséges beavatkozás", "Végrehajtási határ"]
    assert any(b["k"] == "p" and "Beavatkozások:" in "".join(q.get("t", "") for q in b["runs"]) for b in blocks)


def test_field_boxes_under_row_line_are_not_table_columns():
    """Рамки полей ввода в ячейке бланка начинаются в 3 pt под линейкой строки — не столбцы таблицы: доходят до своей
    короткой нижней кромки, а не до линейки строки (30RDGT: 3 столбца → 7, 15 листов → 29, 22.10)."""
    hs = [(81, 42, 574), (117, 42, 574), (300, 42, 574), (769, 42, 574),
          (120, 92.5, 206), (136, 92.5, 206), (120, 209, 298), (136, 209, 298)]   # две рамки полей под линейкой 117
    vs = [(42, 81, 769), (90, 81, 769), (469, 81, 769), (574, 81, 769),
          (92.5, 120, 136), (206, 120, 136), (209, 120, 136), (298, 120, 136)]
    tabs = DL._tables(hs, vs)
    big = max(tabs, key=lambda t: t[0][-1] - t[0][0])
    assert [round(x) for x in big[0]] == [42, 90, 469, 574]


def _glued_pdf(path, pages=6):
    """Лист за листом: шапка «Azonosító» + «Oldalszám:N.» (номер приклеен к подписи без пробела), тело разное."""
    doc = fitz.open()
    for pno in range(pages):
        pg = doc.new_page(width=595, height=842)
        put = lambda x, y, t: pg.insert_text((x, y), t, fontsize=10, fontname="tiro")
        put(60, 40, "Azonosító: 30RR2126GT")
        put(440, 40, f"Oldalszám:{pno + 2}.")
        words = ("elsődleges tömítettség szivattyú indítás szelep zárás nyomás hőmérséklet teljesítmény áramlás "
                 "szint vezérlés védelem mérés kapcsoló betáplálás hűtés fűtés csővezeték").split()
        for k in range(8):
            put(60, 120 + 40 * k, f"A {words[(pno * 3 + k) % len(words)]} és a {words[(pno * 5 + k * 2 + 1) % len(words)]}"
                                  f" ellenőrzése megtörtént, az eredmény megfelelő ({words[(pno + k * 7) % len(words)]}).")
        put(60, 812, "Nyomtatva: 2026-10-28")
    doc.save(str(path))


def test_page_number_glued_to_label_is_a_field(tmp_path):
    """Номер листа, приклеенный к подписи без пробела («Oldalszám:7.», 30RR2126GT), — поле PAGE в шапке; прежде слово
    целиком не складывалось в число, и во всех шапках DOCX стояло число одного листа («Номер страницы: 7»)."""
    src = tmp_path / "glued.pdf"
    _glued_pdf(src)
    segs = _segs(src)
    blocks, info = DL.load(str(src), segs, text=True)
    L = D.layout(str(src), blocks, content=info.get("content_boxes"))
    assert L["page_tok"] and L["page_marked"]
    out = tmp_path / "glued.docx"
    pairs = [(s["text"], s["text"].replace("Oldalszám:", "Номер страницы: ").replace("Azonosító", "Идентификатор"))
             for s in segs]
    B.Builder(B.Memory(pairs), lambda img: png(), PROFILE).build(
        blocks, str(out), dict(LAYOUT, **{k: L[k] for k in ("head_grid", "page_tok", "page_marked", "num_start")
                                            if k in L}))
    head = "".join(h._element.xml for s in Document(str(out)).sections for h in (s.header,))
    assert "PAGE" in head and "Номер страницы" in head


def test_notdef_glyph_of_symbol_font_is_a_box():
    """Пустой глиф шрифта Symbol/Wingdings (код 0, в текстовом слое «�») — флажок «□»; в других шрифтах «�» не трогается."""
    import layout as LY
    assert LY.span_text({"text": "�", "font": "SymbolMT"}) == "□"
    assert LY.span_text({"text": "� ok", "font": "Wingdings-Regular"}) == "□ ok"
    assert LY.span_text({"text": "�", "font": "TimesNewRomanPSMT"}) == "�"
    assert LY.span_text({"text": "abc", "font": "SymbolMT"}) == "abc"


def test_row_number_in_middle_of_two_line_cell_opens_paragraph_once():
    """Номер строки («2.8.») по центру ячейки лежит против обеих строк предложения: начало абзаца — первая строка, вторая
    продолжает; прежде предложение переводилось двумя обрывками («…участниками и» / «между блочными щитами», 28.10)."""
    import layout as LY
    fr = lambda t, bb: {"t": t, "bbox": bb, "size": 11.0, "bold": False, "family": "serif"}
    frags = [fr("Folyamatos hírkapcsolatot kell biztosítani a résztvevők és a", [154, 628, 429, 643]),
             fr("2.8.", [81, 635, 101, 650]),
             fr("blokkvezénylő között.", [154, 642, 251, 656])]
    pars = LY.group(frags, [], [])
    texts = sorted(p["text"] for p in pars)
    assert texts == ["2.8.", "Folyamatos hírkapcsolatot kell biztosítani a résztvevők és a blokkvezénylő között."]
    # номер и название на одной строке — по-прежнему два абзаца (оглавление 3SZ19)
    pars = LY.group([fr("1.6.", [60, 100, 80, 112]), fr("FOT szintcsökkenés", [100, 100, 200, 112]),
                     fr("FOT szint folytatás", [100, 113, 200, 125])], [], [])
    assert sorted(p["text"] for p in pars) == ["1.6.", "FOT szintcsökkenés FOT szint folytatás"]


def test_sentence_cut_mid_phrase_continues_despite_row_number_against_last_line():
    """Строка оборвана на середине фразы («…megengedett» / «szintváltozásának…»), номер строки таблицы («0.») стоит против
    последней строки: это продолжение, а не новый пункт (30TH50: «Допустимые гидроемкостей САОЗ изменение изменения
    уровня», 28.10). Пункт списка со своим номером в первой строке и строчной буквой по-прежнему отдельный."""
    import layout as LY
    fr = lambda t, bb: {"t": t, "bbox": bb, "size": 12.0, "bold": False, "family": "serif"}
    frags = [fr("Ciklikus felülvizsgálat. Ismételt új „0” verziós TU", [113, 158, 358, 174]),
             fr("kiadása. A hidroakkumulátorok megengedett", [113, 172, 332, 188]),
             fr("0.", [85, 187, 95, 201]),
             fr("szintváltozásának módosítása 25mm-ről 2.5mm-re.", [113, 186, 362, 202])]
    texts = sorted(p["text"] for p in LY.group(frags, [], []))
    assert texts == ["0.", "Ciklikus felülvizsgálat. Ismételt új „0” verziós TU kiadása. A hidroakkumulátorok megengedett "
                     "szintváltozásának módosítása 25mm-ről 2.5mm-re."]
    # пункты списка: у первой строки каждого — свой номер, текст со строчной, без точки в конце — каждый пункт отдельно
    items = [fr("1.", [60, 100, 70, 112]), fr("a nyomás mérése az első szakaszon", [90, 100, 300, 112]),
             fr("2.", [60, 113, 70, 125]), fr("a szelep zárása a második szakaszon", [90, 113, 300, 125])]
    assert len([p for p in LY.group(items, [], []) if len(p["text"]) > 3]) == 2


def test_tall_line_boxes_with_tight_pitch_stay_one_paragraph():
    """Моноширинный шрифт 3SZ17: рамка строки 15 pt при шаге строк 10 pt — рамки соседних строк перекрываются на 5 pt;
    строка 2 не подходила к абзацу строки 1, а строка 3 — подходила, и абзац шёл через строку («valaséhez», 28.10)."""
    import layout as LY
    fr = lambda t, y: {"t": t, "bbox": [323, y, 557, y + 15], "size": 9.5, "bold": False, "family": "mono"}
    frags = [fr("MEGJEGYZÉS: A 30UJ04S201 vala-", 108), fr("mint a 30UJ04D001 szivattyú BV-i működteté-", 118),
             fr("séhez az ENGEDÉLYEZÉS gomb megnyomása,", 128), fr("valamint az érintett működtető kapcsolók", 138)]
    pars = LY.group(frags, [], [])
    assert len(pars) == 1
    assert pars[0]["text"] == ("MEGJEGYZÉS: A 30UJ04S201 valamint a 30UJ04D001 szivattyú BV-i működtetéséhez az "
                               "ENGEDÉLYEZÉS gomb megnyomása, valamint az érintett működtető kapcsolók")
    # два кусочка на одной строке (разные столбцы) — по-прежнему два абзаца
    side = [{"t": "bal", "bbox": [71, 100, 120, 115], "size": 9.5, "bold": False, "family": "mono"},
            {"t": "jobb", "bbox": [330, 100, 380, 115], "size": 9.5, "bold": False, "family": "mono"}]
    assert len(LY.group(side, [], [])) == 2


def test_trailing_space_does_not_hide_column_rule():
    """Хвостовой пробел моноширинной строки доходит до линии столбца: рамка «…tudja » кончалась правее линии, и строки
    двух столбцов склеивались в один кусок («…kézzel tudja Az érintett tartályokhoz», 3SZ17, 28.10)."""
    import layout as LY
    f = {"t": "kell lennie, hogy kézzel tudja ", "bbox": [71, 283, 321, 299]}
    assert LY._ink_x1(f) < 319.1 and LY._vrule_in([(234, 316, 319.1)], LY._ink_x1(f), 323, 283, 299)
    assert LY._ink_x1({"t": "tudja", "bbox": [71, 283, 321, 299]}) == 321


def test_page_number_found_by_place_when_label_differs_per_page(tmp_path):
    """Подпись номера листа разная на листах («Oldalszám:» / «Oldal» «:» / «Old.» …, 30BLFKSZ): общего набора слов нет,
    число ищется по месту — прежде оно застывало числом одного листа во всех шапках."""
    src = tmp_path / "labels.pdf"
    doc = fitz.open()
    labels = ["Oldalszám:", "Oldal :", "Old.", "Oldalsz.:", "Lap", "Lapszám:"]
    words = ("elsődleges tömítettség szivattyú indítás szelep zárás nyomás hőmérséklet teljesítmény áramlás "
             "szint vezérlés védelem mérés kapcsoló betáplálás hűtés fűtés csővezeték").split()
    for pno in range(6):
        pg = doc.new_page(width=595, height=842)
        put = lambda x, y, t: pg.insert_text((x, y), t, fontsize=10, fontname="tiro")
        if pno:                                             # лист 1 — титул без шапки
            put(60, 40, "Azonosító: 30BLFKSZ")
            put(420, 40, labels[pno])
            put(520, 40, str(pno + 1))
        for k in range(8):
            put(60, 120 + 40 * k, f"A {words[(pno * 3 + k) % len(words)]} és a {words[(pno * 5 + k * 2 + 1) % len(words)]}"
                                  f" ellenőrzése megtörtént ({words[(pno + k * 7) % len(words)]}).")
    doc.save(str(src))
    segs = _segs(src)
    blocks, info = DL.load(str(src), segs, text=True)
    L = D.layout(str(src), blocks, content=info.get("content_boxes"))
    assert L["page_tok"] and L["page_marked"]


def test_header_label_column_widens_for_longer_russian_label():
    """Столбец подписи шапки («Azonosító:» 55 pt) не вмещает «Идентификатор:» — слово рвалось посреди («Идентифика / тор:»);
    ширина добавляется за счёт самого широкого столбца, общая ширина таблицы та же; без нужды ширины не меняются."""
    rows = [{"cells": ["Azonosító:", "30KBFPRBIGT", "Verziószám:", "0."], "bold": True}]
    mem = B.Memory([("Azonosító:", "Идентификатор:"), ("Verziószám:", "Номер версии:"), ("30KBFPRBIGT", "30KBFPRBIGT"),
                    ("0.", "0.")])
    b = B.Builder(mem, lambda img: png(), PROFILE)
    w = b._fit_widths(rows, [55, 300, 70, 40], 11.0)
    assert w[0] > 55 and w[2] > 70 and abs(sum(w) - 465) < 1 and w[1] >= 0.6 * 300
    assert b._fit_widths(rows, [120, 200, 100, 45], 11.0) == [120, 200, 100, 45]


def test_header_words_keep_glyphs_of_fallback_font_raised_by_two_points():
    """«ő» запасного шрифта стоит на 2,1 pt выше соседей: при допуске по высоте 2 pt он отрывался «словом» «ő ő ő», а
    в названии не хватало букв («gzfejleszt», «ellenrz»; шапка 30GF118GT, 28.10)."""
    chars, x = [], 100.0
    for ch in "gőzfejlesztő":
        raised = ch == "ő"
        chars.append({"text": ch, "x0": x, "x1": x + 5.5, "top": 54.6 if raised else 56.7, "bottom": 65.5 if raised else 67.6,
                      "size": 10.9, "upright": True, "doctop": 0, "fontname": "Times"})
        x += 5.5
    assert [w["text"] for w in D._words(chars)] == ["gőzfejlesztő"]
    # слово с «ő» выше соседей по строке — верх выровнен, строка колонтитула не рвётся при округлении верха
    row = [dict(c, x0=c["x0"] + 200, x1=c["x1"] + 200) for c in chars]
    plain = [{"text": t, "x0": 20.0 + 6 * i, "x1": 25.5 + 6 * i, "top": 56.7, "bottom": 67.6, "size": 10.9,
              "upright": True, "doctop": 0, "fontname": "Times"} for i, t in enumerate("blokk")]
    got = D._words(plain + row)
    assert len({round(w["top"] / 3) for w in got}) == 1 and [w["text"] for w in got] == ["blokk", "gőzfejlesztő"]
    assert D._words([]) == []


def test_full_page_row_is_capped_below_body_height(tmp_path):
    """Строка таблицы во весь лист оригинала задана «не ниже» его высоты и вместе с шапкой выше поля тела (линии оригинала
    доходят ниже измеренного поля); абзац в 1 pt перед таблицей раздела — и таблица уходила на следующий лист, оставляя
    пустой (30THHAGT, 28.10). Потолок высоты строки — поле тела минус 60 pt; обычные строки не трогаются."""
    g = lambda top: {"page": 1, "top": top, "bottom": top + 10, "x0": 80, "x1": 300, "first_x0": 80, "size": 10,
                     "lines": 1}
    cellp = lambda t, top: {"blocks": [{"k": "p", "runs": [{"t": t, "font": "Tahoma", "size": 10, "b": False,
                                                            "i": False, "color": "000000"}], "geom": g(top)}],
                            "th": False, "geom": g(top)}
    tab = {"k": "table", "rows": [[cellp("1.1.", 100), cellp("Első", 100)], [cellp("1.2.", 112), cellp("Második", 112)]],
           "borders": True, "geom": dict(g(100), x1=500),
           "frame": {"xs": [70, 120, 500], "ys": [78.0, 90.0, 742.0], "row_lines": True}}
    out = tmp_path / "cap.docx"
    B.Builder(B.Memory([("Első", "Первый"), ("Második", "Второй")]), lambda img: png(), PROFILE).build(
        [tab], str(out), dict(LAYOUT, page_tok="2", num_start={"page": 1, "value": 2}, first_is_body=True))
    heights = [int(h.get(W + "val")) / 20 for h in Document(str(out)).element.body.iter(W + "trHeight")]
    body_h = LAYOUT["H"] - LAYOUT["top"] - LAYOUT["bottom"]
    assert heights and max(heights) <= body_h - 60 + 0.1 and min(heights) == 12.0


def test_flowchart_of_many_shape_bitmaps_is_a_whole_page_figure():
    """Блок-схема Word: фигура — десятки картинок (градиент, тень), надписи лежат на них. Рисунок по краске не находился
    (текст вычтен), схема шла таблицами и абзацами россыпью на трёх листах (30GF(135)51GT, л. 19, 28.10): лист с ≥ 60
    картинками, на которых лежит половина надписей, — снимок листа с переведённого PDF; лист с мелкими значками,
    рядом с которыми стоит текст, — нет."""
    import layout as LY
    import docx_layout as DLY
    png_bytes = png()
    doc = fitz.open()
    page = doc.new_page(width=842, height=595)
    segs = []
    for k in range(70):
        x, y = 30 + (k % 10) * 78, 60 + (k // 10) * 60
        page.insert_image(fitz.Rect(x, y, x + 70, y + 30), stream=png_bytes)
        page.insert_text((x + 4, y + 18), "szöveg", fontsize=8)
        bb = [x + 4, y + 10, x + 40, y + 22]
        segs.append({"id": k, "page": 0, "text": "szöveg", "loc": {"bbox": bb, "frags": [bb]}})
    assert DLY._graphic_page(page, segs, []) is True
    icons = doc.new_page(width=842, height=595)
    segs2 = []
    for k in range(70):
        x, y = 30, 20 + k * 7.5
        icons.insert_image(fitz.Rect(x, y, x + 5, y + 5), stream=png_bytes)
        icons.insert_text((x + 12, y + 5), "pont", fontsize=5)
        bb = [x + 12, y, x + 30, y + 6]
        segs2.append({"id": k, "page": 1, "text": "pont", "loc": {"bbox": bb, "frags": [bb]}})
    assert DLY._graphic_page(icons, segs2, []) is False


def test_overlapping_header_texts_are_not_interleaved_by_position():
    """Название бланка залезает под «Dátum:» и дату: по положению знаки двух надписей перемешивались
    («áDtjááturmha:tóságának», 30RR0102, 28.10) — слова берутся в порядке потока PDF."""
    chars, x = [], 100.0
    def put(text, x0, step=5.0):
        nonlocal chars
        for k, ch in enumerate(text):
            chars.append({"text": ch, "x0": x0 + k * step, "x1": x0 + k * step + step, "top": 56.0, "bottom": 66.0,
                          "size": 10.0, "upright": True, "doctop": 0, "fontname": "Times"})
    put("atjarhatosag", 100.0)            # название: 100…160
    put("Datum:", 110.0)                  # метка лежит поверх названия
    got = [w["text"] for w in D._words(chars)]
    assert "atjarhatosag" in got and "Datum:" in got


def test_caps_word_split_between_cell_paragraphs_is_one_paragraph():
    """Шапка таблицы 30YCGT: «ELLENŐR-» и «ZÉS HELYE» — два P одной ячейки; с текстом страницы («ELLENŐR- ZÉS HELYE») они
    не совпадали, и половины переводились «МЕСТО» / «КОНТРОЛЯ» (28.10). Вне ячейки и со строчной второй частью — как было."""
    g = {"page": 4, "top": 100, "bottom": 110, "x0": 80, "x1": 140, "first_x0": 80, "size": 9, "lines": 1}
    para = lambda t: {"k": "p", "runs": [{"t": t, "font": "Tahoma", "size": 9, "b": True, "i": False, "color": "000000"}],
                      "geom": dict(g)}
    cell = lambda *ts: {"blocks": [para(t) for t in ts], "th": True}
    tab = {"k": "table", "rows": [[cell("ELLENŐR-", "ZÉS HELYE"), cell("MEGNEVEZÉS-", "tetel")]], "geom": dict(g)}
    out = D.join_pages([tab, para("ELLENŐR-"), para("ZÉS HELYE")])
    texts = [[D.text_of(x["runs"]) for x in c["blocks"]] for c in out[0]["rows"][0]]
    assert texts == [["ELLENŐR- ZÉS HELYE"], ["MEGNEVEZÉS-", "tetel"]]
    assert len(out) == 3                                        # вне ячейки абзацы не склеиваются

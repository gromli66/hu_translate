# -*- coding: utf-8 -*-
"""Ворота качества готовых переводов (tools/audit_out.py): каждая проверка на малом примере + сквозной прогон по
собранному в tmp_path заданию (JSON, DOCX, PDF перевода, PDF оригинала). Данных проекта и Docker не требуется."""
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import audit_out as A  # noqa: E402

FONT = next((f for f in (r"C:\Windows\Fonts\arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf") if Path(f).exists()), None)
needs_font = pytest.mark.skipif(FONT is None, reason="нет шрифта с кириллицей для синтетических PDF")


def W(x0, y0, x1, y1, t):
    return (x0, y0, x1, y1, t, 0, 0, 0)


# ---------------------------------------------------------------- 2. повтор фразы
def test_repeat_phrase_three_times():
    r = A.find_repeat("МАШИННЫЙ ЗАЛ ряд 28КА " * 3 + "конец")
    assert r == ("МАШИННЫЙ ЗАЛ ряд 28КА", 3)


def test_repeat_twice_is_not_a_defect():
    assert A.find_repeat("МАШИННЫЙ ЗАЛ ряд 28КА МАШИННЫЙ ЗАЛ ряд 28КА конец") is None


def test_repeat_ignores_marks_and_short_words():
    assert A.find_repeat("○ ○ ○ ○ ○ ○ ○ ○ ○ ○") is None            # стопка значков — дело проверки 6
    assert A.find_repeat("Да Да Да Да") is None                     # короче 10 знаков
    assert A.find_repeat("0,0 м 0,0 м 0,0 м 0,0 м") is None         # числа без букв-слов


def test_repeat_inside_long_text():
    r = A.find_repeat("Начало абзаца. " + "30UY01/1 закрыт, " * 4 + "дальше текст")
    assert r and r[1] == 4 and "закрыт" in r[0]


# ---------------------------------------------------------------- 1. потеря содержимого
def seg(i, text, page=0, **loc):
    return {"id": i, "text": text, "page": page, "loc": loc}


def test_lost_segments_found_and_filtered():
    segs = [seg(1, "Zárva vannak az alábbi armatúrák:"), seg(2, "Nyitás (bar)"), seg(3, "Azonosító: 3SZ23"),
            seg(4, "Lapszám: 1."), seg(5, "30RR21S002"), seg(6, "Hőmérséklet", role="header"),
            seg(7, "Állapot tölt.", shot={"t": "x"}), seg(8, "Oldalszám: 4"), seg(9, "Szabályozó szelep zárva", ocr=True),
            seg(10, "Új javítás szükséges", 4)]
    res = {"1": {"ru": "Следующая арматура закрыта:"}, "2": {"ru": "Открытие (бар)"}, "3": {"ru": "Идентификатор: 3SZ23"},
           "4": {"ru": "Лист: 1."}, "5": {"ru": "30RR21S002", "copied": True}, "6": {"ru": "Температура"},
           "7": {"ru": "Сост. заправки"}, "8": {"ru": "Номер страницы: 4"}, "9": {"ru": "Регулирующий клапан закрыт"},
           "10": {"ru": "Требуется новый ремонт"}}
    full = A.norm("Открытие (бар) и ещё текст")                     # нашлось только сегмент 2
    lost = A.lost_segments(segs, res, full)
    assert [s["id"] for s in lost] == [1, 10]                       # заголовок шапки, копия, колонтитул, скрин, скан не в счёт


def test_lost_segment_found_by_polished_text():
    segs = [seg(1, "Ismételt új „0” verziós KU kiadása")]
    res = {"1": {"ru": "Повторный выпуск ИЭ новой версии „0“"}}
    full = A.norm("Повторный выпуск ИЭ новой версии «0»")
    assert A.lost_segments(segs, res, full) == []                   # кавычки не отличают
    assert [s["id"] for s in A.lost_segments(segs, res, A.norm("совсем другое"))] == [1]
    assert A.lost_segments(segs, res, A.norm("Первичный выпуск ИЭ версии 0"), polish={1: "Первичный выпуск ИЭ версии 0"}) == []


def test_lost_segments_repeated_source_counts_occurrences():
    # один венгерский текст дважды; в res у повтора иная формулировка, а сборка (память) ставит везде первую
    segs = [seg(1, "Zárva vannak az alábbi armatúrák:", 0), seg(2, "Zárva vannak az alábbi armatúrák:", 3)]
    res = {"1": {"ru": "Закрыты следующие арматуры:"}, "2": {"ru": "Следующая арматура закрыта:"}}
    assert A.lost_segments(segs, res, A.norm("Закрыты следующие арматуры: ... Закрыты следующие арматуры:")) == []
    lost = A.lost_segments(segs, res, A.norm("Закрыты следующие арматуры: и только один раз"))
    assert [s["id"] for s in lost] == [2]                           # вхождений 1 на два сегмента — не хватило одному
    assert [s["id"] for s in A.lost_segments(segs, res, A.norm("пусто совсем"))] == [1, 2]


# ---------------------------------------------------------------- 3. шапки
HU_HDR = "Azonosító: 3SZ23 \nVerziószám: 1\nMegnevezés: TURBINA\nOldalszám: 5 \n"
RU_HDR = "Идентификатор: 3SZ23 Номер версии: 1\nНаименование: ТУРБИНА Номер \nстраницы: 5\nтекст"


def test_header_of_both_languages():
    assert A.header_of(HU_HDR, "hu") == ("3SZ23", 1, 5)
    assert A.header_of(RU_HDR, "ru") == ("3SZ23", 1, 5)
    assert A.header_of("просто текст листа", "ru") is None
    assert A.header_of("Azonosító:\n 30RL2126GT\n Verziószám:\n 5\n", "hu") == ("30RL2126GT", 5, None)


def test_compare_headers_version_lost():
    ho = [("3SZ23", 0, 1), ("3SZ23", 1, 2), ("3SZ23/INFO", 0, 1)]
    hr = [("3SZ23", 0, 1), ("3SZ23", 0, 2), ("3SZ23", 0, 3)]
    out = A.compare_headers(ho, hr)
    msgs = [m for lvl, m in out if lvl == A.RED]
    assert any("3SZ23, версия 1" in m for m in msgs) and any("3SZ23/INFO" in m for m in msgs)


def test_compare_headers_same_is_clean_and_order_is_yellow():
    ho = [("A", 0, 1), ("A", 1, 2), ("A", 0, 3)]
    assert A.compare_headers(ho, list(ho)) == []
    out = A.compare_headers(ho, [("A", 0, 1), ("A", 0, 2), ("A", 1, 3)])
    assert [lvl for lvl, _ in out] == [A.YELLOW]
    assert A.compare_headers([], [("A", 1, 1)]) == []              # у оригинала шапки нет — сравнивать не с чем
    assert A.compare_headers(ho, [])[0][0] == A.RED                 # а в переводе шапка пропала


def test_compare_headers_sheet_numbers():
    out = A.compare_headers([("A", 0, 1), ("A", 0, 17)], [("A", 0, 1), ("A", 0, 26)])
    assert [lvl for lvl, _ in out] == [A.YELLOW] and "до 17" in out[0][1] and "до 26" in out[0][1]


# ---------------------------------------------------------------- 5. коды
def test_compare_codes():
    co, _ = A.codes_in(["30RR21S002 30RR21S003 30RR21S003", "31UJ21S601"])
    cr, _ = A.codes_in(["30RR21S002 30RR21S002", "30XX99S999 30RR21S003"])
    gone, extra, more, less = A.compare_codes(co, cr)
    assert gone == {"31UJ21S601": 1} and extra == {"30XX99S999": 1}
    assert more == {"30RR21S002": 1} and less == {"30RR21S003": 1}


def test_code_glued_to_word_counts_like_clean():
    # «30TK62P101szerint» слито в оригинале — в переводе «30TK62P101 согласно» не должно быть «лишним»
    co, _ = A.codes_in(["a 30TK62P101szerint p=3 bar"])
    cr, _ = A.codes_in(["по 30TK62P101 согласно p=3 бар"])
    assert A.compare_codes(co, cr) == ({}, {}, {}, {})
    assert A.codes_in(["30TK62P1011"])[0] == {}                      # лишняя цифра — уже не код


# ---------------------------------------------------------------- 6. отметки
def test_marks_and_glued_rows():
    row = [W(100, 100, 107, 108, "O"), W(111, 100, 118, 108, "O"), W(122, 100, 129, 108, "O")]
    other = [W(100, 200, 107, 208, "O"), W(300, 200, 307, 208, "O")]            # две отметки в разных колонках — не слиплись
    assert len(A.marks_on(row + [W(0, 0, 5, 5, "слово")])) == 3
    assert A.glued_rows(A.marks_on(row)) == 1
    assert A.glued_rows(A.marks_on(other)) == 0
    assert A.glued_rows(A.marks_on([W(100, 10 * i, 107, 10 * i + 8, "\uf0a1") for i in range(5)])) == 0   # столбик по строкам


def test_code_row_share():
    words = [W(10, 100, 70, 108, "30RR21S002"), W(300, 101, 307, 109, "○"),
             W(10, 120, 70, 128, "30RR22S002"), W(300, 140, 307, 148, "○")]      # у второго кода отметка на чужой строке
    assert A.code_row_share(words) == (2, 1)
    assert A.code_row_share([W(10, 100, 70, 108, "30RR21S002")]) == (0, 0)


def test_groups_merge_split_pages():
    gs = A.groups([(0, 0), (1, 1), (1, 2), (2, 2)], range(3), range(3))
    assert len(gs) == 1 and sorted(gs[0][0]) == [1, 2] or len(gs) == 2
    only = A.groups([(0, 0), (1, 1)], range(2), range(2))
    assert sorted(map(lambda g: (g[0], g[1]), only)) == [([0], [0]), ([1], [1])]


# ---------------------------------------------------------------- 7. венгерские буквы
def test_hungarian_leftovers():
    f = A.hu_leftovers("Проверить, что процесс nem sikerült начат; Ф.И.О. Bóth Zoltán, НСБ (BLÜGY), ÜFK, 30RR0102ÁTJGT")
    assert f == ["sikerült"]
    assert "Bóth" not in f and "BLÜGY" not in f and "ÜFK" not in f and "30RR0102ÁTJGT" not in f
    assert A.hu_leftovers("Давление в петле: tömörségvizsgálat прошло") == ["tömörségvizsgálat"]
    assert A.hu_leftovers("Péter") == []                                        # ячейка с одним именем
    assert A.hu_leftovers("tömörségvizsgálat", allowed={"tömörségvizsgálat"}) == []


# ---------------------------------------------------------------- DOCX
def test_docx_text_cells_and_units(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Абзац вне таблицы")
    t = d.add_table(rows=2, cols=2)
    for r, text in enumerate(["Ячейка один", "Ячейка один", "Ячейка один"][:2]):
        t.cell(r, 0).text = text
    t.cell(0, 1).text = "Верх"
    t.cell(0, 1).add_paragraph("Низ")
    d.sections[0].header.paragraphs[0].text = "Идентификатор: X Номер версии: 0"
    p = tmp_path / "a_RU.docx"
    d.save(p)
    dt = A.DocxText(p)
    assert "абзацвнетаблицы" in dt.full and "идентификатор" in dt.full
    units = dt.units()
    assert "Верх Низ" in units and "Абзац вне таблицы" in units


# ---------------------------------------------------------------- сквозной прогон
def make_pdf(path, pages, images=()):
    """pages — [[(x, y, текст), …], …]; images — {индекс листа: доля площади}: картинка-заливка."""
    import pymupdf
    doc = pymupdf.open()
    for i, items in enumerate(pages):
        pg = doc.new_page(width=595, height=842)
        for x, y, t in items:
            pg.insert_text((x, y), t, fontsize=10, fontname="F0", fontfile=FONT)
        if i in dict(images):
            share = dict(images)[i]
            h = 842 * share
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
            pix.clear_with(200)
            pg.insert_image(pymupdf.Rect(0, 200, 595, 200 + h), pixmap=pix)
    doc.save(path)
    doc.close()


def make_job(tmp_path, ru_paras, ru_pages, orig_pages, segs, res, ru_images=(), ru_name="doc"):
    import docx
    work = tmp_path / "work"
    for sub in ("json", "out", "in"):
        (work / sub).mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    for t in ru_paras:
        d.add_paragraph(t)
    d.save(work / "out" / f"{ru_name}_RU.docx")
    make_pdf(str(work / "out" / f"{ru_name}_RU.pdf"), ru_pages, ru_images)
    make_pdf(str(work / "in" / f"{ru_name}.pdf"), orig_pages)
    (work / "json" / f"{ru_name}.json").write_text(json.dumps({"file": "x.pdf", "segs": segs, "res": res, "info": {}}), encoding="utf-8")
    return work


def hdr_hu(v, n):
    return [(40, 40, f"Azonosító: 3SZ23 Verziószám: {v}"), (40, 55, f"Oldalszám: {n}")]


def hdr_ru(v, n):
    return [(40, 40, f"Идентификатор: 3SZ23 Номер версии: {v}"), (40, 55, f"Номер страницы: {n}")]


@needs_font
def test_end_to_end_clean(tmp_path):
    segs = [seg(1, "Zárva vannak az armatúrák", 0), seg(2, "Nyitott szelepek vannak", 1)]
    res = {"1": {"ru": "Арматура закрыта"}, "2": {"ru": "Клапаны открыты"}}
    work = make_job(tmp_path, ["Арматура закрыта", "Клапаны открыты"],
                    [hdr_ru(0, 1) + [(40, 100, "Арматура закрыта 30RR21S002")], hdr_ru(1, 2) + [(40, 100, "Клапаны открыты 30RR22S003")]],
                    [hdr_hu(0, 1) + [(40, 100, "Zárva 30RR21S002")], hdr_hu(1, 2) + [(40, 100, "Nyitott 30RR22S003")]], segs, res)
    d = A.Doc(work / "out" / "doc_RU.docx", work, [work / "in"])
    assert d.js and d.orig and d.pdf
    assert [f.as_dict() for f in d.run()] == []
    assert A.main([str(work / "out"), "--work", str(work), "--orig", str(work / "in")]) == 0


@needs_font
def test_end_to_end_defects(tmp_path, capsys):
    segs = [seg(1, "Zárva vannak az armatúrák", 0), seg(2, "Nyitott szelepek vannak", 1), seg(3, "Vezérlő szelep zárva", 1)]
    res = {"1": {"ru": "Арматура закрыта"}, "2": {"ru": "Клапаны открыты"}, "3": {"ru": "Регулирующий клапан закрыт"}}
    boom = "МАШИННЫЙ ЗАЛ ряд 28КА " * 4
    work = make_job(
        tmp_path, ["Арматура закрыта", boom, "Клапаны открыты tömörség"],                       # сегмента 3 в DOCX нет
        [hdr_ru(0, 1) + [(40, 100, "Арматура закрыта 30RR21S002")], hdr_ru(0, 2) + [(40, 100, "Клапаны открыты tömörség")],
         hdr_ru(0, 3) + [(40, 100, "хвост")], [(40, 100, "ещё")]],
        [hdr_hu(0, 1) + [(40, 100, "Zárva 30RR21S002 30RR22S003")], hdr_hu(1, 2) + [(40, 100, "Nyitott")]], segs, res)
    d = A.Doc(work / "out" / "doc_RU.docx", work, [work / "in"])
    found = {f.check: f for f in d.run()}
    assert "3 сегм" not in found["1 потеря содержимого"].msg and "1 сегм" in found["1 потеря содержимого"].msg
    assert "×4" in found["2 повтор фразы"].msg
    assert "версия 1" in found["3 шапки"].msg and found["3 шапки"].level == A.RED
    assert "×2.00" in found["4 листы"].msg                                                       # 4 листа при 2
    assert "30RR22S003" in found["5 коды оборудования"].msg
    assert "tömörség" in found["7 венгерские буквы"].msg
    assert A.main([str(work / "out"), "--work", str(work), "--orig", str(work / "in")]) == 1
    out = capsys.readouterr().out
    assert "🔴" in out and "стр." in out and "doc" in out
    assert A.main([str(work / "out"), "--work", str(work), "--orig", str(work / "in"), "--fail-on", "never"]) == 0


@needs_font
def test_end_to_end_picture_page(tmp_path):
    segs = [seg(1, "Zárva vannak az armatúrák", 0)]
    work = make_job(tmp_path, ["Арматура закрыта"], [[(40, 100, "Арматура закрыта")], [(40, 100, "x")]],
                    [[(40, 100, "Zárva vannak")]], segs, {"1": {"ru": "Арматура закрыта"}}, ru_images=[(1, 0.5)])
    d = A.Doc(work / "out" / "doc_RU.docx", work, [work / "in"])
    f = [x for x in d.run() if x.check == "8 снимок листа"]
    assert f and f[0].page == 2 and f[0].level == A.RED


@needs_font
def test_empty_table_rows_page(tmp_path):
    import pymupdf
    doc = pymupdf.open()
    pg = doc.new_page(width=595, height=842)
    pg.insert_text((60, 70), "Заголовок", fontsize=10, fontname="F0", fontfile=FONT)
    for i in range(16):                                             # сетка 16×3 без текста почти на весь лист
        for j in range(3):
            pg.draw_rect(pymupdf.Rect(40 + j * 170, 100 + i * 44, 40 + (j + 1) * 170, 100 + (i + 1) * 44), width=0.8)
    assert A.empty_row_share(pg, pg.get_text("words")) > A.EMPTY_TABLE_SHARE
    pg2 = doc.new_page(width=595, height=842)
    assert A.empty_row_share(pg2, []) == 0.0


def test_main_without_documents(tmp_path, capsys):
    assert A.main([str(tmp_path)]) == 2
    assert "Нечего проверять" in capsys.readouterr().out

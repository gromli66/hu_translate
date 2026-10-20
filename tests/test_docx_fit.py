# Подгонка перевода DOCX «на месте» под листы оригинала (docx_fit, 20.10; 3PR04: раздел = лист, 514 → 651 листов).
import zipfile

import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Pt
from lxml import etree

import docx_fit as F

W = F.W


def _doc(path, texts, size=11):
    d = Document()
    d.styles["Normal"].font.size = Pt(size)
    for i, t in enumerate(texts):
        if i:
            d.add_section(WD_SECTION.NEW_PAGE)
        d.add_paragraph(t)
    d.save(str(path))


def _body(path):
    return etree.fromstring(zipfile.ZipFile(path).read("word/document.xml")).find(W + "body")


def test_sections_and_marks(tmp_path):
    src, dst = tmp_path / "a.docx", tmp_path / "m.docx"
    _doc(src, ["Első lap.", "Második lap.", "Harmadik lap."])
    assert len(F._sections(_body(src))) == 3
    assert F._marked(src, dst) == 3
    text = "".join(_body(dst).itertext())
    assert all(F.MARK.format(i) in text for i in range(3))


def test_shrink_scales_style_size_and_keeps_rpr_order(tmp_path):
    src, dst = tmp_path / "a.docx", tmp_path / "s.docx"
    _doc(src, ["Első lap.", "Második lap."], size=11)
    d = Document(str(src))
    [q for q in d.paragraphs if q.text][1].runs[0].underline = True      # w:u после кегля по схеме
    d.save(str(src))
    F._shrink(src, dst, {1: 0.9})
    p0, p1 = [p for p in _body(dst).iter(W + "p") if "".join(p.itertext()).strip()][:2]
    assert p0.find(f".//{W}sz") is None                                       # раздел 0 не тронут
    rpr = p1.find(f"{W}r/{W}rPr")
    assert rpr.find(W + "sz").get(W + "val") == "20"                          # 22 полупункта × 0,9
    names = [etree.QName(c).localname for c in rpr]
    assert names.index("sz") < names.index("szCs") < names.index("u")


def test_fit_shrinks_only_overflowing_sections(tmp_path):
    """Поддельная печать: раздел 1 перевода без уменьшения кегля занимает два листа, остальные — по одному."""
    orig, ru = tmp_path / "o.docx", tmp_path / "r.docx"
    _doc(orig, ["A", "B", "C"])
    _doc(ru, ["А", "Б длинный", "В"])

    def printer(docx, pdf):
        body = _body(docx)
        pdfd = fitz.open()
        for i, sec in enumerate(F._sections(body)):
            szs = [int(e.get(W + "val")) for e in sec[0].iter(W + "sz") if e.get(W + "val") != "2"]
            pages = 2 if docx.name.startswith("ru_m") and i == 1 and not (szs and min(szs) < 22) else 1
            for k in range(pages):
                pg = pdfd.new_page()
                pg.insert_text((72, 120), (F.MARK.format(i) + " ") if k == 0 else "folytatás", fontsize=10)
                pg.insert_text((72, 700 if k == 0 else 150), "szöveg", fontsize=10)
        pdfd.save(str(pdf))
    info = F.fit(orig, ru, printer, tmp_path / "w")
    assert info == {"sections": 3, "over": 1, "fixed": 1, "left": 0}
    sz = [e.get(W + "val") for e in _body(ru).iter(W + "sz")]
    assert sz and all(int(v) < 22 for v in sz)

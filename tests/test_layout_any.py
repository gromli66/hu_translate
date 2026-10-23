# Вёрстка без языка (22.10): лист-плакат и таблица-график — с переведённого PDF целиком; альбомные листы — альбомными
# разделами; стенд tools/layout_any.py (поддельный перевод любой длины).
import fitz
from docx import Document

import docx_build as B
import docx_layout as DL

from test_struct_docx import LAYOUT, PROFILE, p, png


def _seg(i, x0, y0, x1, text="címke"):
    return {"id": i, "page": 0, "text": text, "loc": {"bbox": [x0, y0, x1, y0 + 10]}}


def test_graphic_page_poster_and_fragmented_art():
    page = fitz.open().new_page(width=842, height=595)
    big = [(fitz.Rect(20, 20, 800, 560), {})]
    labels = [_seg(i, 100 + 40 * i, 200, 130 + 40 * i) for i in range(10)]
    assert DL._graphic_page(page, labels, big)                                  # плакат: схема с подписями
    pieces = [(fitz.Rect(20 + 90 * k, 20, 100 + 90 * k, 100), {}) for k in range(6)]
    assert DL._graphic_page(page, labels[:4], pieces)                           # картинка обрывками, текста мало
    prose = [_seg(i, 60, 100 + 14 * i, 780, "szöveg " * 20) for i in range(30)]
    shot = [(fitz.Rect(60, 300, 500, 560), {})]
    assert not DL._graphic_page(page, prose, shot)                              # абзацы и один снимок экрана


def test_chart_table_needs_shapes_inside_not_cell_fills():
    xs, ys = [40, 200, 300, 400, 500], [100, 120, 140, 160, 180]
    t = (xs, ys, [], [])
    bars = [{"rect": (210 + 20 * k, 103 + 20 * k, 290 + 20 * k, 112 + 20 * k), "solid": True} for k in range(3)]
    fills = [{"rect": (40, 100, 200, 120), "solid": True}, {"rect": (200, 120, 300, 140), "solid": True},
             {"rect": (300, 140, 400, 160), "solid": True}]
    assert DL._chart_table(t, bars) and not DL._chart_table(t, fills)


def test_landscape_pages_get_landscape_sections(tmp_path):
    blocks = [p("Első", 1, 100), p("Második", 2, 100), p("Harmadik", 3, 100)]
    mem = B.Memory([("Első", "Первый"), ("Második", "Второй"), ("Harmadik", "Третий")])
    L = dict(LAYOUT, first_is_body=True, page_sizes=[(595.3, 841.9), (841.9, 595.3), (595.3, 841.9)])
    out = tmp_path / "o.docx"
    B.Builder(mem, lambda img: png(), PROFILE).build(blocks, str(out), L)
    secs = Document(str(out)).sections
    assert [s.page_width > s.page_height for s in secs] == [False, True, False]


def test_word_export_recognizes_pdfmaker():
    """PDF из Word через Acrobat PDFMaker — экспорт Word (разметка по его соглашениям), как «Microsoft Word»."""
    assert B.word_export({"creator": "Acrobat PDFMaker 5.0 for Word", "producer": "Acrobat Distiller 5.0"})
    assert B.word_export({"producer": "Microsoft® Office Word 2007"})
    assert not B.word_export({"creator": "PDF24 Creator", "producer": "GPL Ghostscript"})


def test_builder_survives_control_chars_wide_logos_and_bad_margins(tmp_path):
    """Сборка не падает: управляющие знаки в тексте (NUL, FF), логотипы шапки шире полосы набора, подвал ниже листа
    (отрицательное поле) — №6, №27, №40, №48 партии 2 (22.10)."""
    blocks = [p("Első" + chr(12) + " sor" + chr(1), 2, 100)]
    mem = B.Memory([("Első sor", "Первая строка")])
    logo = {"xref": 1, "x0": 0, "x1": 595, "top": 10, "bottom": 60}
    L = dict(LAYOUT, first_is_body=True, head_imgs=[logo, dict(logo, x0=300, x1=595)], bottom=-200, footer=-5)
    out = tmp_path / "x.docx"
    B.Builder(mem, lambda img: png(), PROFILE).build(blocks, str(out), L)
    sec = Document(str(out)).sections[-1]
    assert sec.bottom_margin >= 0 and chr(12) not in "".join(par.text for par in Document(str(out)).paragraphs)


def test_upright_copy_of_rotated_pages(tmp_path):
    """Лист с /Rotate 90 без разметки — копия, где слова идут строкой слева направо в координатах видимого листа (HAKSER:
    таблицы альбомных листов собирались узкими столбцами, E48)."""
    import segments as S
    src = tmp_path / "rot.pdf"
    d = fitz.open()
    pg = d.new_page(width=612, height=792)
    pg.insert_text((100, 100), "Mérés eredménye", fontsize=11, rotate=270)
    pg.insert_text((100, 720), "Helyszín", fontsize=11, rotate=270)     # у дальнего края: show_pdf_page его срезала
    pg.set_rotation(90)
    d.save(str(src))
    up = S.upright(src, tmp_path / "up")
    assert up != src
    page = fitz.open(str(up))[0]
    w = page.get_text("words")
    x0, y0, x1, y1 = w[0][:4]
    assert x1 - x0 > y1 - y0                                    # слово лежит строкой, а не столбиком
    # все слова — там же, где их видно на повёрнутом листе
    seen = sorted((t[4], round((fitz.Rect(t[:4]) * pg.rotation_matrix).x0)) for t in fitz.open(str(src))[0].get_text("words"))
    assert page.rect == fitz.Rect(0, 0, 792, 612)
    assert sorted((t[4], round(t[0])) for t in w) == seen
    plain = tmp_path / "plain.pdf"
    fitz.open().new_page().insert_text((50, 50), "a")
    d2 = fitz.open(); d2.new_page().insert_text((50, 50), "szöveg"); d2.save(str(plain))
    assert S.upright(plain, tmp_path / "up") == plain           # без поворотов — исходный файл


def test_ocr_layer_only_invisible_text_over_scan():
    """Скан с невидимым слоем распознавания — не плакат; живой текст поверх фона-картинки (брошюра) — может быть."""
    import io
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (60, 80), "lightblue").save(buf, "PNG")
    doc = fitz.open()
    for mode in (3, 0):
        pg = doc.new_page()
        pg.insert_image(pg.rect, stream=buf.getvalue())
        pg.insert_text((72, 100), "Szöveg a háttér felett", render_mode=mode)
    assert DL._ocr_layer(doc[0]) and not DL._ocr_layer(doc[1])


def test_nested_table_in_cell_keeps_its_text(tmp_path):
    """Таблица в ячейке таблицы (листовка Word) — вложенной таблицей, а не пропуском вместе с текстом."""
    inner = {"k": "table", "rows": [[{"blocks": [p("Belső szöveg")], "th": False}],
                                    [{"blocks": [p("Második sor")], "th": False}]]}
    outer = {"k": "table", "rows": [[{"blocks": [p("Külső"), inner, p("Utána")], "th": False},
                                     {"blocks": [p("Jobb")], "th": False}]]}
    mem = B.Memory([("Belső szöveg", "Внутренний текст"), ("Második sor", "Вторая строка"), ("Külső", "Внешний"),
                    ("Utána", "После"), ("Jobb", "Справа")])
    out = tmp_path / "n.docx"
    B.Builder(mem, lambda img: png(), PROFILE).build([outer], str(out), LAYOUT)
    doc = Document(str(out))
    cell = doc.tables[0].cell(0, 0)
    assert [c.text for r in cell.tables[0].rows for c in r.cells] == ["Внутренний текст", "Вторая строка"]
    texts = [q.text for q in cell.paragraphs if q.text]
    assert texts == ["Внешний", "После"] and doc.tables[0].cell(0, 1).text == "Справа"


def test_paragraph_styled_like_chapter_heading_starts_new_page():
    """E41: абзац с видом заголовков 1-го уровня (те — с нового листа), первым на листе на их высоте, — с нового листа;
    обычный абзац первым на листе и такой же жирный ниже по листу — нет."""
    import docx_struct as D

    def blk(k, text, page, top, size, bold, **kw):
        return dict({"k": k, "runs": [{"t": text, "size": size, "b": bold}],
                     "geom": {"page": page, "top": top, "bottom": top + 14, "x0": 70}}, **kw)
    blocks = [blk("h", "1. BEVEZETÉS", 2, 104, 14, True, level=1), blk("p", "Szöveg.", 2, 140, 12, False),
              blk("h", "2. FOGALMAK", 3, 104, 14, True, level=1), blk("p", "Szöveg.", 3, 140, 12, False),
              blk("p", "ELŐSZÓ", 4, 104, 14, True), blk("p", "Folytatás a lap tetején", 5, 104, 12, False),
              blk("p", "Kiemelt sor", 5, 300, 14, True)]
    tops = {2: 104, 3: 104, 4: 104, 5: 104}
    bots = {1: 760, 2: 760, 3: 760, 4: 760, 5: 760}
    D.mark_breaks(blocks, tops, bots)
    for b in blocks[:4]:
        b["break"] = b["k"] == "h"                    # заголовки 1-го уровня — с нового листа (как стилем у Word)
    D.mark_breaks(blocks, tops, bots)
    assert [bool(b.get("break")) for b in blocks[4:]] == [True, False, False]


def test_toc_first_on_page_after_short_page_starts_new_page(tmp_path):
    """E41: оглавление первым на листе после листа, занятого на четверть, — с нового листа (и в DOCX)."""
    import docx_struct as D
    body = p("Előszó szövege.", 5, 100)
    body["geom"]["bottom"] = 250
    toc = {"k": "toc", "info": {"page": 6, "levels": 2}, "geom": {"page": 6}}
    blocks = [body, toc]
    D.mark_breaks(blocks, {5: 100}, {1: 760, 2: 760, 3: 760, 4: 760, 5: 250})
    assert toc.get("break")
    out = tmp_path / "t.docx"
    B.Builder(B.Memory([("Előszó szövege.", "Текст предисловия.")]), lambda img: png(), PROFILE).build(
        blocks, str(out), LAYOUT)
    doc = Document(str(out))
    head = next(q for q in doc.paragraphs if q.style.name.startswith("TOC Title"))
    assert head.paragraph_format.page_break_before


def test_header_logo_is_not_a_screenshot():
    """E42: небольшая картинка у верха листа (логотип шапки) — не скриншот для перевода надписей; та же в теле — да."""
    import io
    from PIL import Image
    import shots as SH
    buf = io.BytesIO()
    Image.new("RGB", (240, 150), "navy").save(buf, "PNG")
    doc = fitz.open()
    pg = doc.new_page(width=842, height=595)
    pg.insert_image(fitz.Rect(61, 54, 169, 120), stream=buf.getvalue())
    pg2 = doc.new_page(width=842, height=595)
    pg2.insert_image(fitz.Rect(61, 300, 169, 366), stream=buf.getvalue())
    got = SH.find_shots(doc)
    assert [n for places in got.values() for n, _ in places] == [1]


def test_upright_sideways_text_without_rotate(tmp_path):
    """Таблица, напечатанная «набок» на книжном листе без /Rotate (Danfoss л. 155–159), — выпрямляется, как /Rotate:
    повёрнутый текст сегменты пропускают, и лист терялся целиком."""
    import segments as S
    src = tmp_path / "side.pdf"
    d = fitz.open()
    pg = d.new_page(width=595, height=842)
    for k in range(12):
        pg.insert_text((80 + 30 * k, 700), f"6-{30 + k} X30/11-es csatlakozó, alsó feszültség", fontsize=9, rotate=90)
    d.new_page(width=595, height=842).insert_text((72, 72), "Normál szöveg " * 20, fontsize=9)
    d.save(str(src))
    up = S.upright(src, tmp_path / "up")
    assert up != src
    out = fitz.open(str(up))
    assert out[0].rect.width > out[0].rect.height and out[1].rect.width < out[1].rect.height
    w = out[0].get_text("words")
    assert len(w) == len(d[0].get_text("words")) and all(t[2] - t[0] > t[3] - t[1] for t in w if len(t[4]) > 3)


def test_review_xlsx_survives_control_chars(tmp_path):
    """Управляющий знак в тексте PDF (брошюра HERCA) — не роняет таблицу вычитки (а с ней сборку задачи)."""
    import assemble as A
    from openpyxl import load_workbook
    segs = [{"id": 0, "page": 0, "text": "Elkerülhető\x03 vizsgálat", "loc": {}}]
    res = {"0": {"ru": "Можно\x0b избежать", "issues": [], "tries": 1, "gloss_miss": []}}
    out = tmp_path / "v.xlsx"
    A.review_xlsx(out, segs, res, "t")
    row = [c.value for c in load_workbook(out).active[2]]
    assert row[2] == "Elkerülhető vizsgálat" and row[3] == "Можно избежать"


def test_text_layer_figure_with_text_inside_comes_from_translated_pdf(tmp_path):
    """Текстовый слой: рисунок, внутри которого текст листа (цитата поверх фото), — вырезкой с переведённого PDF
    (clip), а не картинкой оригинала; рисунок без текста — как был."""
    import io
    from PIL import Image
    import docx_layout as DL
    import segments as S
    import numpy as np
    buf = io.BytesIO()                              # пёстрая картинка (фото): сплошную заливку _figures берёт за линии
    Image.fromarray(((np.indices((200, 300)) // 20).sum(0) % 2 * 255)
                    .astype(np.uint8)).convert("RGB").save(buf, "PNG")
    d = fitz.open()
    pg = d.new_page()
    pg.insert_image(fitz.Rect(72, 300, 372, 500), stream=buf.getvalue())
    pg.insert_text((90, 420), "Ha már rendelkezésre áll a diagnózis", fontsize=12, color=(1, 1, 1))
    pg.insert_image(fitz.Rect(72, 560, 372, 760), stream=buf.getvalue())
    for k in range(6):
        pg.insert_text((72, 80 + 20 * k), f"A klinikai vizsgálatok eredménye {k}. bekezdés szövege.", fontsize=11)
    src = tmp_path / "f.pdf"
    d.save(str(src))
    segs, _ = S.pdf_segments(str(src), ocr=False, shots=False)
    blocks, _ = DL.load(str(src), segs, text=True)
    imgs = [r["img"] for b in blocks if b["k"] == "p" for r in b.get("runs") or [] if "img" in r]
    assert sorted("clip" in i for i in imgs) == [False, True]


def test_textless_copy_keeps_pictures_drops_text():
    """Копия листа без текста сегментов: картинка цела, текст в рамках снят, вне рамок — остался."""
    import io
    import numpy as np
    from PIL import Image
    import docx_layout as DL
    buf = io.BytesIO()
    Image.fromarray(((np.indices((100, 100)) // 20).sum(0) % 2 * 255).astype(np.uint8)).convert("RGB").save(buf, "PNG")
    d = fitz.open()
    pg = d.new_page()
    pg.insert_image(fitz.Rect(72, 300, 272, 500), stream=buf.getvalue())
    pg.insert_text((80, 400), "akkor nincs szükség", fontsize=12)
    pg.insert_text((80, 100), "Fejléc marad", fontsize=12)
    c = DL._textless(d, 0, [(78, 388, 300, 404)])
    assert c[0].get_text().strip() == "Fejléc marad" and len(c[0].get_images()) == 1


def test_on_lines_rect_is_text_not_picture():
    """Рамка почти целиком на строке сегмента — строка текста (растр с «запечённым» текстом под ней), не рисунок."""
    import docx_layout as DL
    seg = {"loc": {"bbox": [65, 452, 554, 532], "frags": [[65, 452, 554, 480], [84, 504, 535, 532]]}}
    assert DL._on_lines(fitz.Rect(83, 508, 535, 532), [seg])
    assert not DL._on_lines(fitz.Rect(32, 439, 73, 538), [seg])


def test_split_columns_of_borderless_table_not_justified_text():
    """E43: «Dátum ⇥ Határozat száma» (просвет 27 pt при прочих 3 pt) — два куска; выключка по ширине (просветы
    поровну), номер пункта и номер листа у просвета — один кусок."""
    import layout as LY
    R = fitz.Rect

    def frag(words):
        t = " ".join(w for _, w in words)
        return {"t": t, "bbox": list(R(words[0][0]) | R(words[-1][0])), "chars": [(len(t), 11, False, "serif")]}
    cols = [((71, 177, 109, 192), "Dátum"), ((136, 177, 191, 192), "Határozat"), ((194, 177, 229, 192), "száma")]
    just = [((70, 300, 110, 312), "Az"), ((130, 300, 170, 312), "MVM"), ((190, 300, 240, 312), "Paksi")]
    num = [((70, 400, 90, 412), "1.2"), ((120, 400, 170, 412), "Bevezető"), ((173, 400, 220, 412), "rész")]
    words = [(R(b), w) for b, w in cols + just + num]
    out = LY._split_columns([frag(cols), frag(just), frag(num)], words)
    assert [f["t"] for f in out] == ["Dátum", "Határozat száma", "Az MVM Paksi", "1.2 Bevezető rész"]


def test_two_text_columns_read_column_by_column_borderless_table_not():
    """E46: две колонки сплошного текста — просвет найден, абзацы колонка за колонкой; таблица без линий (абзацы слева
    и справа на одной высоте) — не колонки."""
    import layout as LY

    def par(x0, y0, x1, y1, t):
        return {"bbox": [x0, y0, x1, y1], "frags": [{"bbox": [x0, y0, x1, y1], "t": t}], "text": t}
    head = par(57, 40, 540, 52, "Fejezet címe")
    left = [par(57, 80 + 60 * k, 290, 120 + 60 * k, f"bal {k}") for k in range(4)]
    right = [par(310, 95 + 60 * k, 540, 135 + 60 * k, f"jobb {k}") for k in range(4)]
    pars = [head] + [p for pair in zip(left, right) for p in pair]
    g = LY.gutter(pars, 595)
    assert g is not None and 290 < g < 310
    assert [p["text"] for p in LY.column_order(pars, g)] == ["Fejezet címe"] + [f"bal {k}" for k in range(4)] + \
        [f"jobb {k}" for k in range(4)]
    table = [head] + [p for k in range(4) for p in (par(57, 80 + 30 * k, 200, 92 + 30 * k, f"a{k}"),
                                                    par(320, 80 + 30 * k, 500, 92 + 30 * k, f"b{k}"))]
    assert LY.gutter(table, 595) is None
    # список «сокращение ⇥ расшифровка» (u_T3): узкий столбец слева — таблица, а не колонка текста
    abbr = [head] + [par(57, 80 + 45 * k, 100, 92 + 45 * k, f"LA{k}") for k in range(4)] +         [par(160, 86 + 45 * k, 540, 120 + 45 * k, f"rendszer {k}") for k in range(4)]
    assert LY.gutter(abbr, 595) is None


def test_cut_span_pieces_do_not_rejoin_across_column_line(tmp_path):
    """Строка PDF через линии столбцов («47,05281 19,11048 Kunszentmiklós», HAKSER) — тремя кусками: рамка куска
    без пробела у линии, склейка видит линию."""
    import layout as LY
    d = fitz.open()
    pg = d.new_page(width=792, height=612)
    pg.insert_text((108, 150), "47,05281 19,11048 Kunszentmiklós", fontsize=10)
    xs = [fitz.get_text_length("47,05281 ", fontsize=10), fitz.get_text_length("47,05281 19,11048 ", fontsize=10)]
    vs = [(139, 165, 108 + x - 1.5) for x in xs]            # линия — у пробела, вплотную к началу следующего куска
    for _, _, x in vs:
        pg.draw_line((x, 139), (x, 165))
    fr = LY.frags_from_text(pg, vs)
    assert [f["t"] for f in fr] == ["47,05281", "19,11048", "Kunszentmiklós"]


def test_each_landscape_page_own_section_with_fit_want(tmp_path):
    """Каждый альбомный лист — своим разделом; листов оригинала у альбомных разделов — fit_want (docx_fit ужимает
    перелившийся на второй лист): таблица, идущая через два альбомных листа одним блоком, — раздел в 2 листа."""
    blocks = [p("Első", 1, 100), p("Második", 2, 100), p("Harmadik", 3, 100), p("Negyedik", 5, 100),
              p("Ötödik", 6, 100)]
    mem = B.Memory([("Első", "Первый"), ("Második", "Второй"), ("Harmadik", "Третий"), ("Negyedik", "Четвёртый"),
                    ("Ötödik", "Пятый")])
    land, port = (841.9, 595.3), (595.3, 841.9)
    L = dict(LAYOUT, first_is_body=True, page_sizes=[port, land, land, land, land, port])
    out = tmp_path / "o.docx"
    bld = B.Builder(mem, lambda img: png(), PROFILE)
    bld.build(blocks, str(out), L)
    secs = Document(str(out)).sections
    assert [s.page_width > s.page_height for s in secs] == [False, True, True, True, False]
    assert bld.fit_want == {1: 1, 2: 2, 3: 1}


def test_foot_top_text_layer_ignores_trailing_space():
    """Текстовый слой: пробел в конце строки таблицы выше низа тела — не подвал (таблица решений OAH: нижнее поле
    165 pt); у размеченного PDF (cbot None) — как было."""
    import docx_struct as D
    space = {"text": " ", "top": 443.0, "bottom": 455.0, "x0": 598.0, "x1": 601.0}
    assert D._foot_top([space], [], 595.0, 470.0) == 595.0
    assert D._foot_top([space], [], 595.0, None) < 450


def test_fit_words_picks_donor_with_fewest_new_lines():
    """Ширину длинному слову отдаёт столбец, у которого меньше прибавится строк, а не самый просторный по слову:
    «0,0042 (5)» остаётся в строку, отдаёт столбец с одиночными словами."""
    cell = lambda t: {"blocks": [p(t)], "th": False}
    rows = [[cell("гамма-спектрометрия"), cell("0,0042 (5)"), cell("Да")], [cell("x"), cell("< 0,46"), cell("Нет")]]
    mem = B.Memory([(t, t) for t in ("гамма-спектрометрия", "0,0042 (5)", "Да", "x", "< 0,46", "Нет")])
    bld = B.Builder(mem, lambda img: png(), PROFILE)
    widths = [60.0, 70.0, 64.0]
    bld._fit_words(rows, widths, 10)
    assert widths[1] == 70.0 and widths[2] < 64.0 and abs(sum(widths) - 194.0) < 0.01

# Сборка DOCX текстового слоя без разметки (аудит 29.10): листы одной сетки бланка — таблица, а не снимок листа;
# подписи строк вверху листа — тело; шапка у серии листов со своей версией — свой раздел.
import fitz
from docx import Document
from docx.oxml.ns import qn

import docx_build as B
import docx_layout as DL
import docx_struct as D

from test_struct_docx import LAYOUT, PROFILE, png
from test_text_forms import _segs


def _sheets_pdf(path, pages=5, head=None, labels=False, pieces=False):
    """Листы бланка с одной и той же сеткой: рамка во весь лист, внутри — две вертикали ниже и короче рамки (столбцы),
    линия под шапкой, подвал-рамка; шапка листа и строки в ячейках. head(pno) → (версия, лист); labels — подписи строк
    вверху таблицы («Berendezés alfanumerika» — на каждом листе одни и те же); pieces — вертикали столбцов рисованы по
    строкам: стопка отрезков в 11 pt, как рамку таблицы Word."""
    doc = fitz.open()
    for pno in range(pages):
        pg = doc.new_page(width=595, height=842)
        sh = pg.new_shape()
        sh.draw_rect(fitz.Rect(60, 90, 540, 700))
        sh.draw_line((60, 130), (540, 130))
        sh.draw_rect(fitz.Rect(60, 720, 540, 760))
        for x in (140, 440):
            if pieces:
                for k in range(54):
                    sh.draw_rect(fitz.Rect(x - 0.5, 130 + 10.5 * k, x + 0.5, 130 + 10.5 * k + 10))
            else:
                sh.draw_line((x, 130), (x, 700))
        sh.finish(color=(0, 0, 0), fill=None if not pieces else (0, 0, 0), width=0.6)
        sh.commit()
        put = lambda x, y, t: pg.insert_text((x, y), t, fontsize=9, fontname="tiro")
        ver, num = head(pno) if head else (0, pno + 1)
        put(62, 40, "Azonosító: 30XYZ")
        put(400, 40, f"Verziószám: {ver}")
        put(62, 56, "Megnevezés: Próba")
        put(400, 56, f"Oldalszám: {num}")
        put(64, 124, "Sorsz.")
        put(200, 124, "Művelet, ellenőrzés")
        put(450, 124, "Igazolás")
        y = 150
        if labels:
            put(64, y, "Berendezés alfanumerika")
            put(200, y, f"YA1{pno}S001")
            put(64, y + 14, "Beépítés helye")
            put(200, y + 14, f"A30{pno}")
            y += 40
        for k in range(14):
            put(64, y + 22 * k, f"{pno + 1}.{k + 1}.")
            put(200, y + 22 * k, f"A {pno + 1}. lap {k + 1}. lépése megtörtént.")
        put(300, 745, "A lapot kitöltötte")
    doc.save(str(path))
    return _segs(path)


def test_form_grid_repeated_on_every_sheet_is_a_table_not_a_sheet_snapshot(tmp_path):
    """Одна и та же сетка бланка на каждом листе (30PR132GT: 37 листов из 68; 30PR25GT) — таблица: её вертикали не рамка
    листа и не снимаются, а без таблицы рамка уходила снимком листа с переведённого PDF (_vector_figures)."""
    src = tmp_path / "sheets.pdf"
    segs = _sheets_pdf(src, pieces=True)
    blocks, _ = DL.load(str(src), segs, text=True)
    assert not [b for b in blocks if b.get("figure")]
    tabs = [b for b in blocks if b["k"] == "table"]
    assert tabs and any("lépése" in "".join(r.get("t", "") for c in row for x in c["blocks"] for r in x["runs"])
                        for t in tabs for row in t["rows"])
    d = fitz.open(str(src))
    hs, vs = DL._page_rules(d, True)[1][2]
    assert sorted(round(v[0]) for v in vs if v[2] - v[1] > 100) == [60, 140, 440, 540]     # столбцы не сняты как рамка


def test_piece_of_long_rule_is_not_a_vector_shape():
    """Вертикаль таблицы Word — стопка отрезков по строкам (1×11 pt): _rules склеивает их в линию, а как «схема» они
    давали область листа снимком (3PR10: оглавление и тело листа с рамкой) — кусок длинной линии фигурой не считается."""
    rules = ([(100.0, 60.0, 540.0)], [(140.0, 130.0, 700.0)])
    assert DL._on_rule(fitz.Rect(139.5, 140, 140.5, 151) + (-0.5, -0.5, 0.5, 0.5), rules)
    assert DL._on_rule(fitz.Rect(70, 100, 90, 100) + (-0.5, -0.5, 0.5, 0.5), rules)
    assert not DL._on_rule(fitz.Rect(200, 140, 201, 151), rules)                    # не на линии
    assert not DL._on_rule(fitz.Rect(130, 140, 150, 160), rules)                    # рамка, не отрезок линии
    assert not DL._on_rule(fitz.Rect(139.5, 60, 140.5, 71), ([], [(140.0, 60.0, 80.0)]))   # линия короче 40 pt


def test_row_labels_at_sheet_top_stay_in_table_on_every_sheet(tmp_path):
    """«Berendezés alfanumerika», «Beépítés helye»: подписи строк таблицы вверху каждого листа (3PR10: 9 листов) — тело,
    не «название раздела, повторяющее верх прошлого листа»: в переводе оставались на одном листе из девяти."""
    src = tmp_path / "labels.pdf"
    segs = _sheets_pdf(src, labels=True)
    blocks, _ = DL.load(str(src), segs, text=True)
    cells = " ".join("".join(r.get("t", "") for r in x["runs"]) for b in blocks if b["k"] == "table"
                     for row in b["rows"] for c in row for x in c["blocks"])
    assert cells.count("Berendezés alfanumerika") == 5 and cells.count("Beépítés helye") == 5


def test_sheets_with_other_version_get_own_header_section(tmp_path):
    """Версия «1» на одних листах и «0» на других (3PR10, 3SZ17, 30PR25GT, 30RVR*: шапка одна с образца) — серия листов
    со своей шапкой и подвалом, номер листа — полем и в шапке серии; шапка тела — с самого частого листа."""
    src = tmp_path / "versions.pdf"
    ver = {0: 0, 1: 0, 2: 0, 3: 1, 4: 1, 5: 0, 6: 0}
    segs = _sheets_pdf(src, pages=7, head=lambda pno: (ver[pno], pno + 1))
    blocks, info = DL.load(str(src), segs, text=True)
    L = D.layout(str(src), blocks, content=info.get("content_boxes"))
    blocks = D.split_at_zones(blocks, L["variants"])                    # таблица через лист режется по сериям
    assert [(v["from"], v["to"]) for v in L["variants"]] == [(4, 5)]
    tabs = [b for b in blocks if b["k"] == "table"]
    assert [b["geom"]["page"] for b in tabs] == [1, 4, 6]
    v = L["variants"][0]
    assert v["page_marked"] and "Verziószám: 1" in str(v["head_grid"]).replace("", "")
    assert "Verziószám: 0" in str(L["head_grid"]).replace("", "") and L["page_marked"]
    out = tmp_path / "versions.docx"
    B.Builder(B.Memory([(s["text"], s["text"]) for s in segs]), lambda img: png(), PROFILE).build(
        blocks, str(out), dict(LAYOUT, **{k: v for k, v in L.items() if k in (
            "head_grid", "variants", "page_marked", "page_tok", "stamp", "stamp_widths", "stamp_box", "stamp_inner",
            "stamp_paras", "stamp_aligns", "stamp_merge", "head_rule", "head_imgs", "head_lines", "first_is_body")}))
    doc = Document(str(out))
    heads = ["".join(t.text for t in s.header._element.iter(qn("w:t"))) for s in doc.sections]
    assert len(heads) >= 3 and any("Verziószám: 1" in h for h in heads) and "Verziószám: 0" in heads[0]
    assert all("PAGE" in sec.header._element.xml for sec in doc.sections)       # номер листа — полем в каждой шапке


def test_text_lost_in_docx_is_reported_with_its_count(tmp_path):
    """Инвариант сборки (аудит 29.10): перевод сегмента есть в DOCX. Потерян целиком — в списке; вошёл на одном листе из
    двух — в списке с числом потерянных; колонтитул на пяти листах — один раз в колонтитуле, не потеря; листы-снимки и
    надписи на картинках не считаются."""
    import assemble as A
    seg = lambda i, page, text, y=300.0, **loc: {"id": i, "page": page, "text": text,
                                                 "loc": dict({"bbox": [60.0, y, 200.0, y + 10]}, **loc)}
    segs = [seg(0, 0, "Állapot ellenőrzése"), seg(1, 1, "Állapot ellenőrzése"),         # две таблицы — один перевод
            seg(2, 0, "Berendezés alfanumerika"),                                      # потерян
            seg(3, 2, "Beépítés helye megjelölése"),                                   # лист-снимок
            seg(4, 0, "Állapot ellenőrzése", shot={"xref": 7})]                        # надпись на картинке
    segs += [seg(10 + k, k, "Azonosító: Próba azonosítója", y=30.0) for k in range(5)]  # колонтитул ×5
    ru = {"Állapot ellenőrzése": "Проверка состояния", "Berendezés alfanumerika": "Буквенный код оборудования",
          "Beépítés helye megjelölése": "Место установки обозначено", "Azonosító: Próba azonosítója": "Идентификатор пробы"}
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Идентификатор пробы"
    doc.add_paragraph("Проверка состояния")
    out = tmp_path / "lost.docx"
    doc.save(str(out))
    got = A.lost_segments(str(out), segs, lambda s: ru[s["text"]], skip_pages={2})
    assert sorted((s["id"], n) for s, _, n in got) == [(0, 1), (2, 1)]


def test_header_cell_glued_from_two_segments_is_translated_by_parts():
    """Ячейка шапки у листа с другой версией — «Azonosító: 3SZ23 Verziószám: 1» одной строкой, а сегментами листа — две
    подписи: перевод по частям (прежде оставалась по-венгерски — 3SZ23, лист 5); неизвестное и чужие склейки — не трогаем."""
    mem = B.Memory([("Azonosító: 3SZ23", "Идентификатор: 3SZ23"), ("Verziószám: 1", "Номер версии: 1"),
                    ("Oldalszám: 3", "Номер страницы: 3")])
    assert mem.get("Azonosító: 3SZ23 Verziószám: 1") == "Идентификатор: 3SZ23 Номер версии: 1"
    assert mem.get("Azonosító: 3SZ23") == "Идентификатор: 3SZ23"
    assert mem.get("Azonosító: 3SZ23 Verziószám: 2") is None and mem.get("Elsőszám: 3SZ23") is None


def test_page_label_of_other_sheet_is_translated_by_known_number():
    """Строка шапки листа серии — «Oldalszám: 2.»: в памяти есть только «Oldalszám: 13.» (и склейка «Verziószám: 1 Oldalszám:
    2.»); число ставит поле листа, перевод — тот же с другим числом. Число, которого в переводе нет, — не по образцу."""
    mem = B.Memory([("Oldalszám: 13.", "Номер страницы: 13."), ("Verziószám: 0", "Номер версии: 0"),
                    ("Dátum: 2017-04-20", "Дата: 20.04.2017")])
    assert mem.renumbered("Verziószám: 1") == "Номер версии: 1"
    assert mem.renumbered("Oldalszám: 2.") == "Номер страницы: 2." and mem.renumbered("Oldalszám: 112.") == "Номер страницы: 112."
    assert mem.renumbered("Dátum: 2018-05-21") is None and mem.renumbered("Megjegyzés: 2.") is None
    # «Oldalszám: 1 2» и «Oldalszám: 12» — один ключ без пробелов, а чисел разное число: не падать (30RL2126GT, 29.10)
    mem = B.Memory([("Oldalszám: 1 2", "Номер страницы: 1 2")])
    assert mem.renumbered("Oldalszám: 12") is None and mem.renumbered("Oldalszám: 3 4") is None
    # строка шапки — половина «подпись: значение» известного сегмента: подпись и значение в разных ячейках
    mem = B.Memory([("Azonosító: 30GF(246)51GT", "Идентификатор: 30GF(246)51GT"),
                    ("Megnevezés: tömörség ellenőrzése.", "Наименование: проверка плотности.")])
    assert [mem.zone(x) for x in ("Azonosító:", "Megnevezés:", "tömörség ellenőrzése.", "30GF(246)51GT", "Dátum:")] == [
        "Идентификатор:", "Наименование:", "проверка плотности.", "30GF(246)51GT", None]


def test_repeated_table_header_row_is_not_lost():
    """Строка таблицы с w:tblHeader печатается на каждом листе таблицы: в файле она одна, а сегментов шапки — по числу листов
    (30GF(246)51GT: «Feltétel, művelet, ellenőrzés» ×9; 30PR132GT: легенда ×49) — не потеря; без tblHeader — потеря."""
    import tempfile
    from docx.oxml import OxmlElement
    import assemble as A
    segs = [{"id": k, "page": k, "text": "Feltétel, művelet, ellenőrzés", "loc": {"bbox": [60.0, 300.0, 200.0, 310.0]}}
            for k in range(3)]
    for header in (True, False):
        doc = Document()
        t = doc.add_table(rows=2, cols=1)
        t.rows[0].cells[0].text = "Условие, операция, проверка"
        if header:
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        with tempfile.TemporaryDirectory() as tmp:
            doc.save(tmp + "/h.docx")
            got = A.lost_segments(tmp + "/h.docx", segs, lambda s: "Условие, операция, проверка")
        assert (not got) if header else [n for _, _, n in got] == [2]



def test_footer_phrase_wrapped_by_narrow_cell_is_translated_whole():
    """«A lapot kitöltötte» / «(név, dátum, aláírás):» — одна фраза в узкой ячейке штампа: по строкам её в памяти нет
    (30HU375GT: подвал листов оставался по-венгерски), целиком — есть; строки, известные по отдельности, не склеиваются."""
    mem = B.Memory([("A lapot kitöltötte (név, dátum, aláírás):", "Лист заполнил (Ф. И. О., дата, подпись):"),
                    ("Azonosító: 3SZ23", "Идентификатор: 3SZ23"), ("Verziószám: 1", "Номер версии: 1")])
    b = B.Builder(mem, lambda img: None, PROFILE)
    assert b._wrapped(["A lapot kitöltötte", "(név, dátum, aláírás):"]) == ["A lapot kitöltötte (név, dátum, aláírás):"]
    assert b._wrapped(["Azonosító: 3SZ23", "Verziószám: 1"]) == ["Azonosító: 3SZ23", "Verziószám: 1"]

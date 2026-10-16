# -*- coding: utf-8 -*-
"""Стенд синтетики с известным ответом: DOCX с заданной структурой → PDF через Word (размеченный, как у заказчика) →
сборка DOCX по разметке (assemble.build_struct_docx, без перевода) → сравнение собранного DOCX с исходным по признакам
оформления. Исходный DOCX и есть правильный ответ: шапка, подвал, нумерация листов, заголовки и их уровни, оглавление,
списки, таблицы (объединённые ячейки, линии), отступ первой строки, выключка, длина.

    python tools/synth_bench.py gen <папка> <число> [зерно]   — корпус DOCX + паспорта признаков
    python tools/synth_bench.py pdf <папка>                    — PDF через Word (PowerShell + COM, один сеанс)
    python tools/synth_bench.py build <папка>                  — DOCX по разметке PDF
    python tools/synth_bench.py score <папка>                  — оценки по признакам: score.json и score.md

Зачем: правило, починенное по одному документу, ломалось на следующем (16.10: пять чужих PDF из Word вскрыли шапку
с логотипом, уровни оглавления, плавающие надписи). Стенд меряет всё сразу и ловит, что правка сломала."""
import json
import random
import re
import subprocess
import sys
from collections import Counter
from difflib import SequenceMatcher
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

WORDS = ("система контроль регулятор давление турбина насос клапан режим работа персонал проверка значение уставка "
         "сигнал защита блок агрегат подача вода пар температура уровень расход датчик измерение канал шкаф панель "
         "оператор управление испытание выполнение операция порядок документ требование условие параметр нагрузка "
         "мощность частота вращение отключение включение пуск останов автоматический ручной основной резервный "
         "технический эксплуатация обслуживание состояние положение срабатывание выдержка время задание отклонение "
         "необходимо следует допускается выполняется производится обеспечивается контролируется проверяется").split()


def _sent(rng, lo=6, hi=16):
    s = " ".join(rng.choice(WORDS) for _ in range(rng.randint(lo, hi)))
    return s[0].upper() + s[1:] + "."


def _para(rng, lo=1, hi=4):
    return " ".join(_sent(rng) for _ in range(rng.randint(lo, hi)))


def _title(rng):
    s = " ".join(rng.choice(WORDS) for _ in range(rng.randint(2, 5)))
    return s[0].upper() + s[1:]


def _logo_png(path):
    import fitz
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 160, 50), False)
    pix.set_rect(pix.irect, (255, 255, 255))
    pix.set_rect(fitz.IRect(4, 4, 46, 46), (230, 110, 20))
    pix.set_rect(fitz.IRect(56, 14, 156, 36), (20, 40, 120))
    pix.save(str(path))


# ---------------------------------------------------------------- генератор
def gen_one(seed, out):
    """Документ с признаками по зерну; паспорт — что в нём должно восстановиться."""
    from docx import Document
    from docx.shared import Pt, Mm
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
    from docx.enum.section import WD_SECTION
    from docx_build import _el, _field, W
    rng = random.Random(seed)
    P = {"seed": seed, "font": rng.choice(["Times New Roman", "Arial", "Calibri"]), "size": rng.choice([10, 11, 12]),
         "justify": rng.random() < 0.6, "first_line": rng.choice([0, 0, 18, 36]), "after": rng.choice([0, 6]),
         "header": rng.choice(["none", "text", "table_logo"]), "footer": rng.choice(["page", "stamp_grid", "stamp_box"]),
         "title": rng.random() < 0.6, "heads": rng.choice(["text", "auto", "none"]), "depth": rng.choice([2, 3, 4])}
    P["toc"] = P["heads"] != "none" and rng.random() < 0.75
    P["toc_levels"] = min(P["depth"], rng.choice([2, 3]))
    P["front"] = P["title"] and P["toc"] and rng.random() < 0.6
    doc = Document()
    st = doc.styles
    st["Normal"].font.name, st["Normal"].font.size = P["font"], Pt(P["size"])
    pf = st["Normal"].paragraph_format
    pf.space_after, pf.first_line_indent = Pt(P["after"]), Pt(P["first_line"]) if P["first_line"] else None
    pf.alignment = AL.JUSTIFY if P["justify"] else AL.LEFT
    for lv in range(1, 5):
        h = st[f"Heading {lv}"]
        h.font.name, h.font.size, h.font.bold = P["font"], Pt(P["size"] + max(0, 4 - lv)), lv <= 2
        h.font.color.rgb = None
        h.paragraph_format.first_line_indent = Pt(0)
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(12), Pt(6)
    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2", "Caption"):
        st[name].paragraph_format.first_line_indent = Pt(0)
    if P["heads"] == "auto":                    # автонумерация заголовков — многоуровневый список стилей
        num = doc.part.numbering_part.element
        an = _el("w:abstractNum", w__abstractNumId=880)
        an.append(_el("w:multiLevelType", w__val="multilevel"))
        for lv in range(4):
            l = _el("w:lvl", w__ilvl=lv)
            l.append(_el("w:start", w__val=1)); l.append(_el("w:numFmt", w__val="decimal"))
            l.append(_el("w:pStyle", w__val=f"Heading{lv + 1}"))
            l.append(_el("w:lvlText", w__val=".".join(f"%{i + 1}" for i in range(lv + 1)) + "."))
            l.append(_el("w:lvlJc", w__val="left"))
            ppr = _el("w:pPr"); ppr.append(_el("w:ind", w__left=720, w__hanging=720)); l.append(ppr)
            an.append(l)
        first = num.find(W + "num")
        (first.addprevious if first is not None else num.append)(an)
        n = _el("w:num", w__numId=880); n.append(_el("w:abstractNumId", w__val=880)); num.append(n)
        for lv in range(1, 5):
            ppr = st[f"Heading {lv}"].element.get_or_add_pPr()
            npr = _el("w:numPr"); npr.append(_el("w:ilvl", w__val=lv - 1)); npr.append(_el("w:numId", w__val=880))
            ppr.insert(0, npr)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(rng.choice([20, 25]))
    logo = out.with_suffix(".logo.png")
    _logo_png(logo)
    code = f"ИН-{rng.randint(100, 999)}-{rng.randint(1000, 9999)}"
    title = _title(rng)

    def borders(t, inner=True):
        b = _el("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            on = inner or side in ("top", "left", "bottom", "right")
            b.append(_el(f"w:{side}", w__val="single" if on else "nil", w__sz=4, w__space=0, w__color="000000"))
        t._tbl.tblPr.append(b)

    def zone(part, kind):
        part.is_linked_to_previous = False
        p0 = part.paragraphs[0]
        if kind == "none":
            return
        if kind == "text":
            p0.text = f"{title} — {code}"
            p0.alignment = AL.CENTER
            return
        if kind == "page":
            p0.alignment = AL.CENTER
            p0.add_run("Лист ")
            _field(p0, "PAGE", "1")
            return
        if kind == "table_logo":                # шапка-таблица: логотип | номер листа и название | код и гриф
            t = part.add_table(rows=2, cols=3, width=Mm(170))
            t.style = doc.styles["Table Grid"]
            for c, w in zip(t.columns, (40, 95, 35)):
                for cell in c.cells:
                    cell.width = Mm(w)
            t.cell(0, 0).merge(t.cell(1, 0)).paragraphs[0].add_run().add_picture(str(logo), width=Mm(34))
            mid = t.cell(0, 1).merge(t.cell(1, 1))
            mid.paragraphs[0].alignment = AL.CENTER
            _field(mid.paragraphs[0], "PAGE", "1")
            q = mid.add_paragraph(title); q.alignment = AL.CENTER
            t.cell(0, 2).paragraphs[0].text = code
            t.cell(1, 2).paragraphs[0].text = "Для служебного пользования"
            p0._p.addprevious(t._tbl)
            return
        t = part.add_table(rows=2, cols=3, width=Mm(170))  # штамп: подписи и значения, номер листа в углу
        borders(t, inner=kind == "stamp_grid")
        for c, w in zip(t.columns, (45, 90, 35)):
            for cell in c.cells:
                cell.width = Mm(w)
        t.cell(0, 0).paragraphs[0].text = "Объект:"
        t.cell(0, 1).paragraphs[0].text = "Наименование:"
        t.cell(0, 2).paragraphs[0].text = "Лист:"
        t.cell(1, 0).paragraphs[0].text = "Блок 3"
        t.cell(1, 1).paragraphs[0].text = title
        q = t.cell(1, 2).paragraphs[0]
        q.alignment = AL.RIGHT
        _field(q, "PAGE", "1")
        p0._p.addprevious(t._tbl)

    if P["title"]:
        for _ in range(6):
            doc.add_paragraph("")
        for t in (title.upper(), code, f"Редакция {rng.randint(1, 9)}"):
            q = doc.add_paragraph(t); q.alignment = AL.CENTER
            q.paragraph_format.first_line_indent = Pt(0)
            for r in q.runs:
                r.bold, r.font.size = True, Pt(P["size"] + 6)
        nsec = doc.add_section(WD_SECTION.NEW_PAGE)
    else:
        nsec = sec
    if P["toc"]:
        q = doc.add_paragraph("Содержание"); q.paragraph_format.first_line_indent = Pt(0)
        q.runs[0].bold = True
        _field(doc.add_paragraph(), f'TOC \\o "1-{P["toc_levels"]}" \\h \\z \\u', "[оглавление]")
        if P["front"]:
            nsec = doc.add_section(WD_SECTION.NEW_PAGE)
            P["front_sections"] = True
    heads, counters = [], [0] * 5
    nchap = rng.randint(2, 3)
    feats = Counter()

    def heading(lv, text):
        counters[lv - 1] += 1
        for i in range(lv, 5):
            counters[i] = 0
        num = ".".join(str(c) for c in counters[:lv]) + "."
        t = f"{num} {text}" if P["heads"] == "text" else text
        doc.add_heading(t, level=lv)
        heads.append((lv, text))

    def block():
        k = rng.choice(["p", "p", "p", "bullets", "numbers", "labels", "table", "figure"])
        if k == "p":
            doc.add_paragraph(_para(rng))
        elif k == "bullets":
            feats["bullets"] += 1
            for i in range(rng.randint(2, 4)):
                doc.add_paragraph(_sent(rng, 4, 10), style="List Bullet")
                if rng.random() < 0.3:
                    feats["nested"] += 1
                    doc.add_paragraph(_sent(rng, 3, 8), style="List Bullet 2")
        elif k == "numbers":
            feats["numbers"] += 1
            for i in range(rng.randint(2, 4)):
                doc.add_paragraph(_sent(rng, 4, 10), style="List Number")
        elif k == "labels":                     # метки прямо в тексте, как у ÜFK: «a.) …»
            feats["labels"] += 1
            for i in range(rng.randint(2, 4)):
                q = doc.add_paragraph(f"{'abcdef'[i]}.) {_sent(rng, 5, 12)}")
                q.paragraph_format.left_indent, q.paragraph_format.first_line_indent = Pt(36), Pt(-18)
        elif k == "table":
            kind = rng.choice(["grid", "merged", "hlines", "long"])
            feats["table_" + kind] += 1
            nrow = rng.randint(14, 24) if kind == "long" else rng.randint(3, 6)
            ncol = rng.randint(3, 4)
            t = doc.add_table(rows=nrow, cols=ncol)
            if kind == "hlines":
                b = _el("w:tblBorders")
                for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                    on = side in ("top", "bottom", "insideH")
                    b.append(_el(f"w:{side}", w__val="single" if on else "nil", w__sz=4, w__space=0, w__color="000000"))
                t._tbl.tblPr.append(b)
            else:
                t.style = doc.styles["Table Grid"]
            for c in range(ncol):
                t.cell(0, c).paragraphs[0].add_run(_title(rng)).bold = True
            if kind == "long":
                t.rows[0]._tr.get_or_add_trPr().append(_el("w:tblHeader"))
            for r in range(1, nrow):
                for c in range(ncol):
                    t.cell(r, c).paragraphs[0].text = _sent(rng, 1, 4) if c else f"{r}"
            if kind == "merged":
                a = t.cell(1, 1).merge(t.cell(1, 2))
                a.paragraphs[-1]._p.getparent().remove(a.paragraphs[-1]._p) if len(a.paragraphs) > 1 else None
            for cell in t._cells:
                for q in cell.paragraphs:
                    q.paragraph_format.first_line_indent = Pt(0)
            doc.add_paragraph("")
        else:
            feats["figure"] += 1
            q = doc.add_paragraph()
            q.alignment = AL.CENTER
            q.paragraph_format.first_line_indent = Pt(0)
            q.add_run().add_picture(str(logo), width=Mm(rng.choice([60, 90])))
            c = doc.add_paragraph(f"Рисунок {feats['figure']} — {_title(rng)}", style="Caption")
            c.alignment = AL.CENTER

    for _ in range(nchap):
        heading(1, _title(rng))
        doc.add_paragraph(_para(rng))
        for _ in range(rng.randint(2, 3)):
            heading(2, _title(rng))
            for _ in range(rng.randint(1, 3)):
                block()
            if P["depth"] >= 3:
                for _ in range(rng.randint(0, 2)):
                    heading(3, _title(rng))
                    block()
                    if P["depth"] >= 4 and rng.random() < 0.5:
                        heading(4, _title(rng))
                        block()
    # разделы: титул — без колонтитулов; вводная часть (оглавление) — свой подвал без номера; тело — нумерация с 1
    secs = doc.sections
    body = secs[-1]
    for s in secs:
        s.left_margin = s.right_margin = sec.left_margin
    if P["title"]:
        zone(secs[0].header, "none"); zone(secs[0].footer, "none")
    if P.get("front_sections"):
        front = secs[1]
        zone(front.header, P["header"])
        front.footer.is_linked_to_previous = False
        front.footer.paragraphs[0].text = "Содержание"
        front.footer.paragraphs[0].alignment = AL.CENTER
    zone(body.header, P["header"])
    zone(body.footer, P["footer"])
    if P["title"]:
        sp = body._sectPr
        el = _el("w:pgNumType", w__start=1)
        nxt = next((c for c in sp if c.tag in {W + n for n in ("cols", "titlePg", "docGrid")}), None)
        nxt.addprevious(el) if nxt is not None else sp.append(el)
        P["num_start"] = 1
    doc.save(str(out))
    P["heads_list"], P["feats"] = heads, dict(feats)
    out.with_suffix(".json").write_text(json.dumps(P, ensure_ascii=False, indent=1), encoding="utf-8")
    return P


# ---------------------------------------------------------------- печать через Word
PS = r'''param([string]$List)
$ErrorActionPreference = "Continue"
$w = New-Object -ComObject Word.Application
$w.Visible = $false
$w.DisplayAlerts = 0
foreach ($line in Get-Content -Encoding UTF8 $List) {
    $src, $dst = $line -split "`t"
    try {
        $d = $w.Documents.Open($src, $false, $true, $false)
        foreach ($t in $d.TablesOfContents) { $null = $t.Update() }
        $null = $d.Fields.Update()
        $d.Repaginate()
        # 17 — PDF; закладки по заголовкам, разметка структуры (DocStructureTags) — как экспорт у заказчика
        $d.ExportAsFixedFormat($dst, 17, $false, 0, 0, 1, 1, 0, $true, $true, 1, $true, $true, $false)
        $d.Close(0)
        Write-Output "ok`t$dst"
    } catch { Write-Output "fail`t$src`t$($_.Exception.Message)" }
}
$w.Quit()
'''


def to_pdf(folder):
    folder = Path(folder).resolve()
    jobs = [(p, p.with_suffix(".pdf")) for p in sorted(folder.glob("*.docx")) if ".rebuilt" not in p.name]
    lst = folder / "_to_pdf.txt"
    lst.write_text("\n".join(f"{a}\t{b}" for a, b in jobs), encoding="utf-8-sig")
    ps = folder / "_to_pdf.ps1"
    ps.write_text(PS, encoding="utf-8-sig")
    r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ps), "-List", str(lst)],
                       capture_output=True, text=True, timeout=1800)
    print(r.stdout[-3000:], r.stderr[-2000:])


# ---------------------------------------------------------------- сборка по разметке
def build(folder):
    import assemble as A
    import time
    for pdf in sorted(Path(folder).glob("*.pdf")):
        if ".rebuilt" in pdf.name:
            continue
        out = pdf.with_suffix(".rebuilt.docx")
        t0 = time.time()
        try:
            info, _ = A.build_struct_docx(pdf, out, [], {}, None)
            print(pdf.name, "ok", round(time.time() - t0, 1), "с", flush=True)
        except Exception as e:                      # noqa: BLE001 — провал сборки — тоже оценка
            print(pdf.name, "ОШИБКА", type(e).__name__, str(e)[:200], flush=True)


# ---------------------------------------------------------------- признаки DOCX
NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _norm(t):
    t = re.sub(r"^\s*(?:\d+(?:\.\d+)*\.?|[a-zа-я]\.?\)|[ivx]+\.)\s+", "", t.strip().lower())
    return re.sub(r"[^\w]", "", t)


def _style_chain(style):
    while style is not None:
        yield style
        style = style.base_style


def _eff(par, get):
    v = get(par.paragraph_format)
    if v is not None:
        return v
    for s in _style_chain(par.style):
        v = get(s.paragraph_format)
        if v is not None:
            return v
    return None


def _numpr(par):
    ppr = par._p.pPr
    n = ppr.find(NS + "numPr") if ppr is not None else None
    if n is None:
        for s in _style_chain(par.style):
            sp = s.element.find(NS + "pPr")
            n = sp.find(NS + "numPr") if sp is not None else None
            if n is not None:
                break
    if n is None:
        return None
    ilvl = n.find(NS + "ilvl")
    nid = n.find(NS + "numId")
    return (int(ilvl.get(NS + "val")) if ilvl is not None else 0, nid.get(NS + "val") if nid is not None else None)


def _numfmt(doc, nid, ilvl):
    try:
        num = doc.part.numbering_part.element
    except Exception:                               # noqa: BLE001 — нумерации нет
        return None
    for n in num.findall(NS + "num"):
        if n.get(NS + "numId") == nid:
            aid = n.find(NS + "abstractNumId").get(NS + "val")
            for an in num.findall(NS + "abstractNum"):
                if an.get(NS + "abstractNumId") == aid:
                    for l in an.findall(NS + "lvl"):
                        if l.get(NS + "ilvl") == str(ilvl):
                            f = l.find(NS + "numFmt")
                            return f.get(NS + "val") if f is not None else None
    return None


def _kind(fmt):
    return {"bullet": "bullet", "decimal": "decimal", "lowerLetter": "alpha", "upperLetter": "alpha",
            "lowerRoman": "roman", "upperRoman": "roman"}.get(fmt or "", fmt)


def _zone(part):
    """Колонтитул: вид (таблица / текст / пусто), поле номера листа, картинка, форма таблицы."""
    x = part._element.xml
    tbls = part._element.findall(".//" + NS + "tbl")
    text = "".join(t.text or "" for t in part._element.iter(NS + "t")).strip()
    shape = None
    if tbls:
        t = tbls[0]
        shape = (len(t.findall(NS + "tr")), len(t.findall(NS + "tblGrid/" + NS + "gridCol")))
    return {"kind": "table" if tbls else "text" if text else "none", "page": "PAGE" in x, "img": "<w:drawing" in x
            or "graphicData" in x, "shape": shape}


def features(path):
    from docx import Document
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    doc = Document(str(path))
    items = []
    for el in doc.element.body:
        if el.tag == NS + "p":
            p = Paragraph(el, doc)
            x = el.xml
            if "TOC \\o" in x or "TOC \\h" in x:
                m = re.search(r'TOC \\o "1-(\d)"', x)
                items.append({"k": "toc", "levels": int(m.group(1)) if m else None})
                continue
            text = p.text.strip()
            if not text and "<w:drawing" not in x:
                continue
            name = p.style.name
            m = re.match(r"Heading (\d)", name)
            npr = _numpr(p)
            d = {"k": "h" if m else "p", "text": text, "norm": _norm(text), "style": name,
                 "img": "<w:drawing" in x}
            if m:
                d["level"] = int(m.group(1))
            if npr and not m:
                d["k"], d["ilvl"] = "li", npr[0]
                d["kind"] = _kind(_numfmt(doc, npr[1], npr[0]))
            fl = _eff(p, lambda f: f.first_line_indent)
            d["first_line"] = round(fl.pt, 1) if fl is not None else 0.0
            al = _eff(p, lambda f: f.alignment)
            d["justify"] = al is not None and int(al) == 3
            items.append(d)
        elif el.tag == NS + "tbl":
            t = Table(el, doc)
            spans = sum(1 for tc in el.iter(NS + "tc") if tc.find(NS + "tcPr/" + NS + "gridSpan") is not None)
            vm = sum(1 for tc in el.iter(NS + "tc") if tc.find(NS + "tcPr/" + NS + "vMerge") is not None)
            ncol = len(el.findall(NS + "tblGrid/" + NS + "gridCol"))
            x = el.xml
            # линии: таблица целиком (стиль «Сетка», tblBorders) или по ячейкам (tcBorders)
            tb = el.find(NS + "tblPr/" + NS + "tblBorders")
            inner_v = None
            if tb is not None and tb.find(NS + "insideV") is not None:
                inner_v = tb.find(NS + "insideV").get(NS + "val") not in ("nil", "none")
            elif "TableGrid" in x and tb is None:
                inner_v = True
            cell_v = [c.get(NS + "val") for c in el.iter(NS + "left")] + [c.get(NS + "val") for c in el.iter(NS + "right")]
            if "tcBorders" in x:
                inner_v = sum(v not in ("nil", "none") for v in cell_v) > 0.5 * max(1, len(cell_v))
            texts = [_norm(c.text) for r in t.rows for c in r.cells]
            items.append({"k": "table", "rows": len(t.rows), "cols": ncol, "spans": spans, "vmerge": vm,
                          "inner_v": inner_v, "texts": texts})
    secs = []
    for s in doc.sections:
        pg = s._sectPr.find(NS + "pgNumType")
        secs.append({"header": _zone(s.header), "footer": _zone(s.footer),
                     "start": int(pg.get(NS + "start")) if pg is not None and pg.get(NS + "start") else None})
    return {"items": items, "sections": secs}


# ---------------------------------------------------------------- оценка
def score_one(src, out):
    a, b = features(src), features(out)
    A, B = a["items"], b["items"]
    S = {}
    ta = [i["norm"] for i in A if i["k"] in ("p", "h", "li") and len(i["norm"]) >= 3]
    tb = [i["norm"] for i in B if i["k"] in ("p", "h", "li") and len(i["norm"]) >= 3]
    tb_par = list(tb)                                 # порядок — по абзацам: текст ячейки, совпавший с заголовком, — не вставка
    for i in B:                                       # текст ячеек и абзацев внутри таблиц — тоже текст
        if i["k"] == "table":
            tb += [t for t in i["texts"] if len(t) >= 3]
    ca = Counter(t for i in A if i["k"] == "table" for t in i["texts"] if len(t) >= 3)
    cb = Counter(tb)
    found = sum(min(n, cb[t]) for t, n in Counter(ta).items())
    S["text"] = found / max(1, len(ta))
    S["order"] = SequenceMatcher(None, ta, [t for t in tb_par if t in set(ta)], autojunk=False).ratio()
    S["cells"] = sum(min(n, cb[t]) for t, n in ca.items()) / max(1, sum(ca.values())) if ca else None
    by = {}
    for i in B:
        by.setdefault(i.get("norm"), []).append(i)
    ha = [i for i in A if i["k"] == "h"]
    ok = sum(1 for i in ha if any(j["k"] == "h" and j.get("level") == i["level"] for j in by.get(i["norm"], [])))
    S["headings"] = ok / len(ha) if ha else None
    la = [i for i in A if i["k"] == "li"]
    ok = sum(1 for i in la if any(j["k"] == "li" and j.get("kind") == i["kind"] and j.get("ilvl") == i["ilvl"]
                                  for j in by.get(i["norm"], [])))
    S["lists"] = ok / len(la) if la else None
    # метки в тексте «a.) …» — должны стать списком (в исходнике это абзацы с меткой)
    lab = [i for i in A if i["k"] == "p" and re.match(r"^[a-f]\.\)\s", i["text"])]
    ok = sum(1 for i in lab if any(j["k"] == "li" and j.get("kind") == "alpha" for j in by.get(i["norm"], [])))
    S["labels"] = ok / len(lab) if lab else None
    toa = [i for i in A if i["k"] == "toc"]
    tob = [i for i in B if i["k"] == "toc"]
    S["toc"] = None if not toa else float(len(tob) == 1 and tob[0]["levels"] == toa[0]["levels"])
    tA = [i for i in A if i["k"] == "table"]
    tB = [i for i in B if i["k"] == "table"]
    if tA:
        sh = sum(1 for x, y in zip(tA, tB) if (x["rows"], x["cols"]) == (y["rows"], y["cols"]))
        mg = sum(1 for x, y in zip(tA, tB) if (x["spans"] > 0) == (y["spans"] > 0))
        bd = sum(1 for x, y in zip(tA, tB) if x["inner_v"] == y["inner_v"])
        S["tables_count"] = float(len(tA) == len(tB))
        S["tables_shape"] = sh / len(tA)
        S["tables_merge"] = mg / len(tA)
        S["tables_lines"] = bd / len(tA)
    body_a, body_b = a["sections"][-1], b["sections"][-1]
    for z in ("header", "footer"):
        x, y = body_a[z], body_b[z]
        S[z] = float(x["kind"] == y["kind"] and x["page"] == y["page"] and x["img"] == y["img"] and
                     (x["kind"] != "table" or x["shape"] == y["shape"]))
    S["sections"] = float(len(a["sections"]) == len(b["sections"]))
    # номер первого листа тела: без явного начала у единственного раздела — 1, у последующего — продолжение
    eff = lambda secs: secs[-1]["start"] if secs[-1]["start"] is not None else (1 if len(secs) == 1 else None)
    S["num_start"] = float(eff(a["sections"]) == eff(b["sections"]))
    # тело — без абзацев с висячим отступом (метки «a.)» текстом): сборка верно делает их пунктами списка
    body = [i for i in A if i["k"] == "p" and len(i["text"]) > 80 and i["first_line"] >= 0]
    fb = Counter(round(i["first_line"]) > 3 for i in body).most_common(1)
    rb = [j for i in body for j in by.get(i["norm"], []) if j["k"] == "p"]
    fr = Counter(round(j["first_line"]) > 3 for j in rb).most_common(1)
    S["first_line"] = float(bool(fb) and bool(fr) and fb[0][0] == fr[0][0]) if body else None
    jb = Counter(i["justify"] for i in body).most_common(1)
    jr = Counter(j["justify"] for j in rb).most_common(1)
    S["justify"] = float(bool(jb) and bool(jr) and jb[0][0] == jr[0][0]) if body else None
    return S


def lo_print(folder):
    """Пересборки → PDF LibreOffice в контейнере hut (как печатает сервис): X.rebuilt.docx → X.rebuilt.pdf."""
    import subprocess
    for out in sorted(Path(folder).glob("*.rebuilt.docx")):
        tmp = f"/tmp/bench_{out.stem}"
        subprocess.run(["docker", "cp", str(out), f"hut:{tmp}.docx"], check=True, capture_output=True)
        r = subprocess.run(["docker", "exec", "hut", "/usr/bin/python3", "/app/src/lo_pdf.py", f"{tmp}.docx", f"{tmp}.pdf"],
                           capture_output=True, text=True)
        if r.returncode == 0:
            subprocess.run(["docker", "cp", f"hut:{tmp}.pdf", str(out.with_suffix(".pdf"))], check=True, capture_output=True)
        print("ok" if r.returncode == 0 else "ОШИБКА " + r.stderr[-200:], out.name, flush=True)
        subprocess.run(["docker", "exec", "hut", "rm", "-f", f"{tmp}.docx", f"{tmp}.pdf"], capture_output=True)


def _hf_orig(pdf, n=3):
    """Колонтитулы первых листов оригинала (Word, PDF с тегами) — знаки-артефакты: шапка выше середины листа,
    подвал ниже; текст без пробелов."""
    import pdfplumber
    out = []
    with pdfplumber.open(pdf) as d:
        for pg in d.pages[:n]:
            art = [c for c in pg.chars if c.get("mcid") is None and c["text"].strip()]
            top = "".join(c["text"] for c in sorted(art, key=lambda c: (round(c["top"]), c["x0"])) if c["top"] < pg.height / 2)
            bot = "".join(c["text"] for c in sorted(art, key=lambda c: (round(c["top"]), c["x0"])) if c["top"] >= pg.height / 2)
            out.append((top.lower(), bot.lower()))
    return out


def _hf_built(pdf, docx_path, n=3):
    """Колонтитулы первых листов пересборки (LibreOffice, без тегов) — знаки в полях листа по разделам DOCX:
    1-й лист — 1-й раздел, дальше — следующий (титул отдельным разделом) или тот же."""
    import pdfplumber
    from docx import Document
    secs = Document(docx_path).sections
    out = []
    with pdfplumber.open(pdf) as d:
        for i, pg in enumerate(d.pages[:n]):
            sec = secs[0] if i == 0 or len(secs) == 1 else secs[1]
            t, bm = sec.top_margin.pt, sec.bottom_margin.pt
            ch = sorted((c for c in pg.chars if c["text"].strip()), key=lambda c: (round(c["top"]), c["x0"]))
            top = "".join(c["text"] for c in ch if c["bottom"] <= t + 1)
            bot = "".join(c["text"] for c in ch if c["top"] >= pg.height - bm - 1)
            out.append((top.lower(), bot.lower()))
    return out


def score(folder):
    folder = Path(folder)
    rows = {}
    for src in sorted(folder.glob("*.docx")):
        if ".rebuilt" in src.name:
            continue
        out = src.with_suffix(".rebuilt.docx")
        if not out.exists():
            rows[src.stem] = {"build": 0.0}
            continue
        try:
            rows[src.stem] = {"build": 1.0, **score_one(src, out)}
            # по отрисовке: шапка и подвал (с номерами) первых трёх листов — как у оригинала (Word) в печати сервиса
            # (LibreOffice): титул «особым первым колонтитулом» и титул отдельным разделом выглядят одинаково
            a_pdf, b_pdf = src.with_suffix(".pdf"), out.with_suffix(".pdf")
            if a_pdf.exists() and b_pdf.exists():
                ha, hb = _hf_orig(a_pdf), _hf_built(b_pdf, out)
                pairs = [(x, y) for pa, pb in zip(ha, hb) for x, y in zip(pa, pb)]
                rows[src.stem]["hf_pages"] = sum(x == y for x, y in pairs) / len(pairs) if pairs else None
        except Exception as e:                      # noqa: BLE001
            rows[src.stem] = {"build": 0.0, "error": f"{type(e).__name__}: {e}"[:200]}
    keys = ["build", "text", "order", "cells", "headings", "lists", "labels", "toc", "tables_count", "tables_shape",
            "tables_merge", "tables_lines", "header", "footer", "sections", "num_start", "hf_pages", "first_line",
            "justify"]
    agg = {}
    for k in keys:
        v = [r[k] for r in rows.values() if r.get(k) is not None]
        agg[k] = (round(sum(v) / len(v), 3), len(v)) if v else None
    (folder / "score.json").write_text(json.dumps({"docs": rows, "agg": agg}, ensure_ascii=False, indent=1),
                                       encoding="utf-8")
    lines = ["| признак | среднее | документов |", "|---|---|---|"]
    for k in keys:
        if agg[k]:
            lines.append(f"| {k} | {agg[k][0]:.3f} | {agg[k][1]} |")
    bad = {k: [d for d, r in rows.items() if r.get(k) is not None and r[k] < 1] for k in keys}
    lines += ["", "Провалы по признакам (документы с оценкой < 1):"]
    lines += [f"- {k}: {', '.join(v[:12])}{' …' if len(v) > 12 else ''}" for k, v in bad.items() if v]
    (folder / "score.md").write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))


if __name__ == "__main__":
    cmd, folder = sys.argv[1], Path(sys.argv[2])
    if cmd == "gen":
        folder.mkdir(parents=True, exist_ok=True)
        n = int(sys.argv[3]) if len(sys.argv) > 3 else 20
        base = int(sys.argv[4]) if len(sys.argv) > 4 else 1000
        for i in range(n):
            gen_one(base + i, folder / f"syn_{base + i}.docx")
        print("документов:", n)
    elif cmd == "pdf":
        to_pdf(folder)
    elif cmd == "build":
        build(folder)
    elif cmd == "lo":
        lo_print(folder)
    elif cmd == "score":
        score(folder)
